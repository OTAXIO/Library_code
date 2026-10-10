"""User-requested cleanup of old failed intake notes, without browser actions.

Planning is read-only. Applying requires an authored/verified candidate and the
original checksum. Only two requested cells per failed owned row are cleared;
other known exception notes are shortened without clearing their skip flags.
The original XLSX package is narrowly patched, not reserialized by the preview.
"""
from __future__ import annotations

import argparse
import json
import sys
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import HEADERS, SafetyStop, read_roster
from operation_log import OperationLog
from roster_write import cleanup_skipped_notes
from skip_notes import known_skip_note

OWNER = '谭勋策'


def failed_intake(record):
    return (record.owner == OWNER and record.matches == 0 and record.skipped and not record.done
            and record.remark.startswith(('WOS 下载未完成：', 'WOS导入失败：', '导入失败：'))
            and known_skip_note(record.remark) not in ('wos未收录', '非交大', '交大署名待核验'))


def cleanup_plan(roster):
    if not roster.status_separate or roster.remark_column != 1:
        raise SafetyStop('必须先确认首列备注与独立是否识别列；未清理。')
    roster.assert_unchanged()
    result = []
    for record in roster.records:
        if record.owner != OWNER or not record.skipped or record.done:
            continue
        short = known_skip_note(record.remark)
        if failed_intake(record):
            result.append({'sa_id': record.sa_id, 'row': record.row, 'before': record.remark,
                           'after': '', 'clear_flag': True})
        elif short and short != record.remark:
            result.append({'sa_id': record.sa_id, 'row': record.row, 'before': record.remark,
                           'after': short, 'clear_flag': False})
    return result


def apply_candidate(roster, candidate, expected_sha):
    if roster.sha256 != expected_sha:
        raise SafetyStop('名单已变化，清理计划失效；未修改。')
    plan = cleanup_plan(roster)
    intended = {change['sa_id']: change for change in plan}
    expected = [replace(record, remark=intended[record.sa_id]['after'],
                        skipped=not intended[record.sa_id]['clear_flag'])
                if record.sa_id in intended else record for record in roster.records]
    # Never use an authoring tool's reserialized metadata as task facts. Some
    # exporters turn an empty shared string into its index (e.g. "323"). Only
    # the authored note/status cells are consumed; SA IDs and all note/status
    # rows must still align. The final narrow write verifies every original fact.
    from openpyxl import load_workbook
    authored = load_workbook(candidate, read_only=True, data_only=False)
    original = load_workbook(roster.path, read_only=True, data_only=False)
    try:
        if authored.sheetnames != original.sheetnames:
            raise SafetyStop('候选工作表与名单不一致，未清理。')
        sheet = authored[roster.sheet_name]
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1,
                          max_col=roster.header_column_count))]
        before_header = [cell.value for cell in next(original[roster.sheet_name].iter_rows(
                          min_row=1, max_row=1, max_col=roster.header_column_count))]
        if header != before_header:
            raise SafetyStop('候选表头与名单不一致，未清理。')
        id_column = header.index(HEADERS['sa_id'])
        rows = {number: cells for number, cells in enumerate(sheet.iter_rows(
                min_row=2, max_col=roster.header_column_count), 2)}
        for record in expected:
            cells = rows.get(record.row)
            if cells is None:
                raise SafetyStop('候选名单缺少目标行，未清理。')
            id_cell, note_cell, flag_cell = (cells[id_column], cells[roster.remark_column - 1],
                                            cells[roster.completion_column - 1])
            flag = 1 if record.done else 2 if record.skipped else None
            if (any(cell.data_type == 'f' for cell in (id_cell, note_cell, flag_cell)) or
                    id_cell.value != record.sa_id or (note_cell.value or '') != record.remark or
                    flag_cell.value != flag or isinstance(flag_cell.value, bool)):
                raise SafetyStop('候选备注、状态或行对应关系不符合批准的范围，原名单未修改。')
    finally:
        authored.close()
        original.close()
    if not plan:
        return None
    clears = [record for record in roster.records if record.sa_id in intended and
              intended[record.sa_id]['clear_flag']]
    notes = {change['sa_id']: change['after'] for change in plan if not change['clear_flag']}
    result = cleanup_skipped_notes(roster, clears, notes)
    if result.roster.records != expected:
        raise SafetyStop('清理后的名单不符合预览，请保留备份并核验。')
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--roster', type=Path, default=Path(__file__).resolve().parents[1] / 'list.xlsx')
    parser.add_argument('--plan', type=Path, required=True)
    parser.add_argument('--apply-authored', type=Path)
    parser.add_argument('--expected-sha')
    args = parser.parse_args()
    roster = read_roster(args.roster)
    changes = cleanup_plan(roster)
    # This private audit preserves the full old diagnostic before shortening it.
    data = {'sha256': roster.sha256, 'sheet': roster.sheet_name,
            'remark_column': roster.remark_column, 'status_column': roster.completion_column,
            'changes': changes}
    if args.apply_authored:
        if not args.expected_sha:
            parser.error('--apply-authored requires --expected-sha')
        result = apply_candidate(roster, args.apply_authored, args.expected_sha)
        data['after_sha256'] = result.roster.sha256 if result else roster.sha256
        data['backup'] = str(result.backup) if result else None
        data['applied'] = bool(result)
        if result:
            log = OperationLog(roster.path.parent / 'log.txt')
            for change in changes:
                action = ('清理旧下载/导入失败标记：' if change['clear_flag'] else '简化跳过原因：') + change['before']
                log.record(action, '已执行', change['sa_id'])
    args.plan.parent.mkdir(parents=True, exist_ok=True)
    args.plan.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'cleared': sum(c['clear_flag'] for c in changes),
                      'shortened': sum(not c['clear_flag'] for c in changes),
                      'applied': data.get('applied', False)}, ensure_ascii=False))


if __name__ == '__main__':
    main()

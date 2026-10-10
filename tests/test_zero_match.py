"""Disposable rosters, fake WOS and fake production UI; never user data."""
import copy
import tempfile
import threading
import unittest
from unittest.mock import patch
from pathlib import Path

from openpyxl import load_workbook
from automation import ImportStore
from core import HEADERS, SafetyStop, file_hash, read_roster
from tests.test_automation import sample
from tests.test_roster_write import make_roster
from zero_match import WorkflowStore, run_batch, select_records, resolved_item


class Backend:
    def __init__(self, records):
        self.rows = {r.sa_id: {"saLzkId": r.sa_id, "titleValue": r.title,
            "doiValue": r.doi, "wosValue": r.wos, "gh": r.staff_id,
            "markStatus": "待处理", "matchCount": 0, "itemId": "", "reason": "", "remark": ""} for r in records}
        self.items = []
        self.calls = []
        self.phase = 0
        self.candidate = None
        self.lose = None
        self.author_issue = False
        self.comparison_override = {}
        self.ambiguous_author = False

    def call(self, action, payload, timeout=75):
        self.calls.append(action)
        row = self.rows[payload["sa_id"]]
        if action == "import_capabilities":
            return {"zero_match_protocol": 1, "library_resolution": True}
        if action in ("search", "status"):
            fields = [] if not row["itemId"] else [
                {"label": label, "sa": row[source], "library": self.candidate[key]}
                for label, key, source in (("题名", "title", "titleValue"), ("DOI", "doi", "doiValue"), ("WOS记录号", "wos", "wosValue"))]
            if fields:
                info = f"工号：{row['gh']}\n是否第一作者：否\n是否通讯作者：否"
                fields += [{"label": "作者信息", "sa": info, "library": info},
                           {"label": "认领状态", "sa": f"测试员({row['gh']})", "library": "已认领"},
                           {"label": "交大是否第一单位", "sa": "是", "library": "是"}]
                for field in fields:
                    field.update(self.comparison_override.get(field['label'], {}))
            return {"row": copy.deepcopy(row), "comparison": fields, "non_sjtu_completion_protocol": 1,
                    "import_completion_protocol": 1}
        if action == "import_resolve":
            return {"verified": True, "items": copy.deepcopy(self.items)}
        if action == 'prepare_claim':
            current = self.call('search', payload)
            prepared = {'staff_id': row['gh'], 'sa_text': payload['sa_text'],
                        'person': {'id': 'scholar-1', 'wno': row['gh'], 'name': '测试员', 'names': ['Test, Alice']},
                        'authors': [{'index': 0, 'order': 2, 'fullname': 'Test, Alice', 'eligible': True, 'scholarId': ''}]}
            return {**current, 'prepared': prepared, 'suggested_index': None if self.ambiguous_author else 0}
        if action == 'submit_claim':
            self.comparison_override['认领状态'] = {'library': '已认领'}
            if self.lose == action:
                self.lose = None
                raise SafetyStop('认领已提交，回传丢失')
            return {'verified': True, 'claimed': True, 'row': copy.deepcopy(row), 'staff_id': row['gh'],
                    'scholar_id': 'scholar-1', 'author': 'Test, Alice', 'order': 2}
        if action == "import_scan":
            return {"batches": []}
        if action == "import_upload":
            self.candidate = payload["candidate"]
            return {"uploaded": True, "sha256": self.candidate["sha256"]}
        if action == "import_submit":
            self.phase = 1
            return {"submitted": True}
        if action == "import_check":
            return {"verified": True, "batch": {"id": "batch-1", "status": self.phase,
                "batchNumber": "B-1", "modelId": "article", "source": "WOS",
                "instructions": "SA补充-" + payload["sa_id"], "total": 1,
                "actual": 1 if self.phase == 2 else 0, "fail": 0}}
        if action == "import_push":
            self.phase = 2
            c = payload["candidate"]
            self.items = [{"id": "1234567890123456789", **{k: c[k] for k in ("title", "doi", "wos")}}]
            return {"submitted": True}
        if action == "link":
            assert payload["expected"] == row
            row.update(itemId=payload["item_id"], matchCount=1)
            if self.author_issue:
                row["reason"] = "通讯作者不一致"
        elif action == "complete":
            assert payload["expected"] == row
            row.update(markStatus="已处理", remark=payload["note"])
        else:
            raise AssertionError(action)
        if self.lose == action:
            self.lose = None
            raise SafetyStop("已发出写入，回传结果丢失")
        return {"verified": True, "row": copy.deepcopy(row)}


class Download:
    def __init__(self, folder):
        self.store = ImportStore(folder / "downloads")
        self.path = folder / "record.txt"
        self.path.write_bytes(sample())
        self.fail = {}
        self.calls = []

    def call(self, action, payload, timeout=120):
        self.calls.append(action)
        if payload["sa_id"] in self.fail:
            raise SafetyStop(self.fail[payload["sa_id"]])
        return {"path": str(self.path), "sa_id": payload["sa_id"], "ready": True,
                "record_url": "https://www.webofscience.com/wos/woscc/full-record/WOS:000123456789012"}


class ZeroMatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.path = self.folder / "list.xlsx"
        make_roster(self.path, flags=[None] * 3)
        book = load_workbook(self.path)
        sheet = book["名单"]
        sheet.cell(1, 14, "数据来源")
        sheet.cell(1, 15, "是否识别")
        for row in range(2, 5):
            for key, value in {"owner": "谭勋策", "title": "Synthetic paper", "doi": "10.1234/test",
                               "matches": 0, "item_ids": "", "reason": ""}.items():
                sheet.cell(row, 2 + list(HEADERS).index(key)).value = value
        sheet.cell(4, 2, "其他人")
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)
        self.sa = Backend(self.roster.records)
        self.download = Download(self.folder)
        self.imports = ImportStore(self.folder / "imports")
        self.workflow = WorkflowStore(self.folder / "workflow")
        self.stop = threading.Event()

    def run_flow(self, limit=2):
        records = select_records(self.roster, "谭勋策", limit)
        result = run_batch(self.roster, records, self.sa, self.download, self.imports, self.workflow, self.stop)
        self.roster = result.roster
        return result

    def existing_paper(self):
        from automation import parse_wos
        self.sa.candidate = parse_wos(sample())
        self.sa.items = [{"id": "1234567890123456789",
            **{key: self.sa.candidate[key] for key in ("title", "doi", "wos")}}]

    def test_confirmed_intake_uses_exact_imported_note_and_closes_without_claim(self):
        self.sa.comparison_override = {'认领状态': {'library': '未认领'},
                                      '交大是否第一单位': {'library': '未知'}}
        self.sa.author_issue = True
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        record = result.roster.records[0]
        self.assertTrue(record.done)
        self.assertFalse(record.skipped)
        self.assertEqual(record.remark, '已入库')
        self.assertEqual(self.sa.rows[record.sa_id]['markStatus'], '已处理')
        self.assertEqual(self.sa.rows[record.sa_id]['remark'], '已入库')
        self.assertNotIn('prepare_claim', self.sa.calls)
        self.assertNotIn('submit_claim', self.sa.calls)
        self.assertEqual(self.sa.calls.count('complete'), 1)
        self.assertEqual(self.workflow.get(record)['phase'], 'done')

    def test_intake_closure_requires_successful_batch_readback(self):
        original = self.sa.call
        checks = 0
        invalid_batch = {}
        def call(action, payload, timeout=75):
            nonlocal checks
            answer = original(action, payload, timeout)
            if action == 'import_check':
                checks += 1
                if checks >= 3:
                    answer['batch'].update(invalid_batch)
            return answer
        self.sa.call = call
        for invalid_batch in ({'fail': 1}, {'status': 1}, {'total': 2, 'actual': 2},
                              {'actual': 0}, {'source': 'CNKI'}, {'instructions': 'SA补充-other'},
                              {'id': 'another-batch'}, {'total': True}):
            with self.subTest(batch=invalid_batch):
                result = self.run_flow(1)
                self.assertTrue(result.halted)
                self.assertFalse(result.roster.records[0].done)
                self.assertNotIn('complete', self.sa.calls)
                self.assertEqual(self.sa.calls.count('import_submit'), 1)

    def test_intake_closure_refuses_old_extension_without_protocol(self):
        original = self.sa.call
        def call(action, payload, timeout=75):
            answer = original(action, payload, timeout)
            answer.pop('import_completion_protocol', None)
            return answer
        self.sa.call = call
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertNotIn('complete', self.sa.calls)
        self.assertFalse(result.roster.records[0].done)
        self.assertNotIn('import_upload', self.sa.calls)

    def test_explicit_retry_finishes_previously_imported_skipped_item_without_reimport(self):
        original = self.sa.call
        checks = 0
        def call(action, payload, timeout=75):
            nonlocal checks
            answer = original(action, payload, timeout)
            if action == 'import_check':
                checks += 1
                if checks == 3:
                    raise SafetyStop('模拟入库结果回读暂不可用')
            return answer
        self.sa.call = call
        self.assertTrue(self.run_flow(1).halted)
        from roster_write import mark_skipped_many
        record = self.roster.records[0]
        self.roster = mark_skipped_many(self.roster, [record],
            reasons={record.sa_id: '第一单位待核验'}).roster
        record = self.roster.records[0]
        self.sa.rows[record.sa_id]['reason'] = '第一单位不一致'
        self.sa.comparison_override['交大是否第一单位'] = {'library': '未知'}
        self.sa.call = original
        result = run_batch(self.roster, [record], self.sa, self.download, self.imports,
                           self.workflow, self.stop, retry_skipped=True)
        self.assertFalse(result.halted, result.reason)
        self.assertEqual(result.roster.records[0].remark, '已入库')
        self.assertTrue(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)
        for action in ('import_upload', 'import_submit', 'import_push', 'complete'):
            self.assertEqual(self.sa.calls.count(action), 1)
        self.assertNotIn('submit_claim', self.sa.calls)

    def test_full_two_row_pipeline_one_download_one_import_two_closures(self):
        result = self.run_flow()
        self.assertFalse(result.halted, result.reason)
        self.assertEqual(result.remaining, 0)
        self.assertEqual([r.done for r in result.roster.records], [True, True, False])
        for action, count in (("wos_search", 1), ("wos_export", 1)):
            self.assertEqual(self.download.calls.count(action), count)
        for action, count in (("import_upload", 1), ("import_submit", 1), ("import_push", 1), ("complete", 2)):
            self.assertEqual(self.sa.calls.count(action), count)
        book = load_workbook(self.path)
        self.assertEqual(book["名单"]["O2"].value, 1)
        self.assertEqual(book["名单"]["A2"].value, "已入库")
        self.assertEqual(book["名单"]["N2"].value, "WOS")
        self.assertEqual(book["名单"]["M2"].value, "网页备注保留")
        self.assertIsNone(book["名单"]["O4"].value)
        book.close()

    def test_executor_does_not_process_already_skipped_or_done_even_from_explicit_queue(self):
        from dataclasses import replace
        skipped = replace(self.roster.records[0], skipped=True, remark="已跳过原因")
        done = replace(self.roster.records[1], done=True)
        self.roster.records[:2] = [skipped, done]
        before = file_hash(self.path)
        result = run_batch(self.roster, [skipped, done], self.sa, self.download,
                           self.imports, self.workflow, self.stop)
        self.assertFalse(result.halted)
        self.assertEqual(result.remaining, 0)
        self.assertEqual([x["status"] for x in result.outcomes], ["already_skipped", "already_done"])
        self.assertEqual(self.sa.calls, [])
        self.assertEqual(self.download.calls, [])
        self.assertEqual(file_hash(self.path), before)

    def test_verification_pauses_without_marking_current_or_tail_two(self):
        before = file_hash(self.path)
        self.download.fail["demo-001"] = "WOS 显示人工验证"
        result = self.run_flow()
        self.assertTrue(result.halted)
        self.assertEqual(result.remaining, 2)
        self.assertEqual(file_hash(self.path), before)
        self.assertNotIn("import_upload", self.sa.calls)
        self.assertEqual(self.sa.calls, ["status", "import_capabilities"])

    def test_zero_result_marks_only_current_two_and_continues(self):
        self.download.fail["demo-001"] = "WOS 未找到记录"
        result = self.run_flow()
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark, "wos未收录")
        self.assertTrue(result.roster.records[1].done)

    def test_extension_zero_result_saves_requested_note_and_flag_before_next_paper(self):
        self.download.fail["demo-001"] = (
            "[扩展 0.4.1] [WOS 已暂停] WOS 未找到记录；这不等于未发表，也不自动标记完成")
        result = self.run_flow()
        self.assertFalse(result.halted, result.reason)
        self.assertEqual(result.remaining, 0)
        self.assertEqual([outcome["status"] for outcome in result.outcomes], ["skip", "done"])
        self.assertEqual(result.outcomes[0]["message"], "wos未收录")
        self.assertIn("WOS 未找到记录", result.outcomes[0]["detail"])
        self.assertFalse(result.roster.records[0].done)
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        self.assertEqual(book["名单"]["A2"].value, "wos未收录")
        self.assertEqual(book["名单"]["O2"].value, 2)
        self.assertEqual(book["名单"]["O3"].value, 1)
        self.assertIsNone(book["名单"]["A4"].value)
        self.assertIsNone(book["名单"]["O4"].value)

    def test_refresh_timeout_after_zero_preserves_current_and_unexecuted_records(self):
        self.download.fail["demo-001"] = "[WOS 已暂停] WOS 未找到记录；不自动标记完成"
        self.download.fail["demo-002"] = (
            "WOS 刷新后的检索页尚未就绪，上一条提示未清除或输入区仍在加载；未提交当前论文检索")
        result = self.run_flow()
        self.assertTrue(result.halted)
        self.assertEqual(result.remaining, 1)
        self.assertEqual([outcome["status"] for outcome in result.outcomes], ["skip", "halted"])
        book = load_workbook(self.path)
        self.addCleanup(book.close)
        self.assertEqual(book["名单"]["A2"].value, "wos未收录")
        self.assertEqual(book["名单"]["O2"].value, 2)
        for row in (3, 4):
            self.assertIsNone(book["名单"].cell(row, 1).value)
            self.assertIsNone(book["名单"].cell(row, 15).value)
        self.assertNotIn("wos_export", self.download.calls)
        self.assertNotIn("import_upload", self.sa.calls)

    def test_skip_save_failure_retains_prior_completion_and_stops_tail(self):
        records = select_records(self.roster, '谭勋策', 2)
        self.download.fail['demo-002'] = 'WOS 未找到记录'
        # Different genuine identity prevents a shared-file shortcut for row 2.
        from dataclasses import replace
        second = replace(records[1], title='Other synthetic paper')
        book = load_workbook(self.path)
        book['名单'].cell(second.row, 2 + list(HEADERS).index('title'), second.title)
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)
        self.sa.rows[second.sa_id]['titleValue'] = second.title
        with patch('zero_match.mark_skipped_many', side_effect=SafetyStop('synthetic locked workbook')):
            result = self.run_flow()
        self.assertTrue(result.halted)
        self.assertTrue(result.roster.records[0].done)
        self.assertFalse(result.roster.records[1].skipped)
        self.assertEqual(result.remaining, 1)

    def test_audit_failure_does_not_turn_a_verified_write_into_an_unknown_result(self):
        def broken_log(*args):
            raise OSError('synthetic logging failure')
        result = run_batch(self.roster, select_records(self.roster, '谭勋策', 1), self.sa,
                           self.download, self.imports, self.workflow, self.stop, audit=broken_log)
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].done)

    def test_non_sjtu_closes_website_and_roster_without_import(self):
        self.download.path.write_bytes(sample(AF="Test, Alice", C1="[Test, Alice] Other University, China"))
        result = self.run_flow()
        self.assertFalse(result.halted, result.reason)
        self.assertEqual(result.remaining, 0)
        self.assertEqual([e['status'] for e in result.outcomes], ['done', 'done'])
        self.assertEqual(self.sa.calls.count('complete'), 2)
        for record in result.roster.records[:2]:
            self.assertTrue(record.done)
            self.assertFalse(record.skipped)
            self.assertEqual(record.remark, "非交大")
            self.assertEqual(record.source, 'WOS')
            self.assertEqual(self.sa.rows[record.sa_id]['markStatus'], '已处理')
            self.assertEqual(self.sa.rows[record.sa_id]['remark'], '非交大')
            self.assertEqual(self.workflow.get(record)['phase'], 'non_sjtu_done')
        for action in ('import_resolve', 'import_upload', 'import_submit', 'import_push', 'link', 'submit_claim'):
            self.assertNotIn(action, self.sa.calls)
        book = load_workbook(self.path)
        self.assertEqual([book['名单'][cell].value for cell in ('A2', 'A3', 'O2', 'O3')], ['非交大', '非交大', 1, 1])
        self.assertEqual(book['名单']['M2'].value, '网页备注保留')
        self.assertIsNone(book['名单']['O4'].value)
        book.close()

    def non_sjtu(self):
        self.download.path.write_bytes(sample(AF='Test, Alice', C1='[Test, Alice] Other University, China'))

    def test_non_sjtu_lost_ack_reconciles_only_without_a_second_write(self):
        self.non_sjtu()
        self.sa.lose = 'complete'
        first = self.run_flow(1)
        self.assertTrue(first.halted)
        self.assertFalse(first.roster.records[0].done)
        self.assertFalse(first.roster.records[0].skipped)
        self.assertEqual(first.roster.records[0].source, 'WOS')
        self.assertEqual(self.workflow.get(first.roster.records[0])['phase'], 'non_sjtu_complete_intent')
        second = self.run_flow(1)
        self.assertFalse(second.halted, second.reason)
        self.assertTrue(second.roster.records[0].done)
        self.assertEqual(second.roster.records[0].remark, '非交大')
        self.assertEqual(second.roster.records[0].source, 'WOS')
        self.assertEqual(self.sa.calls.count('complete'), 1)
        self.assertNotIn('import_upload', self.sa.calls)

    def test_non_sjtu_unknown_write_stops_tail_and_never_resubmits(self):
        self.non_sjtu()
        original = self.sa.call
        def call(action, payload, timeout=75):
            if action == 'complete':
                self.sa.calls.append(action)
                raise SafetyStop('写入效果未知')
            return original(action, payload, timeout)
        self.sa.call = call
        first = self.run_flow()
        self.assertTrue(first.halted)
        self.assertEqual(first.remaining, 2)
        self.assertFalse(any(r.done or r.skipped for r in first.roster.records[:2]))
        self.assertTrue(self.run_flow().halted)
        self.assertEqual(self.sa.calls.count('complete'), 1)
        self.assertEqual(self.download.calls.count('wos_export'), 1)

    def test_non_sjtu_old_extension_stops_before_intent_or_write(self):
        self.non_sjtu()
        original = self.sa.call
        def call(action, payload, timeout=75):
            value = original(action, payload, timeout)
            value.pop('non_sjtu_completion_protocol', None)
            return value
        self.sa.call = call
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertIn('0.4.5', result.reason)
        self.assertEqual(self.workflow.get(self.roster.records[0]), {})
        self.assertNotIn('complete', self.sa.calls)
        self.assertFalse(result.roster.records[0].skipped)

    def test_non_sjtu_keeps_existing_website_remark(self):
        self.non_sjtu()
        self.sa.rows['demo-001']['remark'] = '其他人工结论'
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertNotIn('complete', self.sa.calls)
        self.assertEqual(self.sa.rows['demo-001']['remark'], '其他人工结论')

    def test_non_sjtu_local_source_failure_stops_before_website_write(self):
        self.non_sjtu()
        with patch('zero_match.record_data_sources', side_effect=SafetyStop('名单正在使用')):
            result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertNotIn('complete', self.sa.calls)
        self.assertEqual(self.workflow.get(self.roster.records[0]), {})
        self.assertFalse(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].source, '')

    def test_non_sjtu_changed_match_before_closure_does_not_write(self):
        self.non_sjtu()
        original = self.sa.call
        reads = 0
        def call(action, payload, timeout=75):
            nonlocal reads
            if action == 'status':
                reads += 1
                if reads == 2:
                    self.sa.rows[payload['sa_id']]['matchCount'] = 1
            return original(action, payload, timeout)
        self.sa.call = call
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertNotIn('complete', self.sa.calls)
        self.assertFalse(result.roster.records[0].done)

    def test_non_sjtu_bad_readback_never_marks_workbook_done(self):
        self.non_sjtu()
        original = self.sa.call
        def call(action, payload, timeout=75):
            value = original(action, payload, timeout)
            if action == 'complete':
                value['row']['remark'] = ''
            return value
        self.sa.call = call
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertFalse(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertEqual(self.sa.calls.count('complete'), 1)

    def test_unknown_affiliation_is_still_skipped_not_closed(self):
        self.download.path.write_bytes(sample(AF='Test, Alice; Other, Bob', C1='[Test, Alice] Other University, China'))
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertFalse(result.roster.records[0].done)
        self.assertEqual(result.roster.records[0].remark, '交大署名待核验')
        self.assertNotIn('complete', self.sa.calls)

    def test_title_only_non_sjtu_is_never_closed(self):
        self.non_sjtu()
        book = load_workbook(self.path)
        book['名单'].cell(2, 2 + list(HEADERS).index('doi')).value = ''
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)
        self.sa = Backend(self.roster.records)
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertFalse(result.roster.records[0].done)
        self.assertNotIn('complete', self.sa.calls)

    def test_remote_processed_is_synced_before_download(self):
        self.sa.rows["demo-001"].update(markStatus="已处理", remark="原网页核验结论", titleValue="人工已纠正题名")
        result = self.run_flow(1)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(self.download.calls, [])
        self.assertEqual(self.sa.calls, ["status"])

    def test_lost_link_ack_resumes_readback_without_second_link_or_import(self):
        self.sa.lose = "link"
        first = self.run_flow(1)
        self.assertTrue(first.halted)
        second = self.run_flow(1)
        self.assertFalse(second.halted, second.reason)
        self.assertTrue(second.roster.records[0].done)
        self.assertEqual(self.sa.calls.count("link"), 1)
        self.assertEqual(self.sa.calls.count("import_submit"), 1)

    def test_lost_completion_ack_only_reconciles_excel_on_restart(self):
        self.sa.lose = "complete"
        self.assertTrue(self.run_flow(1).halted)
        self.assertFalse(self.roster.records[0].done)
        result = self.run_flow(1)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(self.sa.calls.count("complete"), 1)

    def test_unconfirmed_complete_intent_does_not_resubmit(self):
        original = self.sa.call
        def call(action, payload, timeout=75):
            if action == "complete":
                self.sa.calls.append(action)
                raise SafetyStop("提交效果不明")
            return original(action, payload, timeout)
        self.sa.call = call
        self.assertTrue(self.run_flow(1).halted)
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertEqual(self.sa.calls.count("complete"), 1)

    def test_author_issue_after_link_is_not_falsely_closed(self):
        self.existing_paper()  # Existing library paper is not proof of a new import.
        self.sa.author_issue = True
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        self.assertFalse(result.roster.records[0].done)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertNotIn("complete", self.sa.calls)

    def test_unclaimed_precise_person_is_claimed_before_completion(self):
        self.existing_paper()
        self.sa.comparison_override['认领状态'] = {'library': '未认领'}
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(self.sa.calls.count('submit_claim'), 1)
        self.assertLess(self.sa.calls.index('submit_claim'), self.sa.calls.index('complete'))

    def test_unclaimed_ambiguous_person_is_skipped_not_claimed(self):
        self.existing_paper()
        self.sa.comparison_override['认领状态'] = {'library': '未认领'}
        self.sa.ambiguous_author = True
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertNotIn('submit_claim', self.sa.calls)
        self.assertNotIn('complete', self.sa.calls)

    def test_lost_claim_ack_resumes_only_by_readback(self):
        self.existing_paper()
        self.sa.comparison_override['认领状态'] = {'library': '未认领'}
        self.sa.lose = 'submit_claim'
        self.assertTrue(self.run_flow(1).halted)
        result = self.run_flow(1)
        self.assertFalse(result.halted, result.reason)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(self.sa.calls.count('submit_claim'), 1)

    def test_role_or_unit_difference_is_not_closed_with_empty_reason(self):
        self.existing_paper()
        staff = self.roster.records[0].staff_id
        for label, value in [('作者信息', f'工号：{staff}\n是否第一作者：是\n是否通讯作者：否'),
                             ('交大是否第一单位', '否')]:
            with self.subTest(label=label):
                self.sa.comparison_override = {label: {'library': value}}
                result = self.run_flow(1)
                self.assertFalse(result.halted, result.reason)
                self.assertTrue(result.roster.records[0].skipped)
                self.assertNotIn('complete', self.sa.calls)
                self.assertEqual(result.roster.records[0].remark,
                    '第一单位待核验' if label == '交大是否第一单位' else '作者角色待核验')
                self.assertIn('不自动', result.outcomes[0]['detail'])
                # Explicit retry is separate from the default pending queue.
                from roster_write import clear_skipped_many
                self.roster = clear_skipped_many(self.roster, [self.roster.records[0]]).roster

    def test_lookup_rejects_numeric_platform_id_or_conflicting_ut(self):
        from automation import parse_wos
        c = parse_wos(sample())
        for changed in ({"id": 1234567890123456789}, {"wos": "WOS:999999999999999"}):
            item = {"id": "123", **{k: c[k] for k in ("title", "doi", "wos")}, **changed}
            with self.assertRaises(SafetyStop):
                resolved_item({"verified": True, "items": [item]}, c)

    def test_stale_sa_comparison_blocks_completion_even_if_list_snapshot_is_unchanged(self):
        self.sa.comparison_override['DOI'] = {'sa': '10.1234/different'}
        result = self.run_flow(1)
        self.assertTrue(result.halted)
        self.assertNotIn('complete', self.sa.calls)
        self.assertFalse(result.roster.records[0].done)


if __name__ == "__main__":
    unittest.main()

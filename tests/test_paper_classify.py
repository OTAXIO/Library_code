import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl

from core import SafetyStop, file_hash
from paper_classify import (ClassificationClient, DEFAULT_MODEL, MAX_CONSECUTIVE_FAILURES,
                            read_papers, run, validate_result, InvalidClassification,
                            ModelRequestError)


def answer(paper, **changes):
    return dict({'id':paper['id'],'type':'期刊论文','databases':[],
                 'confidence':'低','reason':'仅有题名，需要出版依据。',
                 'missing_evidence':['刊名'],'evidence_ids':[],
                 'import_channel':'WOS','channel_reason':'候选渠道，需核实收录并获取对应导出文件。'},**changes)


class ClassificationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root/'list.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(['负责人','题名','DOI','工号'])
        sheet.append(['private-owner','Paper | A','10.1000/a','secret-staff'])
        sheet.append(['private-owner','Paper | A','10.1000/a','secret-staff'])
        sheet.append(['private-owner','Paper | A','10.1000/b','secret-staff'])
        sheet.append(['private-owner','Paper | A',None,'secret-staff'])
        sheet.append(['private-owner','Paper B',None,'secret-staff'])
        book.create_sheet('统计').append(['负责人','数量'])
        book.save(self.path)
        book.close()
        self.papers = read_papers(self.path)['papers']

    def client(self):
        return Mock(classify=Mock(side_effect=lambda papers,model:{p['id']:answer(p) for p in papers}))

    def test_read_groups_only_matching_title_and_doi_preserves_rows(self):
        self.assertEqual(len(self.papers),4)
        self.assertEqual(self.papers[0]['rows'],[2,3])
        self.assertEqual(sorted(r for p in self.papers for r in p['rows']),list(range(2,7)))
        encoded=json.dumps(self.papers)
        self.assertNotIn('private-owner',encoded)
        self.assertNotIn('secret-staff',encoded)

    def test_reject_formula_and_nonempty_row_without_title(self):
        for value in ('=A1',None):
            book=openpyxl.load_workbook(self.path)
            book.active['B2']=value
            book.save(self.path)
            book.close()
            with self.assertRaises(SafetyStop):
                read_papers(self.path)

    def test_reject_missing_extra_duplicate_and_unknown_results(self):
        good=answer(self.papers[0])
        for values in ([],[good,good],[answer(self.papers[1])],[dict(good,type='SCI论文')],
                       [dict(good,databases=['WOS'])],[dict(good,evidence_ids=['invented'])],
                       [dict(good,confidence='100%')]):
            with self.subTest(values=values),self.assertRaises(SafetyStop):
                validate_result(json.dumps({'results':values}),self.papers[:1])

    def test_supplied_evidence_is_required_for_database_suggestion(self):
        paper={**self.papers[0],'evidence':[{'id':'E1','text':'用户提供的收录摘录'}]}
        value=answer(paper,databases=['DBLP'],evidence_ids=['E1'])
        self.assertEqual(validate_result(json.dumps({'results':[value]}),[paper])[paper['id']]['databases'],['DBLP'])

    def test_undeclared_fields_discarded_but_missing_required_still_rejected(self):
        paper=self.papers[0]
        value=answer(paper,confidence_note='中',execute='do something',verified=True)
        cleaned=validate_result(json.dumps({'results':[value]}),[paper])[paper['id']]
        self.assertEqual(cleaned,answer(paper))
        del value['channel_reason']
        with self.assertRaises(SafetyStop):
            validate_result(json.dumps({'results':[value]}),[paper])

    def test_import_channel_not_database_or_file_format(self):
        for channel in ('PUBMED','SCOPUS','DBLP','WOS数据导入(Excel)',{},'SPO'):
            with self.subTest(channel=channel),self.assertRaises(SafetyStop):
                validate_result(json.dumps({'results':[answer(self.papers[0],import_channel=channel)]}),self.papers[:1])
        from import_channels import route_fields, CHANNELS
        self.assertEqual(len(CHANNELS),9)
        self.assertEqual(sum(map(len,CHANNELS.values())),10)
        value=answer(self.papers[0])
        route=route_fields(value)
        self.assertEqual(len(route['available_buttons']),2)
        self.assertIsNone(route['selected_button'])
        self.assertEqual(value['databases'],[])
        self.assertIsNone(route_fields(None))

    def test_transport_sends_only_allowed_fields_and_rejects_truncation(self):
        client=ClassificationClient(Mock())
        reply={'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'results':[answer(self.papers[0])]}),'reasoning_content':'private-chain'}}]}
        client._request=Mock(return_value=reply)
        client.classify(self.papers[:1],DEFAULT_MODEL)
        args=client._request.call_args.args
        self.assertEqual(args[:2],('POST','/api/v1/chat/completions'))
        sent=json.loads(args[2]['messages'][1]['content'])['papers'][0]
        self.assertEqual(set(sent),{'id','title','doi'})
        reply['choices'][0]['finish_reason']='length'
        with self.assertRaises(SafetyStop):
            client.classify(self.papers[:1],DEFAULT_MODEL)

    def test_full_run_and_resume_do_not_recall_or_change_source(self):
        before=file_hash(self.path)
        client=self.client()
        folder=run(self.path,self.root/'out',client=client,batch_size=5)
        self.assertEqual(client.classify.call_count,1)
        client.classify.reset_mock()
        self.assertEqual(run(self.path,self.root/'out',client=client),folder)
        client.classify.assert_not_called()
        self.assertEqual(file_hash(self.path),before)
        exported=json.loads((folder/'分类结果.json').read_text(encoding='utf-8'))
        self.assertEqual(len(exported['records']),4)
        self.assertTrue(all(p['status']=='AI建议待复核' for p in exported['records']))
        self.assertTrue(all(p['import_route']['recommended_channel']=='WOS' for p in exported['records']))
        self.assertTrue(all(p['import_route']['selected_button'] is None for p in exported['records']))
        self.assertIn('WOS数据导入(Excel)',(folder/'导入渠道分类.md').read_text(encoding='utf-8'))
        self.assertIn('Paper &#124; A',(folder/'分类建议.md').read_text(encoding='utf-8'))
        self.assertFalse((folder/'run.lock').exists())

    def test_cancel_saves_current_batch_resume_only_remaining(self):
        stop=threading.Event()
        client=self.client()
        def classify(papers,model):
            stop.set()
            return {p['id']:answer(p) for p in papers}
        client.classify.side_effect=classify
        folder=run(self.path,self.root/'out',client=client,stop=stop,batch_size=2)
        saved=json.loads((folder/'progress.json').read_text(encoding='utf-8'))
        self.assertEqual(len(saved['results']),2)
        fresh=self.client()
        run(self.path,self.root/'out',client=fresh)
        self.assertEqual(len(fresh.classify.call_args.args[0]),2)

    def test_failed_batch_not_cached_and_lock_released(self):
        client=Mock(classify=Mock(side_effect=SafetyStop('synthetic failure')))
        with self.assertRaises(SafetyStop):
            run(self.path,self.root/'out',client=client)
        folder=next((self.root/'out').iterdir())
        self.assertFalse((folder/'progress.json').exists())
        self.assertFalse((folder/'run.lock').exists())
        self.assertTrue((folder/'分类建议.md').exists())
        run(self.path,self.root/'out',client=self.client())

    def test_changed_model_and_evidence_get_new_checkpoints(self):
        client=self.client()
        first=run(self.path,self.root/'out',client=client)
        second=run(self.path,self.root/'out',model='qwen',client=client)
        evidence=self.root/'evidence.json'
        evidence.write_text(json.dumps({self.papers[0]['id']:[{'id':'E1','text':'出版材料'}]}),encoding='utf-8')
        third=run(self.path,self.root/'out',client=client,evidence_path=evidence)
        self.assertEqual(len({first,second,third}),3)

    def test_tampered_checkpoint_rejected(self):
        folder=run(self.path,self.root/'out',client=self.client())
        checkpoint=folder/'progress.json'
        data=json.loads(checkpoint.read_text(encoding='utf-8'))
        data['results'][self.papers[0]['id']]['id']=self.papers[1]['id']
        checkpoint.write_text(json.dumps(data),encoding='utf-8')
        client=self.client()
        with self.assertRaises(SafetyStop):
            run(self.path,self.root/'out',client=client)
        client.classify.assert_not_called()

    def test_large_checkpoint_resumes_without_api_response_size_limit(self):
        book=openpyxl.Workbook()
        book.active.append(['题名','DOI'])
        for i in range(112):
            book.active.append([f'Paper {i}',None])
        book.save(self.path)
        book.close()
        client=self.client()
        client.classify.side_effect=lambda papers,model:{p['id']:answer(p,reason='依据'*250,channel_reason='条件'*250) for p in papers}
        stop=Mock(is_set=Mock(return_value=False),wait=Mock(return_value=False))
        folder=run(self.path,self.root/'out',client=client,batch_size=10,stop=stop)
        self.assertGreater((folder/'progress.json').stat().st_size,100000)
        fresh=self.client()
        run(self.path,self.root/'out',client=fresh)
        fresh.classify.assert_not_called()
        self.assertEqual(json.loads((folder/'status.json').read_text(encoding='utf-8'))['completed'],112)

    def test_same_task_lock_prevents_second_run(self):
        folder=run(self.path,self.root/'out',client=self.client())
        (folder/'run.lock').write_text('active',encoding='utf-8')
        client=self.client()
        with self.assertRaises(SafetyStop):
            run(self.path,self.root/'out',client=client)
        client.classify.assert_not_called()
        self.assertTrue((folder/'run.lock').exists())

    def test_auto_splits_bad_batch_and_isolates_one_failure(self):
        bad_id=self.papers[1]['id']
        client=self.client()
        def classify(papers,model):
            if any(p['id']==bad_id for p in papers):
                raise InvalidClassification('bad shape')
            return {p['id']:answer(p) for p in papers}
        client.classify.side_effect=classify
        stop=Mock(is_set=Mock(return_value=False),wait=Mock(return_value=False))
        folder=run(self.path,self.root/'out',client=client,resilient=True,stop=stop)
        status=json.loads((folder/'status.json').read_text(encoding='utf-8'))
        self.assertEqual((status['state'],status['completed'],status['failed']),('partial',3,1))
        self.assertLessEqual(client.classify.call_count,7)
        self.assertIn('分类失败',(folder/'分类建议.md').read_text(encoding='utf-8'))
        fresh=self.client()
        run(self.path,self.root/'out',client=fresh,resilient=True)
        self.assertEqual(len(fresh.classify.call_args.args[0]),1)
        self.assertEqual(json.loads((folder/'failures.json').read_text(encoding='utf-8')),{})

    def test_timeout_splits_batch_instead_of_resending_it(self):
        # Reproduces the reported production failure: the model cannot answer a large
        # batch inside the socket timeout, but answers one or two papers well inside it.
        def classify(papers,model):
            if len(papers)>2:
                raise ModelRequestError('synthetic timeout',retryable=True,retry_after=10)
            return {p['id']:answer(p) for p in papers}
        client=Mock(classify=Mock(side_effect=classify))
        stop=Mock(is_set=Mock(return_value=False),wait=Mock(return_value=False))
        folder=run(self.path,self.root/'split',client=client,resilient=True,stop=stop,batch_size=5)
        status=json.loads((folder/'status.json').read_text(encoding='utf-8'))
        self.assertEqual((status['state'],status['completed'],status['failed']),('completed',4,0))
        sizes=[len(call.args[0]) for call in client.classify.call_args_list]
        # The oversized batch is sent once and then split, never resent at full size.
        self.assertEqual(sizes,[4,2,2])

    def test_exhausted_single_retry_is_recorded_and_run_continues(self):
        bad=self.papers[0]['id']
        def classify(papers,model):
            if papers[0]['id']==bad:
                raise ModelRequestError('synthetic transport',retryable=True,retry_after=10)
            return {p['id']:answer(p) for p in papers}
        client=Mock(classify=Mock(side_effect=classify))
        stop=Mock(is_set=Mock(return_value=False),wait=Mock(return_value=False))
        folder=run(self.path,self.root/'single',client=client,resilient=True,stop=stop,batch_size=1)
        status=json.loads((folder/'status.json').read_text(encoding='utf-8'))
        self.assertEqual((status['state'],status['completed'],status['failed']),('partial',3,1))
        self.assertIn(bad,json.loads((folder/'failures.json').read_text(encoding='utf-8')))
        # A transport failure is retried twice, then recorded, without stopping the run.
        self.assertEqual(client.classify.call_count,6)
        exported=json.loads((folder/'分类结果.json').read_text(encoding='utf-8'))
        failed=[r for r in exported['records'] if r['id']==bad][0]
        self.assertEqual(failed['status'],'分类失败')
        self.assertIn('synthetic transport',failed['failure']['error'])

    def test_consecutive_failures_stop_the_run_early(self):
        client=Mock(classify=Mock(side_effect=ModelRequestError('synthetic',retryable=True,retry_after=10)))
        stop=Mock(is_set=Mock(return_value=False),wait=Mock(return_value=False))
        with self.assertRaises(SafetyStop):
            run(self.path,self.root/'circuit',client=client,resilient=True,stop=stop,batch_size=1)
        self.assertLessEqual(client.classify.call_count,MAX_CONSECUTIVE_FAILURES)
        folder=next((self.root/'circuit').iterdir())
        status=json.loads((folder/'status.json').read_text(encoding='utf-8'))
        self.assertEqual(status['state'],'failed')

    def test_auto_retries_transient_only_and_limits_attempts(self):
        stop=Mock(is_set=Mock(return_value=False),wait=Mock(return_value=False))
        # A non-retryable failure still stops immediately: no split, no retry.
        client=Mock(classify=Mock(side_effect=ModelRequestError('synthetic',retryable=False,retry_after=20)))
        with self.assertRaises(ModelRequestError):
            run(self.path,self.root/'terminal',client=client,resilient=True,stop=stop)
        self.assertEqual(client.classify.call_count,1)
        folder=next((self.root/'terminal').iterdir())
        status=json.loads((folder/'status.json').read_text(encoding='utf-8'))
        self.assertEqual(status['state'],'failed')
        self.assertEqual(status['completed'],0)

    def test_cancellation_during_retry_is_saved_as_stopped(self):
        stopped=False
        stop=Mock()
        stop.is_set.side_effect=lambda:stopped
        def wait(delay):
            nonlocal stopped
            stopped=True
            return True
        stop.wait.side_effect=wait
        client=Mock(classify=Mock(side_effect=ModelRequestError('synthetic',retryable=True,retry_after=30)))
        folder=run(self.path,self.root/'out',client=client,resilient=True,stop=stop)
        self.assertEqual(client.classify.call_count,1)
        self.assertEqual(json.loads((folder/'status.json').read_text(encoding='utf-8'))['state'],'stopped')


class DesktopTests(unittest.TestCase):
    def test_saved_results_filter_details_and_model_switch(self):
        import tkinter as tk
        from classify_app import ClassifyApp
        root=tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        with patch('classify_app.read_papers',return_value={'sha256':'fixture','papers':[{'rows':[2,3]}]}):
            app=ClassifyApp(root)
        with tempfile.TemporaryDirectory() as tmp:
            app.output=Path(tmp)
            folder=app.output/'fixture'
            folder.mkdir()
            records=[{'id':'a','rows':[2],'title':'Alpha paper','doi':'10.1/a','status':'AI建议待复核','classification':{'type':'期刊论文','reason':'期刊资料','missing_evidence':['收录证据']},'import_route':{'recommended_channel':'WOS','reason':'核实导出文件','available_buttons':['WOS数据导入(Excel)']}},
                     {'id':'b','rows':[3],'title':'Beta paper','doi':'','status':'AI建议待复核','classification':{'type':None},'import_route':{'recommended_channel':None}}]
            (folder/'分类结果.json').write_text(json.dumps({'model':'deepseek-chat','source':{'sha256':'fixture'},'records':records}),encoding='utf-8')
            (folder/'导入渠道分类.md').write_text('fixture',encoding='utf-8')
            app.load_results()
            self.assertEqual(app.metrics['已处理'].get(),'2')
            self.assertEqual(app.metrics['渠道待判定'].get(),'1')
            self.assertEqual(float(app.bar['value']),100)
            app.search.set('10.1/A')
            self.assertEqual(app.tree.get_children(),('0',))
            app.tree.selection_set('0')
            app.show_detail()
            self.assertIn('核实导出文件',app.detail.get('1.0','end'))
            app.search.set('')
            app.channel.set('待判定')
            self.assertEqual(app.tree.get_children(),('1',))
            app.model.set('qwen')
            app.load_results()
            self.assertEqual(app.records,[])
            self.assertEqual(str(app.channel_button['state']),'disabled')
            self.assertEqual(float(app.bar['value']),0)

    def test_model_selects_batch_default_and_failures_are_visible(self):
        import tkinter as tk
        from classify_app import ClassifyApp
        root=tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        with patch('classify_app.read_papers',return_value={'sha256':'fixture','papers':[{'rows':[2,3]}]}):
            app=ClassifyApp(root)
        self.assertEqual(app.batch.get(),'5')
        app.model.set('deepseek-reasoner')
        app.change_model()
        # The reasoning model must not keep the chat model's batch size.
        self.assertEqual(app.batch.get(),'2')
        with tempfile.TemporaryDirectory() as tmp:
            app.output=Path(tmp)
            folder=app.output/'fixture'
            folder.mkdir()
            records=[{'id':'a','rows':[2],'title':'Failed paper','doi':'10.1/a','status':'分类失败',
                      'classification':None,'import_route':None,
                      'failure':{'rows':[2],'error':'模型连接失败、超时或响应损坏。'}}]
            (folder/'分类结果.json').write_text(json.dumps({'model':'deepseek-reasoner','source':{'sha256':'fixture'},'records':records}),encoding='utf-8')
            app.load_results()
            self.assertEqual(app.metrics['分类失败'].get(),'1')
            self.assertEqual(app.metrics['已处理'].get(),'0')
            app.channel.set('分类失败')
            self.assertEqual(app.tree.get_children(),('0',))
            app.tree.selection_set('0')
            app.show_detail()
            self.assertIn('超时',app.detail.get('1.0','end'))

    def test_stop_button_is_available_during_a_batch_export(self):
        import tkinter as tk
        from classify_app import ClassifyApp
        root=tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        with patch('classify_app.read_papers',return_value={'papers':[{'rows':[2,3]}]}):
            app=ClassifyApp(root,on_export=lambda:None)
        self.assertEqual(str(app.stop_button['state']),'disabled')
        app.set_exporting(True)
        # The export runs through the app's shared busy state, so the classifier's own
        # stop button must light up or the user has no way out of a long batch.
        self.assertEqual(str(app.stop_button['state']),'normal')
        self.assertEqual(str(app.export_button['state']),'disabled')
        app.set_external_busy(True)
        self.assertEqual(str(app.stop_button['state']),'normal')
        app.set_external_busy(False)
        self.assertFalse(app.exporting)
        self.assertEqual(str(app.stop_button['state']),'disabled')

    def test_export_clears_the_stop_flag_before_starting(self):
        import tkinter as tk
        from classify_app import ClassifyApp
        root=tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        started=[]
        with patch('classify_app.read_papers',return_value={'papers':[{'rows':[2,3]}]}):
            app=ClassifyApp(root,on_export=lambda:started.append(app.stop.is_set()))
        app.stop.set()
        app.export_wos()
        self.assertEqual(started,[False],'a previous stop must not cancel the new batch')

    def test_window_controls_progress_and_stop(self):
        import tkinter as tk
        from classify_app import ClassifyApp
        root=tk.Tk()
        root.withdraw()
        self.addCleanup(root.destroy)
        with patch('classify_app.read_papers',return_value={'papers':[{'rows':[2,3]}]}):
            app=ClassifyApp(root)
        self.assertIn('2 条记录',app.summary.get())
        app.set_busy(True)
        self.assertEqual(str(app.start_button['state']),'disabled')
        app.request_stop()
        self.assertTrue(app.stop.is_set())
        app.queue.put(('progress','已保存 1/1'))
        app.queue.put(('done',Path('synthetic-result')))
        app.poll()
        self.assertFalse(app.busy)
        self.assertEqual(app.status.get(),'已保存 1/1')
        self.assertEqual(app.result_folder,Path('synthetic-result'))


if __name__=='__main__':
    unittest.main()

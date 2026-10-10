"""Native UI integration with synthetic files and a fake backend only."""
import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from app import App
from core import HEADERS, Journal, file_hash, read_roster
from operation_log import OperationLog
from settings_panel import DEFAULTS
from tests.test_automation import sample
from tests.test_roster_write import make_roster
from tests.test_wos_import import Bridge


class ImportUITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.path = self.folder / "list.xlsx"
        make_roster(self.path, flags=(None, None))
        book = load_workbook(self.path)
        for key, value in {"owner": "谭勋策", "title": "Synthetic paper", "doi": "10.1234/test",
                           "matches": 0, "item_ids": "", "reason": ""}.items():
            book.active.cell(2, 2 + list(HEADERS).index(key), value)
        book.save(self.path)
        book.close()
        self.inbox = self.folder / "inbox"
        self.inbox.mkdir()
        self.file = self.inbox / "one.txt"
        self.file.write_bytes(sample())
        self.root = tk.Tk()
        self.root.withdraw()
        with patch("settings_panel.read_preferences", return_value=DEFAULTS):
            self.app = App(self.root, Journal(self.folder / "progress.db"), auto_load=False,
                           operation_log=OperationLog(self.folder / "log.txt"))
        self.app.loaded(read_roster(self.path))
        self.app.owner.set("谭勋策")
        self.app.select_owner()
        self.app.automation_panel.runtime = self.folder / "imports"
        self.panel = self.app.wos_import_panel
        self.panel.inbox.set(str(self.inbox))
        self.bridge = Bridge(self.app.roster.records)
        self.bridge.online = True
        self.bridge.close = lambda: None
        self.app.bridge = self.bridge
        self.runner = patch.object(self.app, "run", side_effect=self.run_sync)
        self.runner.start()
        from automation import ImportStore
        self.downloads = ImportStore(self.folder / 'downloads')
        self.download_store_patch = patch('wos_batch.default_store', return_value=self.downloads)
        self.download_store_patch.start()

    def tearDown(self):
        self.runner.stop()
        self.download_store_patch.stop()
        self.app.set_busy(False)
        self.root.update_idletasks()
        self.app.close()
        self.tmp.cleanup()

    def run_sync(self, job, callback, status, **kwargs):
        self.app.set_busy(True)
        try:
            result = job()
        finally:
            self.app.set_busy(False)
        callback(result)

    def test_preview_import_and_restart_do_not_approve_excel(self):
        before = file_hash(self.path)
        self.panel.preview()
        self.assertEqual(self.bridge.calls, [])
        self.assertEqual(self.panel.plan.items[0].status, "ready")
        self.panel.reviewed.set(True)
        with patch("wos_import_panel.messagebox.askyesno", return_value=True):
            self.panel.start()
        self.assertEqual(self.panel.entries["demo-001"]["status"], "pushed")
        self.assertEqual(file_hash(self.path), before)
        self.assertIsNone(self.panel.plan)
        self.assertFalse(self.panel.running)
        self.assertFalse(self.app.busy)
        self.assertFalse(self.panel.reviewed.get())
        self.panel.preview()
        self.assertEqual(self.panel.plan.excluded[0].status, "pushed")
        self.assertEqual(sum(a == "import_submit" for a, _ in self.bridge.calls), 1)

    def test_no_consent_or_cancel_sends_no_commands(self):
        self.panel.preview()
        with patch("wos_import_panel.messagebox.showwarning") as warning:
            self.panel.start()
            self.assertIn("勾选确认", warning.call_args.args[1])
        self.panel.reviewed.set(True)
        with patch("wos_import_panel.messagebox.askyesno", return_value=False):
            self.panel.start()
        self.assertEqual(self.bridge.calls, [])

    def test_owner_and_limit_changes_invalidate_review(self):
        self.panel.preview()
        self.panel.reviewed.set(True)
        self.panel.limit.set("1")
        self.assertIsNone(self.panel.plan)
        self.assertFalse(self.panel.reviewed.get())
        self.panel.preview()
        self.app.owner.set("另一位")
        self.assertIsNone(self.panel.plan)
        with patch("wos_import_panel.messagebox.showwarning") as warning:
            self.panel.start()
            self.assertIn("不操作其他负责人", warning.call_args.args[1])
        self.assertEqual(self.bridge.calls, [])

    def test_single_local_file_read_needs_no_browser_and_does_not_submit(self):
        self.app.next_record()
        self.app.bridge = None
        with patch("automation_panel.filedialog.askopenfilename", return_value=str(self.file)):
            self.app.automation_panel.load_file()
        state = self.app.automation_panel.store.get(self.app.current)
        self.assertEqual(state["phase"], "exported")
        self.assertEqual(self.bridge.calls, [])

    def test_failed_job_resets_stop_button_and_busy_state(self):
        self.panel.running = True
        self.app.set_busy(True)
        self.assertEqual(str(self.panel.stop_button["state"]), "normal")
        self.app.set_busy(False)
        self.assertFalse(self.panel.running)
        self.assertEqual(str(self.panel.stop_button["state"]), "disabled")

    def test_download_handoff_preserves_skip_scope_without_browser_actions(self):
        from wos_import_panel import SCOPES
        with patch('wos_batch.default_inbox', return_value=self.inbox):
            self.panel.receive_downloads('谭勋策', 'skipped')
        self.assertEqual(self.app.owner.get(), '谭勋策')
        self.assertEqual(SCOPES[self.panel.scope.get()], 'skipped')
        self.assertEqual(self.panel.inbox.get(), str(self.inbox))
        self.assertEqual(self.bridge.calls, [])
        self.assertIsNone(self.panel.plan)

    def test_scope_change_invalidates_plan_and_consent(self):
        self.panel.preview()
        self.panel.reviewed.set(True)
        self.panel.scope.set('已跳过论文（备注为 2）')
        self.assertIsNone(self.panel.plan)
        self.assertFalse(self.panel.reviewed.get())
        self.assertEqual(self.panel.entries, {})

    def test_download_button_only_downloads_and_never_imports(self):
        self.app.next_record()
        original = self.bridge.call
        def call(action, payload, timeout=75):
            if action == 'status':
                return original('search', payload, timeout)
            if action == 'wos_search':
                self.bridge.calls.append((action, payload['sa_id']))
                return {}
            if action == 'wos_export':
                self.bridge.calls.append((action, payload['sa_id']))
                return {'sa_id': payload['sa_id'], 'path': str(self.file)}
            return original(action, payload, timeout)
        self.bridge.call = call
        self.app.automation_panel.start(current_wos=True)
        actions = [a for a, _ in self.bridge.calls]
        self.assertIn('wos_export', actions)
        self.assertNotIn('wos_search', actions)
        self.assertFalse(any(a.startswith('import_') for a in actions))
        self.assertEqual(self.panel.store().get(self.app.current)['phase'], 'exported')

    def test_strong_identity_still_requires_approval_to_import(self):
        self.app.next_record()
        from automation import WOSFlow
        WOSFlow(self.bridge, self.panel.store()).prepare_file(self.app.current, self.file)
        with patch('automation_panel.messagebox.askyesno', return_value=False) as confirm:
            self.app.automation_panel.resume()
        confirm.assert_called_once()
        self.assertEqual([action for action, _ in self.bridge.calls], ['status', 'search'])

    def test_single_import_after_approval_reaches_verified_push(self):
        self.app.next_record()
        from automation import WOSFlow
        record = self.app.current
        WOSFlow(self.bridge, self.panel.store()).prepare_file(record, self.file)
        with patch('automation_panel.messagebox.askyesno', return_value=True) as confirm:
            self.app.automation_panel.resume()
        confirm.assert_called_once()
        actions = [action for action, _ in self.bridge.calls]
        self.assertEqual(actions[:2], ['status', 'search'])
        for action in ('import_upload', 'import_submit', 'import_push'):
            self.assertEqual(actions.count(action), 1)
        self.assertEqual(self.panel.store().get(record)['phase'], 'pushed')
        self.assertFalse(read_roster(self.path).records[0].done)

    def test_single_review_without_selection_does_not_reuse_an_old_record(self):
        self.app.next_record()
        selected_tab = self.app.tabs.select()
        with patch('wos_import_panel.messagebox.showinfo') as notice:
            self.panel.open_single()
        notice.assert_called_once()
        self.assertEqual(self.app.tabs.select(), selected_tab)
        self.assertEqual(self.bridge.calls, [])

    def test_weak_archived_download_can_be_loaded_without_download_or_upload(self):
        from automation import parse_wos
        book = load_workbook(self.path)
        book.active.cell(2, 2 + list(HEADERS).index('doi')).value = None
        book.save(self.path)
        book.close()
        self.file.unlink()
        self.app.loaded(read_roster(self.path))
        record = self.app.roster.records[0]
        self.downloads.archive(sample())
        self.downloads.save(record, {'phase': 'downloaded', 'candidate': parse_wos(sample()), 'identity_confirmed': False})
        self.panel.preview()
        self.assertEqual(self.panel.plan.items, ())
        self.assertEqual(self.panel.plan.excluded[0].status, 'deferred')
        self.panel.tree.selection_set(record.sa_id)
        self.panel.open_single()
        state = self.panel.store().get(record)
        self.assertEqual(state['phase'], 'exported')
        self.assertFalse(state['identity_confirmed'])
        self.assertEqual(self.bridge.calls, [])

    def test_one_click_runs_owned_zero_match_to_web_and_excel_without_extra_approval(self):
        from tests.test_zero_match import Backend, Download
        self.app.bridge = Backend(self.app.roster.records)
        self.app.bridge.online = True
        self.app.bridge.close = lambda: None
        transport = Download(self.folder)
        self.panel.limit.set('1')
        with patch('wos_browser.select_transport', return_value=transport), \
             patch('zero_match.SerialWOS', side_effect=lambda transport, *args: transport), \
             patch('wos_import_panel.messagebox.askyesno') as confirm:
            self.panel.one_click()
        confirm.assert_not_called()
        self.assertTrue(self.app.roster.records[0].done)
        self.assertEqual(self.panel.entries['demo-001']['status'], 'done')
        self.assertFalse(self.app.busy)
        self.assertFalse(self.panel.running)
        self.assertIn('已结案 1', self.panel.status.get())

    def test_one_click_global_verification_pause_retains_list_and_has_no_tail_skips(self):
        from tests.test_zero_match import Backend, Download
        self.app.bridge = Backend(self.app.roster.records)
        self.app.bridge.online = True
        self.app.bridge.close = lambda: None
        transport = Download(self.folder)
        transport.fail['demo-001'] = 'WOS 显示人工验证'
        before = file_hash(self.path)
        with patch('wos_browser.select_transport', return_value=transport), \
             patch('zero_match.SerialWOS', side_effect=lambda transport, *args: transport):
            self.panel.one_click()
        self.assertEqual(file_hash(self.path), before)
        self.assertEqual(self.panel.entries['demo-001']['status'], 'halted')
        self.assertIn('已暂停', self.panel.status.get())
        self.assertFalse(self.app.busy)

"""Synthetic workbook only: narrow notes, new-only fill, atomic safe repeats."""
import copy
from dataclasses import replace
from pathlib import Path
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook
from core import SafetyStop,read_roster,file_hash
from roster_write import migrate_status_column
from roster_annotations import annotate_roster,highlight_new_admission,verified_new_admission,format_owned_status
from tests.test_roster_write import make_roster


class AnnotationTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.path=Path(self.temp.name)/'list.xlsx';make_roster(self.path,flags=(None,None))
        book=load_workbook(self.path);book.active['B2']='谭勋策';book.active['H2']=0;book.active['I2']='';book.save(self.path);book.close()
        self.roster=migrate_status_column(read_roster(self.path));self.record=self.roster.records[0]

    def complete(self):
        book=load_workbook(self.path);s=book.active;s['A2']='已入库';s.cell(2,self.roster.completion_column,1)
        s['C2'].number_format='@';s['C2'].alignment=copy.copy(s['A2'].alignment);book.save(self.path);book.close()
        self.roster=read_roster(self.path);self.record=self.roster.records[0]
        return {'phase':'done','item_id':'1234567890123456789','batch':{'id':'batch-test','modelId':'model-test',
          'source':'人工提交','instructions':'SA补充-'+self.record.sa_id,'total':1,'actual':1,'fail':0,'status':2,
          'increase':1,'duplicateSkip':0,'duplicateIncreaseUpdate':0,'duplicateOverallCoverage':0}}

    def test_review_notes_preserve_flags_other_rows_and_existing_remark(self):
        r=annotate_roster(self.roster,notes={self.record.sa_id:'姓名冲突'})
        self.assertEqual(r.roster.records[0].remark,'姓名冲突');self.assertFalse(r.roster.records[0].done)
        self.assertFalse(r.roster.records[0].skipped);self.assertEqual(r.roster.records[1],self.roster.records[1])
        a=annotate_roster(r.roster,notes={self.record.sa_id:'作者单位冲突'})
        self.assertEqual(a.roster.records[0].remark,'姓名冲突；作者单位冲突')
        b=annotate_roster(a.roster,notes={self.record.sa_id:'作者单位冲突'})
        self.assertEqual(b.roster.sha256,a.roster.sha256);self.assertIsNone(b.backup)

    def test_id_yellow_retains_data_font_border_and_package_members(self):
        proof=self.complete();before_bytes=self.path.read_bytes()
        before=load_workbook(self.path);cell=before.active['C2'];original_font=copy.copy(cell.font);original_border=copy.copy(cell.border)
        values=[[c.value for c in row] for row in before.active];before.close()
        result=highlight_new_admission(self.roster,self.record,proof)
        self.assertEqual(result.yellow_cells,('C2',));self.assertEqual(result.note_cells,())
        after=load_workbook(self.path);cell=after.active['C2']
        self.assertEqual(cell.fill.fgColor.rgb,'FFFFFF00');self.assertEqual(cell.font,original_font);self.assertEqual(cell.border,original_border)
        self.assertEqual(cell.number_format,'@');self.assertEqual([[c.value for c in row] for row in after.active],values);after.close()
        with zipfile.ZipFile(result.backup) as a,zipfile.ZipFile(self.path) as b:
            self.assertEqual(a.namelist(),b.namelist())
            self.assertEqual({n for n in a.namelist() if a.read(n)!=b.read(n)},{'xl/worksheets/sheet1.xml','xl/styles.xml'})
        self.assertEqual(result.backup.read_bytes(),before_bytes)
        again=highlight_new_admission(result.roster,result.roster.records[0],proof)
        self.assertEqual(again.roster.sha256,result.roster.sha256);self.assertEqual(again.yellow_cells,())

    def test_only_counted_new_completed_owner_records_can_be_yellow(self):
        proof=self.complete()
        for key,value in [('increase',0),('duplicateSkip',1),('fail',1),('total',True),('actual',0),('status',1)]:
            bad=copy.deepcopy(proof);bad['batch'][key]=value
            self.assertFalse(verified_new_admission(self.record,bad))
            with self.assertRaises(SafetyStop):annotate_roster(self.roster,new_admissions={self.record.sa_id:bad})
        for change in [{'phase':'push_intent'},{'item_id':1234},{'batch':None}]:
            self.assertFalse(verified_new_admission(self.record,{**proof,**change}))
        for record in [replace(self.record,owner='其他人'),replace(self.record,done=False),replace(self.record,remark='已认领')]:
            self.assertFalse(verified_new_admission(record,proof))

    def test_note_targets_cannot_include_completed_other_owners_or_unknown_ids(self):
        for notes in [{self.roster.records[1].sa_id:'待核'}, {'unknown':'待核'}]:
            with self.assertRaises(SafetyStop):annotate_roster(self.roster,notes=notes)
        self.complete()
        with self.assertRaises(SafetyStop):annotate_roster(self.roster,notes={self.record.sa_id:'姓名冲突'})

    def test_changed_original_never_overwrites(self):
        stale=self.roster;self.path.write_bytes(self.path.read_bytes()+b'changed')
        checksum=file_hash(self.path)
        with self.assertRaises(SafetyStop):annotate_roster(stale,notes={self.record.sa_id:'待核'})
        self.assertEqual(file_hash(self.path),checksum)

    def test_unresolved_red_requires_skip_and_preserves_other_rows(self):
        from roster_write import mark_skipped_many
        with self.assertRaises(SafetyStop):
            annotate_roster(self.roster,unresolved={self.record.sa_id:'作者资料冲突'})
        saved=mark_skipped_many(self.roster,[self.record],reasons={self.record.sa_id:'旧原因'}).roster
        result=annotate_roster(saved,unresolved={self.record.sa_id:'作者资料冲突'})
        self.assertEqual(result.red_cells,('C2',));self.assertEqual(result.yellow_cells,())
        self.assertEqual(result.roster.records[0].remark,'作者资料冲突')
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[1:],saved.records[1:])
        book=load_workbook(self.path);self.assertEqual(book.active['C2'].fill.fgColor.rgb,'FFFF0000');book.close()
        again=annotate_roster(result.roster,unresolved={self.record.sa_id:'作者资料冲突'})
        self.assertEqual(again.roster.sha256,result.roster.sha256)
        self.assertIsNone(again.backup)
        self.complete()
        with self.assertRaises(SafetyStop):annotate_roster(self.roster,unresolved={self.record.sa_id:'已完成不可标红'})
        cleared=annotate_roster(self.roster,resolved_ids=[self.record.sa_id])
        self.assertEqual(cleared.restored_cells,('C2',))
        book=load_workbook(self.path);self.assertIsNone(book.active['C2'].fill.patternType);book.close()

    def test_clear_red_does_not_erase_success_yellow_or_clear_pending(self):
        with self.assertRaises(SafetyStop):annotate_roster(self.roster,resolved_ids=[self.record.sa_id])
        proof=self.complete();yellow=highlight_new_admission(self.roster,self.record,proof)
        result=annotate_roster(yellow.roster,resolved_ids=[self.record.sa_id])
        self.assertEqual(result.roster.sha256,yellow.roster.sha256)
        self.assertEqual(result.restored_cells,())

    def test_unified_status_color_does_not_claim_new_admission_or_change_data(self):
        from roster_write import mark_complete,mark_skipped_many
        completed=mark_complete(self.roster,self.record,note='已认领').roster
        result=format_owned_status(completed)
        self.assertEqual(result.roster.records,completed.records)
        book=load_workbook(self.path);self.assertEqual(book.active['C2'].fill.fgColor.rgb,'FFFFFF00');book.close()
        self.assertEqual(result.yellow_cells,('C2',))
        self.assertIsNone(format_owned_status(result.roster).backup)
        self.assertFalse(verified_new_admission(result.roster.records[0],{}))


if __name__=='__main__':unittest.main()

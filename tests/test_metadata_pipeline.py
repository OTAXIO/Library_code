"""Durability and owner/zero-match guard tests; synthetic data only."""
import copy
from dataclasses import replace
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
from core import SafetyStop, file_hash, read_roster
from metadata_sources import MetadataClient, MetadataSkip, merge_metadata
from metadata_pipeline import MetadataReview, claim_state, resolved, run_metadata_batch, verify_metadata_comparison
from roster_write import migrate_status_column
from tests.test_metadata_sources import example
from tests.test_roster_write import make_roster
from zero_match import WorkflowStore


class Backend:
    def __init__(self, record, candidate):
        self.candidate=candidate;self.row={'saLzkId':record.sa_id,'titleValue':record.title,'doiValue':record.doi,
          'wosValue':'','gh':record.staff_id,'markStatus':'待处理','matchCount':0,'itemId':'','reason':'','remark':''}
        self.phase=0;self.calls=[];self.lost='';self.validation_error=False;self.claimed=True;self.roles_match=True
        self.auto_processed=False
        self.auto_processed_claim=False
        self.item_id='1234567890123456789'
        self.model_id='journal'
        self.pushed_counts=None

    def call(self, action, p, timeout=90):
        self.calls.append(action);c=self.candidate
        if action in ('status','search'):
            if action=='search' and not self.row['itemId']:
                raise SafetyStop('零匹配详情为空；关联前只能读取列表状态')
            comparison=[] if not self.row['itemId'] else [{'label':label,'sa':self.row[k],'library':c[f]}
              for label,k,f in [('题名','titleValue','title'),('DOI','doiValue','doi'),('WOS记录号','wosValue','wos')]]
            if comparison:
                comparison.extend([
                    {'label':'认领状态','sa':'Alice Test(00001)','library':'已认领' if self.claimed else '未认领'},
                    {'label':'作者信息','sa':'是否第一作者：否\n是否通讯作者：否',
                     'library':'是否第一作者：否\n是否通讯作者：'+('否' if self.roles_match else '是')},
                    {'label':'交大是否第一单位','sa':'是','library':'是'}])
            return {'metadata_completion_protocol':4,'metadata_detail_protocol':1,
                    'detail_released':action=='search' and p.get('close_detail') is True,
                    'row':copy.deepcopy(self.row),'comparison':comparison}
        if action=='import_excel_resolve':
            return {'verified':True,'library_resolution_protocol':2,'items':[] if self.phase!=2 else [{'id':self.item_id,
              'title':c['title'],'doi':c['doi'],'wos':''}]}
        if action=='import_excel_probe':return {'metadata_excel_protocol':1,'source':'人工提交','model_id':'journal',
          'model_name':c.get('model_name','期刊论文'),
          'fields':[{'label':label,'field':str(index),'required':0} for index,label in enumerate(c['fields'])]}
        if action=='import_excel_scan':return {'batches':[]}
        if action in ('import_excel_upload','import_excel_upload_status'):
            if self.validation_error:raise SafetyStop('格式校验待核验')
            result={'uploaded':True,'validated':True,'sha256':p['xlsx_sha256'],'dataset_id':'sjtu','model_id':'journal'}
        elif action=='import_excel_submit':self.phase=1;result={'submitted':True}
        elif action=='import_excel_check':
            if not self.phase:raise SafetyStop('无批次；未重试提交')
            batch={'id':'batch-one','modelId':self.model_id,'source':'人工提交',
              'instructions':p['instructions'],'total':1,'actual':1 if self.phase==2 else 0,'fail':0,'status':self.phase,
              'increase':1 if self.phase==2 else 0,'duplicateSkip':0,'duplicateIncreaseUpdate':0,'duplicateOverallCoverage':0}
            if self.phase==2 and self.pushed_counts is not None:batch.update(self.pushed_counts)
            return {'verified':True,'batch':batch}
        elif action=='import_excel_push':self.phase=2;result={'submitted':True}
        elif action=='link':
            self.row.update(itemId=p['item_id'],matchCount=1)
            if self.auto_processed:self.row['markStatus']='已处理'
            result={'verified':True,'row':copy.deepcopy(self.row)}
        elif action=='prepare_claim':
            return {**self.call('search',p), 'suggested_index':0, 'prepared':{
                'item_id':self.row['itemId'],'staff_id':'00001','sa_text':'Alice Test(00001)',
                'person':{'id':'person-one','wno':'00001','name':'Alice Test','names':['Alice Test']},
                'authors':[{'index':0,'id':'author-one','order':1,'fullname':'Alice Test',
                            'eligible':True,'scholarId':'','relations':[]}]}}
        elif action=='submit_claim':
            self.claimed=True
            if self.auto_processed_claim:self.row['markStatus']='已处理'
            result={'row':copy.deepcopy(self.row),'verified':True,'claimed':True,'staff_id':'00001',
                    'scholar_id':'person-one','author':'Alice Test','order':1}
        elif action=='complete':self.row.update(markStatus='已处理',remark=p['note']);result={'verified':True,'row':copy.deepcopy(self.row)}
        else:raise AssertionError(action)
        if self.lost==action:
            self.lost='';raise SafetyStop('模拟回执丢失')
        return result


class ClaimStateTests(unittest.TestCase):
    def setUp(self):
        self.record, _, _ = example()

    @staticmethod
    def comparison(first='否', corresponding='否', claimed=True):
        flags = f'是否第一作者：{first}\n是否通讯作者：{corresponding}'
        return [
            {'label':'认领状态','sa':'Alice Test(00001)',
             'library':'已认领' if claimed else '未认领'},
            {'label':'作者信息','sa':flags,'library':flags},
            {'label':'交大是否第一单位','sa':'是','library':'是'},
        ]

    def test_all_explicit_role_combinations_and_claim_states(self):
        for first in ('是', '否'):
            for corresponding in ('是', '否'):
                for claimed in (True, False):
                    with self.subTest(first=first, corresponding=corresponding, claimed=claimed):
                        comparison = self.comparison(first, corresponding, claimed)
                        before = copy.deepcopy(comparison)
                        self.assertEqual(claim_state(comparison, self.record), claimed)
                        self.assertEqual(comparison, before, 'Role verification must be read-only')

    def test_only_exact_stale_role_reasons_are_resolved(self):
        reasons = ('第一作者不一致', '是否第一作者不一致', '通讯作者不一致',
                   '是否通讯作者不一致', '通信作者不一致',
                   '作者不一致； 第一作者不一致;是否通讯作者不一致 ')
        for reason in reasons:
            with self.subTest(reason=reason):
                self.assertTrue(claim_state(self.comparison('是', '是'), self.record, reason))

    def test_unclaimed_empty_native_card_can_claim_without_inventing_role_flags(self):
        comparison=self.comparison('否','是',False);comparison[1]['library']=''
        self.assertFalse(claim_state(comparison,self.record,'作者不一致'))
        self.assertEqual(comparison[1]['library'],'')
        comparison[0]['library']='已认领'
        with self.assertRaises(SafetyStop):claim_state(comparison,self.record,'作者不一致')

    def test_cofirst_order_and_unit_reasons_still_require_review(self):
        reasons = ('共同一作不一致', '共同第一作者不一致', '首位作者待核验', '作者顺序不一致',
                   '交大是否第一单位不一致', '第一署名单位不一致',
                   '是否第一作者不一致；共同一作待核验', '通讯作者不一致（待确认）')
        for reason in reasons:
            with self.subTest(reason=reason):
                with self.assertRaisesRegex(MetadataReview, '作者角色待核验'):
                    claim_state(self.comparison('是', '是'), self.record, reason)

    def test_missing_duplicated_and_non_yes_no_role_flags_fail_closed(self):
        for side in ('sa', 'library'):
            for label in ('是否第一作者', '是否通讯作者'):
                for modification in ('missing', 'duplicate', 'conflicting', 'unknown', 'duplicate_unknown'):
                    with self.subTest(side=side, label=label, modification=modification):
                        comparison = self.comparison()
                        value = comparison[1][side]
                        if modification == 'missing':
                            value = value.replace(label + '：否', '')
                        elif modification == 'unknown':
                            value = value.replace(label + '：否', label + '：未知')
                        elif modification == 'duplicate_unknown':
                            value += '\n' + label + '：未知'
                        else:
                            value += '\n' + label + '：' + ('否' if modification == 'duplicate' else '是')
                        comparison[1][side] = value
                        with self.assertRaisesRegex(SafetyStop, '作者角色结构待核验'):
                            claim_state(comparison, self.record, '第一作者不一致;通讯作者不一致')

    def test_actual_role_or_unit_disagreement_is_not_overridden_by_stale_reason(self):
        for label, note in (('是否第一作者', '第一作者待核验'),
                            ('是否通讯作者', '通讯作者待核验')):
            with self.subTest(label=label):
                comparison = self.comparison()
                comparison[1]['library'] = comparison[1]['library'].replace(label + '：否', label + '：是')
                with self.assertRaisesRegex(MetadataReview, note):
                    claim_state(comparison, self.record, '第一作者不一致;通讯作者不一致')
        comparison = self.comparison()
        comparison[2]['library'] = '否'
        with self.assertRaisesRegex(MetadataReview, '第一单位待核验'):
            claim_state(comparison, self.record, '第一作者不一致;通讯作者不一致')

    def test_explicit_legacy_staff_numbers_must_uniquely_match_the_sa_person(self):
        for side in ('sa', 'library'):
            comparison = self.comparison()
            comparison[1][side] += '\n工号：00001'
            self.assertTrue(claim_state(comparison, self.record))
            for staff_text in ('工号：99999', '工号：1', '工号：未知', '工号：',
                               '工号：00001\n工号：00001', '工号：00001\n工号：其他'):
                with self.subTest(side=side, staff_text=staff_text):
                    comparison = self.comparison()
                    comparison[1][side] += '\n' + staff_text
                    with self.assertRaisesRegex(SafetyStop, '人员编号不唯一或不一致'):
                        claim_state(comparison, self.record, '第一作者不一致;通讯作者不一致')

    def test_invalid_reason_and_detail_shapes_raise_controlled_safety_stop(self):
        for reason in (0, False, [], {}, ['第一作者不一致']):
            with self.subTest(reason=reason):
                with self.assertRaisesRegex(SafetyStop, '原因结构待核验'):
                    claim_state(self.comparison(), self.record, reason)
        for comparison in (None, {}, [None], self.comparison() + [None]):
            with self.subTest(comparison=comparison):
                with self.assertRaises(SafetyStop):
                    claim_state(comparison, self.record)

    def test_duplicate_missing_or_invalid_fields_staff_and_claim_state_are_rejected(self):
        for index in range(3):
            for modification in ('duplicate', 'missing', 'nontext'):
                with self.subTest(index=index, modification=modification):
                    comparison = self.comparison()
                    if modification == 'duplicate':
                        comparison.append(copy.deepcopy(comparison[index]))
                    elif modification == 'missing':
                        comparison.pop(index)
                    else:
                        comparison[index]['library'] = None
                    with self.assertRaises(SafetyStop):
                        claim_state(comparison, self.record)
        for update in ({'sa':'Alice Test(99999)'}, {'library':'查询中'}):
            comparison = self.comparison()
            comparison[0].update(update)
            with self.assertRaises(SafetyStop):
                claim_state(comparison, self.record)


class MetadataComparisonTests(unittest.TestCase):
    def setUp(self):
        record, _, _ = example()
        self.record = replace(record, title='Low-pressure CO2 capture & release')
        self.candidate = {'title':self.record.title, 'doi':self.record.doi, 'wos':''}
        self.comparison = [
            {'label':'题名','sa':self.record.title,'library':self.record.title},
            {'label':'DOI','sa':self.record.doi,'library':self.record.doi},
            {'label':'WOS记录号','sa':'','library':''},
        ]

    def resolve_title(self, title, doi=None):
        return resolved({'verified':True,'library_resolution_protocol':2,'items':[
            {'id':'1234567890123456789','title':title,'doi':doi or self.record.doi}]},self.candidate)

    def test_discovery_resolution_and_comparison_share_display_only_title_rules(self):
        display_title = '<i>Low</i> – pressure CO<sub>2</sub> capture &amp; release.'
        self.assertEqual(self.resolve_title(display_title)['id'],'1234567890123456789')
        for side in ('sa','library'):
            with self.subTest(side=side):
                comparison=copy.deepcopy(self.comparison)
                comparison[0][side]=display_title
                before=copy.deepcopy(comparison)
                verify_metadata_comparison(comparison,self.candidate,self.record)
                self.assertEqual(comparison,before)

    def test_title_words_numbers_and_identifier_conflicts_remain_blocking(self):
        for title in ('Low-pressure CO20 capture & release', 'High-pressure CO2 capture & release',
                      'Low-pressure CO2 capture', 'Low pressure CO2 capture & release'):
            with self.subTest(title=title):
                with self.assertRaises(SafetyStop):self.resolve_title(title)
                for side in ('sa','library'):
                    comparison=copy.deepcopy(self.comparison)
                    comparison[0][side]=title
                    with self.assertRaises(SafetyStop):
                        verify_metadata_comparison(comparison,self.candidate,self.record)
        with self.assertRaises(SafetyStop):self.resolve_title(self.record.title,'10.1234/other')
        for index,value in ((1,'10.1234/other'),(2,'WOS:000000000000001')):
            for side in ('sa','library'):
                comparison=copy.deepcopy(self.comparison)
                comparison[index][side]=value
                with self.assertRaises(SafetyStop):
                    verify_metadata_comparison(comparison,self.candidate,self.record)

    def test_missing_duplicate_and_nontext_comparison_fields_are_not_normalized_away(self):
        for index in range(3):
            for modification in ('missing','duplicate','nontext'):
                with self.subTest(index=index,modification=modification):
                    comparison=copy.deepcopy(self.comparison)
                    if modification=='missing':comparison.pop(index)
                    elif modification=='duplicate':comparison.append(copy.deepcopy(comparison[index]))
                    else:comparison[index]['library']=None
                    with self.assertRaises(SafetyStop):
                        verify_metadata_comparison(comparison,self.candidate,self.record)

    def test_malformed_titles_are_structural_stops_never_metadata_skips(self):
        for malformed in (None, {}, 'x'*16001, 'bad\x01title'):
            with self.subTest(kind=type(malformed).__name__,size=len(malformed) if isinstance(malformed,str) else None):
                with self.assertRaises(SafetyStop) as raised:
                    self.resolve_title(malformed)
                self.assertNotIsInstance(raised.exception,MetadataSkip)
                for side in ('sa','library'):
                    comparison=copy.deepcopy(self.comparison)
                    comparison[0][side]=malformed
                    with self.assertRaises(SafetyStop) as raised:
                        verify_metadata_comparison(comparison,self.candidate,self.record)
                    self.assertNotIsInstance(raised.exception,MetadataSkip)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.base=Path(self.temp.name)
        path=self.base/'list.xlsx';make_roster(path,flags=(None,None))
        book=load_workbook(path);s=book.active
        for col,value in {'B2':'谭勋策','D2':'Synthetic verified paper','E2':'10.1234/example','H2':0,'I2':'','K2':''}.items():s[col]=value
        book.save(path);book.close()
        self.roster=migrate_status_column(read_roster(path));self.record=self.roster.records[0]
        _,cr,oa=example();self.client=MetadataClient(self.base/'runtime/metadata-imports/papers')
        self.client.get=Mock(side_effect=[{'message':cr},oa])
        self.candidate=merge_metadata(self.record,cr,oa);self.backend=Backend(self.record,self.candidate)
        self.stop=threading.Event()

    @staticmethod
    def builder(folder,candidate):
        path=folder/'paper.xlsx';path.write_bytes(b'PK synthetic test workbook')
        return path,file_hash(path)

    def run_once(self):
        return run_metadata_batch(self.base,self.roster,[self.record],self.backend,self.stop,
                                  client=self.client,excel_builder=self.builder)

    def test_success_exact_note_source_and_flags_other_owner_unchanged(self):
        other=self.roster.records[1]
        result=self.run_once();self.assertFalse(result.halted,result.reason)
        current=result.roster.records[0]
        self.assertTrue(current.done);self.assertEqual(current.remark,'已入库');self.assertEqual(current.source,'Crossref+OpenAlex')
        self.assertEqual(result.roster.records[1],other)
        self.assertEqual(self.backend.calls.count('import_excel_upload'),1)
        self.assertEqual(self.backend.calls.count('complete'),1)
        self.assertFalse(any(a.startswith('wos_') or 'claim' in a for a in self.backend.calls))

    def test_explicit_unchanged_claim_rejection_preserves_import_without_unknown_claim_intent(self):
        self.backend.claimed=False;original=self.backend.call
        def call(action,payload,timeout=90):
            if action=='submit_claim':
                self.backend.calls.append(action)
                return {'row':copy.deepcopy(self.backend.row),'rejected':True,'verified':False,
                    'claim_response_protocol':1,'unchanged':True,'staff_id':'00001',
                    'scholar_id':'person-one','order':1,'rejection':{'status':400,'code':400,
                    'message':'Synthetic validation rejection'}}
            return original(action,payload,timeout)
        self.backend.call=call;result=self.run_once()
        self.assertFalse(result.halted,result.reason);self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark,'后台拒绝认领')
        state=WorkflowStore(self.base/'runtime/metadata-imports').get(self.record)
        self.assertEqual(state['phase'],'linked');self.assertEqual(len(state['claim_rejections']),1)
        self.assertNotIn('complete',self.backend.calls)

    def test_malformed_library_title_preserves_prepared_task_without_marking_skip(self):
        for index,malformed in enumerate((None,{},'x'*16001,'bad\x01title')):
            with self.subTest(kind=type(malformed).__name__,size=len(malformed) if isinstance(malformed,str) else None):
                if index:self.setUp()
                original=self.backend.call
                def call(action,payload,timeout=90):
                    if action=='import_excel_resolve':
                        return {'verified':True,'library_resolution_protocol':2,'items':[
                            {'id':self.backend.item_id,'title':malformed,'doi':self.record.doi}]}
                    return original(action,payload,timeout)
                self.backend.call=call
                result=self.run_once()
                self.assertTrue(result.halted)
                self.assertIn('题名结构待核验',result.reason)
                self.assertFalse(result.roster.records[0].done)
                self.assertFalse(result.roster.records[0].skipped)
                self.assertEqual(result.roster.records[0].remark,self.record.remark)
                self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(self.record)['phase'],'prepared')
                self.assertNotIn('import_excel_upload',self.backend.calls)
                self.assertNotIn('import_excel_submit',self.backend.calls)

    def test_lost_ack_all_stages_resume_without_duplicate_writes(self):
        for action in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
            with self.subTest(action=action):
                # Independent synthetic folder and transport per uncertainty.
                self.setUp();self.backend.lost=action
                first=self.run_once();self.assertTrue(first.halted)
                self.roster=first.roster;self.record=self.roster.records[0]
                second=self.run_once();self.assertFalse(second.halted,second.reason)
                self.assertTrue(second.roster.records[0].done)
                for write in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
                    self.assertEqual(self.backend.calls.count(write),1,(action,write))

    def test_nonzero_backend_only_flags_skip_and_does_not_add_note(self):
        self.backend.row.update(matchCount=1,itemId='111',remark='')
        result=self.run_once();self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark,'');self.assertEqual(self.backend.calls,['status'])

    def test_skipped_is_not_repeated_in_normal_batch(self):
        self.record=replace(self.record,skipped=True)
        self.roster.records[0]=self.record
        result=self.run_once();self.assertEqual(self.backend.calls,[])
        self.assertEqual(result.remaining,0)

    def test_source_verified_role_repair_resumes_linked_skip_without_any_reimport(self):
        original=self.backend.call; fixed=False
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='search' and not fixed:
                info=next(v for v in answer['comparison'] if v['label']=='作者信息')
                info['sa']='是否第一作者：否\n是否通讯作者：是'
            elif action=='search' and fixed:
                info=next(v for v in answer['comparison'] if v['label']=='作者信息')
                info['sa']=info['library']='是否第一作者：否\n是否通讯作者：是'
            return answer
        self.backend.call=call
        with patch('publisher_roles.repair_correspondence',return_value=False):
            first=self.run_once()
        self.assertFalse(first.halted,first.reason);self.assertTrue(first.roster.records[0].skipped)
        state=WorkflowStore(self.base/'runtime/metadata-imports').get(first.roster.records[0])
        self.assertEqual(state['phase'],'linked')
        def repair(*args):
            nonlocal fixed
            fixed=True;return True
        with patch('publisher_roles.repair_correspondence',side_effect=repair):
            second=run_metadata_batch(self.base,first.roster,[first.roster.records[0]],self.backend,self.stop,
                retry_skipped=True,client=self.client,excel_builder=self.builder)
        self.assertFalse(second.halted,second.reason);self.assertTrue(second.roster.records[0].done)
        self.assertEqual(second.roster.records[0].remark,'已入库')
        for action in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
            self.assertEqual(self.backend.calls.count(action),1,action)

    def test_other_owner_or_nonzero_roster_rejected(self):
        for bad in (replace(self.record,owner='其他人'),replace(self.record,matches=1)):
            self.roster.records[0]=bad
            with self.assertRaises(SafetyStop):run_metadata_batch(self.base,self.roster,[bad],self.backend,self.stop)
        self.assertEqual(self.backend.calls,[])

    def test_old_protocol_stops_before_source_queries_or_any_write(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            result=original(action,payload,timeout)
            if action=='status':result['metadata_completion_protocol']=2
            return result
        self.backend.call=call
        result=self.run_once();self.assertTrue(result.halted);self.assertIn('0.4.31',result.reason)
        self.assertEqual(self.backend.calls,['status']);self.client.get.assert_not_called()
        self.assertEqual(result.roster.sha256,self.roster.sha256)

    def test_old_detail_protocol_stops_before_upload(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            answer.pop('metadata_detail_protocol',None)
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertTrue(result.halted)
        self.assertEqual(self.backend.calls,['status'])
        self.client.get.assert_not_called()

    def test_boolean_protocol_is_not_a_version_number(self):
        original = self.backend.call
        def call(action, payload, timeout=90):
            answer = original(action, payload, timeout)
            if action == 'status': answer['metadata_detail_protocol'] = True
            return answer
        self.backend.call = call
        result = self.run_once()
        self.assertTrue(result.halted)
        self.assertEqual(self.backend.calls, ['status'])
        self.client.get.assert_not_called()

    def test_boolean_excel_protocol_cannot_start_upload(self):
        original = self.backend.call
        def call(action, payload, timeout=90):
            answer = original(action, payload, timeout)
            if action == 'import_excel_probe': answer['metadata_excel_protocol'] = True
            return answer
        self.backend.call = call
        result = self.run_once()
        self.assertTrue(result.halted)
        self.assertNotIn('import_excel_upload', self.backend.calls)

    def test_unconfirmed_detail_release_never_claims_or_completes(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='search':answer['detail_released']=False
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertTrue(result.halted)
        self.assertFalse(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertNotIn('prepare_claim',self.backend.calls)
        self.assertNotIn('complete',self.backend.calls)

    def test_pause_after_upload_saves_receipt_and_never_submits(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            result=original(action,payload,timeout)
            if action=='import_excel_upload':self.stop.set()
            return result
        self.backend.call=call
        result=self.run_once();self.assertTrue(result.halted)
        self.assertNotIn('import_excel_submit',self.backend.calls)
        self.stop.clear();self.roster=result.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertFalse(second.halted,second.reason)
        self.assertEqual(self.backend.calls.count('import_excel_upload'),1)

    def test_native_auto_processed_after_link_still_gets_intake_note(self):
        self.backend.auto_processed=True
        result=self.run_once()
        self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(result.roster.records[0].remark,'已入库')
        self.assertEqual(self.backend.row['markStatus'],'已处理')
        self.assertEqual(self.backend.row['remark'],'已入库')
        self.assertEqual(self.backend.calls.count('import_excel_submit'),1)
        self.assertEqual(self.backend.calls.count('complete'),1)

    def test_native_auto_processed_lost_note_ack_is_read_only_reconciled(self):
        from zero_match import WorkflowStore
        self.backend.auto_processed=True;self.backend.lost='complete'
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertFalse(second.halted,second.reason)
        self.assertTrue(second.roster.records[0].done)
        state=WorkflowStore(self.base/'runtime/metadata-imports').get(self.record)
        self.assertEqual(state['phase'],'done')
        for action in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
            self.assertEqual(self.backend.calls.count(action),1,action)

    def test_owned_auto_processed_manual_note_is_not_generic_sync_or_overwritten(self):
        self.backend.auto_processed=True
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='link':self.backend.row['remark']='人工保留备注'
            return answer
        self.backend.call=call
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done)
        self.assertEqual(self.backend.row['remark'],'人工保留备注')
        self.assertNotIn('complete',self.backend.calls)
        self.assertEqual(self.backend.calls.count('import_excel_submit'),1)

    def test_auto_processed_unknown_claim_ack_cannot_bypass_claim_intent(self):
        self.backend.claimed=False;self.backend.auto_processed_claim=True;self.backend.lost='submit_claim'
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done)
        self.assertNotIn('complete',self.backend.calls)
        self.assertEqual(self.backend.calls.count('submit_claim'),1)

    def test_auto_processed_changed_owned_link_is_not_generic_completion_sync(self):
        self.backend.auto_processed=True
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='link':self.backend.row['itemId']='9999999999999999999'
            return answer
        self.backend.call=call
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done)
        self.assertIn('原入库关联断点',second.reason)
        self.assertNotIn('complete',self.backend.calls)
        self.assertEqual(self.backend.calls.count('import_excel_submit'),1)

    def test_newly_linked_unclaimed_author_is_claimed_before_closure(self):
        self.backend.claimed=False
        result=self.run_once();self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].done)
        calls=self.backend.calls
        self.assertEqual(calls.count('submit_claim'),1)
        self.assertLess(calls.index('link'),calls.index('submit_claim'))
        self.assertLess(calls.index('submit_claim'),calls.index('complete'))

    def test_unknown_claim_ack_never_reclaims_or_closes(self):
        self.backend.claimed=False;self.backend.lost='submit_claim'
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done)
        self.assertEqual(self.backend.calls.count('submit_claim'),1)
        self.assertNotIn('complete',self.backend.calls)
        self.assertEqual(self.backend.calls.count('import_excel_submit'),1)

    def test_role_disagreement_never_claims_or_closes(self):
        self.backend.claimed=False;self.backend.roles_match=False
        result=self.run_once();self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark,'通讯作者待核验')
        self.assertNotIn('prepare_claim',self.backend.calls)
        self.assertNotIn('complete',self.backend.calls)
        self.assertFalse(result.roster.records[0].done)

    def test_changed_batch_model_is_not_accepted_or_pushed(self):
        self.backend.model_id='unexpected-model'
        result=self.run_once()
        self.assertTrue(result.halted)
        self.assertNotIn('import_excel_push',self.backend.calls)
        self.assertNotIn('link',self.backend.calls)
        self.assertFalse(result.roster.records[0].done)

    def test_explicit_single_duplicate_is_shelved_never_counted_new_or_closed(self):
        from zero_match import WorkflowStore
        for index,key in enumerate(('duplicateSkip','duplicateIncreaseUpdate','duplicateOverallCoverage')):
            with self.subTest(key=key):
                if index:self.setUp()
                self.backend.pushed_counts={'increase':0,key:1}
                result=self.run_once()
                self.assertFalse(result.halted,result.reason)
                self.assertTrue(result.roster.records[0].skipped)
                self.assertFalse(result.roster.records[0].done)
                self.assertEqual(result.roster.records[0].remark,'本库重复待核验')
                state=WorkflowStore(self.base/'runtime/metadata-imports').get(result.roster.records[0])
                self.assertEqual(state['phase'],'duplicate_review')
                self.assertEqual(state['batch'][key],1)
                self.assertNotIn('link',self.backend.calls)
                self.assertNotIn('complete',self.backend.calls)
                for action in ('import_excel_upload','import_excel_submit','import_excel_push'):
                    self.assertEqual(self.backend.calls.count(action),1)

    def test_unknown_or_conflicting_push_counters_never_become_skips(self):
        for index,counts in enumerate(({'increase':0},{'increase':0,'duplicateIncreaseUpdate':True},
                {'increase':0,'duplicateSkip':1,'duplicateIncreaseUpdate':1},
                {'increase':1,'duplicateSkip':1},{'increase':0,'duplicateSkip':''})):
            with self.subTest(counts=counts):
                if index:self.setUp()
                self.backend.pushed_counts=counts
                result=self.run_once()
                self.assertTrue(result.halted)
                self.assertFalse(result.roster.records[0].done)
                self.assertFalse(result.roster.records[0].skipped)
                self.assertNotIn('link',self.backend.calls)
                self.assertNotIn('complete',self.backend.calls)

    def test_lost_push_ack_can_only_read_verified_duplicate_without_replaying(self):
        self.backend.pushed_counts={'increase':0,'duplicateIncreaseUpdate':1}
        self.backend.lost='import_excel_push'
        first=self.run_once()
        self.assertTrue(first.halted)
        self.assertFalse(first.roster.records[0].skipped)
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once()
        self.assertFalse(second.halted,second.reason)
        self.assertTrue(second.roster.records[0].skipped)
        self.assertEqual(self.backend.calls.count('import_excel_push'),1)
        self.assertEqual(self.backend.calls.count('import_excel_submit'),1)

    def test_duplicate_recognition_still_requires_exact_batch_model_and_counts(self):
        from metadata_pipeline import verified_pushed_duplicate
        self.backend.phase=2
        self.backend.pushed_counts={'increase':0,'duplicateIncreaseUpdate':1}
        batch=self.backend.call('import_excel_check',{'instructions':'SA补充-'+self.record.sa_id})['batch']
        self.assertTrue(verified_pushed_duplicate(batch,self.record,'batch-one',model_id='journal'))
        for change in ({'source':'WOS'},{'instructions':'SA补充-other'},{'id':'other'},
                {'modelId':'other'},{'status':1},{'fail':1},{'total':2},{'actual':0}):
            self.assertFalse(verified_pushed_duplicate({**batch,**change},self.record,'batch-one',model_id='journal'))
        self.assertFalse(verified_pushed_duplicate(batch,self.record,'batch-one',model_id=None))

    def test_retry_known_duplicate_never_uploads_pushes_or_closes_again(self):
        self.backend.pushed_counts={'increase':0,'duplicateIncreaseUpdate':1}
        first=self.run_once();self.assertFalse(first.halted,first.reason)
        previous=len(self.backend.calls)
        second=run_metadata_batch(self.base,first.roster,[first.roster.records[0]],self.backend,self.stop,
                                 retry_skipped=True,client=self.client,excel_builder=self.builder)
        self.assertFalse(second.halted,second.reason)
        self.assertTrue(second.roster.records[0].skipped)
        self.assertFalse(second.roster.records[0].done)
        self.assertFalse(set(self.backend.calls[previous:])&{'import_excel_upload','import_excel_submit',
            'import_excel_push','link','submit_claim','complete'})
        self.assertEqual(second.roster.sha256,first.roster.sha256)

    def test_previous_verified_new_batch_cannot_be_reclassified_as_a_duplicate(self):
        from zero_match import WorkflowStore
        original=self.backend.call
        def call(action,payload,timeout=90):
            if action=='import_excel_resolve' and self.backend.phase==2:
                raise SafetyStop('暂停本库索引回读')
            return original(action,payload,timeout)
        self.backend.call=call
        first=self.run_once();self.assertTrue(first.halted)
        saved=WorkflowStore(self.base/'runtime/metadata-imports').get(first.roster.records[0])
        self.assertEqual(saved['phase'],'pushed');self.assertEqual(saved['batch']['increase'],1)
        self.backend.call=original;self.backend.pushed_counts={'increase':0,'duplicateIncreaseUpdate':1}
        self.roster=first.roster;self.record=self.roster.records[0]
        second=self.run_once();self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done);self.assertFalse(second.roster.records[0].skipped)
        self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(self.record),saved)
        self.assertEqual(self.backend.calls.count('import_excel_push'),1)

    def test_completion_with_changed_match_count_does_not_complete_list(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='complete':answer['row']['matchCount']=2
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertTrue(result.halted)
        self.assertFalse(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)

    def test_required_model_field_is_checked_before_upload(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='import_excel_probe':answer['fields'].append({'label':'必填但来源缺失','field':'requiredExtra','required':True})
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark,'导入必填字段缺失')
        self.assertNotIn('import_excel_upload',self.backend.calls)

    def test_unknown_model_field_shape_halts_without_upload_or_false_skip(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='import_excel_probe':answer['fields']=[None]
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertTrue(result.halted)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertNotIn('import_excel_upload',self.backend.calls)

    def test_verified_claim_with_exact_stale_role_reason_can_close(self):
        self.backend.claimed=False
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='submit_claim':self.backend.row['reason']='通讯作者不一致'
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertFalse(result.halted,result.reason)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(self.backend.calls.count('submit_claim'),1)
        self.assertEqual(self.backend.calls.count('complete'),1)

    def test_verified_claim_with_new_actual_role_review_is_shelved_not_closed(self):
        self.backend.claimed=False
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='submit_claim':
                self.backend.row['reason']='通讯作者不一致'
                self.backend.roles_match=False
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertFalse(result.roster.records[0].done)
        self.assertEqual(result.roster.records[0].remark,'通讯作者待核验')
        self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(result.roster.records[0])['phase'],'claimed')
        self.assertEqual(self.backend.calls.count('submit_claim'),1)
        self.assertNotIn('complete',self.backend.calls)

    def test_verified_claim_with_cofirst_reason_still_cannot_close(self):
        self.backend.claimed=False
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='submit_claim':self.backend.row['reason']='共同一作待核验'
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark,'作者角色待核验')
        self.assertNotIn('complete',self.backend.calls)

    def test_unknown_import_and_external_processed_state_cannot_fake_completion(self):
        self.backend.lost='import_excel_submit'
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        self.backend.row.update(markStatus='已处理',remark='人工处理')
        second=self.run_once()
        self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done)
        self.assertFalse(second.roster.records[0].skipped)
        self.assertEqual(self.backend.calls.count('import_excel_submit'),1)
        self.assertNotIn('complete',self.backend.calls)

    def test_unknown_claim_with_changed_roles_is_not_converted_to_skip(self):
        self.backend.claimed=False;self.backend.lost='submit_claim'
        first=self.run_once();self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        self.backend.roles_match=False
        second=self.run_once()
        self.assertTrue(second.halted)
        self.assertFalse(second.roster.records[0].done)
        self.assertFalse(second.roster.records[0].skipped)
        self.assertEqual(self.backend.calls.count('submit_claim'),1)

    def test_unknown_claim_and_closure_intents_keep_role_problems_without_replay_or_skip(self):
        for index,action in enumerate(('submit_claim','complete')):
            with self.subTest(action=action):
                if index:self.setUp()
                self.backend.claimed=False
                self.backend.lost=action
                first=self.run_once()
                self.assertTrue(first.halted)
                self.roster=first.roster;self.record=self.roster.records[0]
                store=WorkflowStore(self.base/'runtime/metadata-imports')
                phase='claim_intent' if action=='submit_claim' else 'complete_intent'
                saved=store.get(self.record)
                self.assertEqual(saved['phase'],phase)
                self.backend.row['reason']='共同一作待核验'
                self.backend.roles_match=False
                previous=len(self.backend.calls)
                second=self.run_once()
                self.assertTrue(second.halted)
                self.assertFalse(second.roster.records[0].done)
                self.assertFalse(second.roster.records[0].skipped)
                self.assertEqual(second.roster.sha256,first.roster.sha256)
                self.assertEqual(store.get(self.record)['phase'],phase)
                self.assertNotIn('deferred_reason',store.get(self.record))
                self.assertFalse(set(self.backend.calls[previous:]) & {'import_excel_upload',
                    'import_excel_submit','import_excel_push','link','submit_claim','complete'})

    def test_malformed_post_link_reason_is_controlled_and_not_a_skip(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            answer=original(action,payload,timeout)
            if action=='search':answer['row']['reason']=['第一作者不一致']
            return answer
        self.backend.call=call
        result=self.run_once()
        self.assertTrue(result.halted)
        self.assertIn('原因结构待核验',result.reason)
        self.assertFalse(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(result.roster.records[0])['phase'],'linked')
        self.assertNotIn('submit_claim',self.backend.calls)
        self.assertNotIn('complete',self.backend.calls)

    def pause_after_preparing(self):
        original=self.backend.call
        def call(action,payload,timeout=90):
            if action=='import_excel_resolve':raise SafetyStop('测试只保存准备断点')
            return original(action,payload,timeout)
        self.backend.call=call
        first=self.run_once()
        self.assertTrue(first.halted)
        self.roster=first.roster;self.record=self.roster.records[0]
        self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(self.record)['phase'],'prepared')
        self.assertNotIn('import_excel_upload',self.backend.calls)
        self.backend.call=original
        return first

    def test_prepared_resume_rechecks_new_author_or_unit_reason_before_upload(self):
        for index,reason in enumerate(('共同一作待核验','作者顺序不一致','第一单位不一致','第一作者不一致')):
            with self.subTest(reason=reason):
                if index:self.setUp()
                self.pause_after_preparing()
                self.backend.row['reason']=reason
                resumed=self.run_once()
                self.assertFalse(resumed.halted,resumed.reason)
                self.assertTrue(resumed.roster.records[0].skipped)
                self.assertEqual(resumed.roster.records[0].remark,'作者角色待核验')
                self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(self.record)['phase'],'prepared')
                self.assertNotIn('import_excel_upload',self.backend.calls)
                self.assertNotIn('import_excel_submit',self.backend.calls)

    def test_nontext_reason_before_upload_halts_new_and_prepared_tasks_without_skip(self):
        for index,(prepared,reason) in enumerate((prepared,reason) for prepared in (False,True)
                for reason in (False,0,[],{},['第一作者不一致'])):
            with self.subTest(prepared=prepared,reason=reason):
                if index:self.setUp()
                if prepared:self.pause_after_preparing()
                before=self.roster.sha256
                self.backend.row['reason']=reason
                result=self.run_once()
                self.assertTrue(result.halted)
                self.assertIn('原因结构待核验',result.reason)
                self.assertEqual(result.roster.sha256,before)
                self.assertFalse(result.roster.records[0].skipped)
                self.assertNotIn('import_excel_upload',self.backend.calls)
                if not prepared:self.client.get.assert_not_called()

    def test_prepared_resume_without_new_role_issue_completes_once(self):
        self.pause_after_preparing()
        result=self.run_once()
        self.assertFalse(result.halted,result.reason)
        self.assertTrue(result.roster.records[0].done)
        for action in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
            self.assertEqual(self.backend.calls.count(action),1)

    def make_five_queue(self):
        make_roster(self.roster.path,flags=(None,)*6)
        book=load_workbook(self.roster.path);sheet=book.active
        for index in range(5):
            row=index+2
            for col,value in {'B':'谭勋策','D':f'Synthetic queue paper {index}',
                    'E':f'10.1234/queue-{index}','G':'00001','H':0,'I':'','K':''}.items():
                sheet[f'{col}{row}']=value
        book.save(self.roster.path);book.close()
        self.roster=migrate_status_column(read_roster(self.roster.path))
        replies=[];backends={}
        for index,record in enumerate(self.roster.records[:5]):
            _,cr,oa=example()
            cr.update(DOI=record.doi,title=[record.title])
            oa.update(doi='https://doi.org/'+record.doi,title=record.title)
            replies.extend([{'message':cr},oa])
            backend=Backend(record,merge_metadata(record,cr,oa))
            backend.item_id=str(1234567890123456789+index)
            backend.auto_processed=True
            backend.claimed=index%2==0
            backends[record.sa_id]=backend
        self.client.get=Mock(side_effect=replies)
        class Queue:
            calls=[]
            def call(queue,action,payload,timeout=90):
                queue.calls.append((payload['sa_id'],action))
                return backends[payload['sa_id']].call(action,payload,timeout)
        return Queue(),backends

    def test_five_consecutive_full_closures_are_serial_and_write_each_once(self):
        queue,backends=self.make_five_queue();untouched=self.roster.records[-1]
        result=run_metadata_batch(self.base,self.roster,self.roster.records[:5],queue,self.stop,
                                 client=self.client,excel_builder=self.builder)
        self.assertFalse(result.halted,result.reason)
        self.assertEqual(result.remaining,0)
        self.assertEqual([out['status'] for out in result.outcomes],['done']*5)
        self.assertTrue(all(r.done and r.remark=='已入库' for r in result.roster.records[:5]))
        self.assertEqual(result.roster.records[-1],untouched)
        for index,record in enumerate(result.roster.records[:5]):
            backend=backends[record.sa_id]
            for action in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
                self.assertEqual(backend.calls.count(action),1,(index,action))
            if index:
                previous=result.roster.records[index-1]
                self.assertLess(queue.calls.index((previous.sa_id,'complete')),
                                queue.calls.index((record.sa_id,'status')))

    def test_one_known_author_review_does_not_stop_remaining_papers(self):
        from zero_match import WorkflowStore
        queue,backends=self.make_five_queue()
        deferred=self.roster.records[1]
        backends[deferred.sa_id].roles_match=False
        result=run_metadata_batch(self.base,self.roster,self.roster.records[:5],queue,self.stop,
                                 client=self.client,excel_builder=self.builder)
        self.assertFalse(result.halted,result.reason)
        self.assertEqual([out['status'] for out in result.outcomes],['done','skip','done','done','done'])
        saved=result.roster.records[1]
        self.assertTrue(saved.skipped);self.assertFalse(saved.done)
        self.assertEqual(saved.remark,'通讯作者待核验')
        self.assertEqual(WorkflowStore(self.base/'runtime/metadata-imports').get(saved)['phase'],'linked')
        self.assertNotIn('submit_claim',backends[deferred.sa_id].calls)
        self.assertNotIn('complete',backends[deferred.sa_id].calls)
        before={key:list(backend.calls) for key,backend in backends.items()}
        again=run_metadata_batch(self.base,result.roster,result.roster.records[:5],queue,self.stop,
                                client=self.client,excel_builder=self.builder)
        self.assertFalse(again.halted,again.reason)
        self.assertEqual(before,{key:backend.calls for key,backend in backends.items()})

    def test_author_review_releases_native_detail_before_next_exact_status(self):
        queue,backends=self.make_five_queue()
        backends[self.roster.records[0].sa_id].roles_match=False
        original=queue.call
        opened=None
        def call(action,payload,timeout=90):
            nonlocal opened
            if action=='status' and opened and opened!=payload['sa_id']:
                raise SafetyStop('网页有未关闭的窗口')
            answer=original(action,payload,timeout)
            if action=='search':
                opened=None if payload.get('close_detail') is True else payload['sa_id']
                answer['detail_released']=payload.get('close_detail') is True
            return answer
        queue.call=call
        result=run_metadata_batch(self.base,self.roster,self.roster.records[:5],queue,self.stop,
                                 client=self.client,excel_builder=self.builder)
        self.assertFalse(result.halted,result.reason)
        self.assertEqual([item['status'] for item in result.outcomes],['skip','done','done','done','done'])
        self.assertIsNone(opened)
        self.assertNotIn('complete',backends[self.roster.records[0].sa_id].calls)

    def test_known_duplicate_does_not_stop_later_papers_or_rewrite_done_rows(self):
        queue,backends=self.make_five_queue()
        first=self.roster.records[0]
        backends[first.sa_id].pushed_counts={'increase':0,'duplicateIncreaseUpdate':1}
        result=run_metadata_batch(self.base,self.roster,self.roster.records[:5],queue,self.stop,
                                 client=self.client,excel_builder=self.builder)
        self.assertFalse(result.halted,result.reason)
        self.assertEqual([out['status'] for out in result.outcomes],['skip','done','done','done','done'])
        self.assertTrue(result.roster.records[0].skipped)
        self.assertNotIn('complete',backends[first.sa_id].calls)
        before={key:list(backend.calls) for key,backend in backends.items()}
        again=run_metadata_batch(self.base,result.roster,result.roster.records[:5],queue,self.stop,
                                client=self.client,excel_builder=self.builder)
        self.assertFalse(again.halted,again.reason)
        self.assertEqual(before,{key:backend.calls for key,backend in backends.items()})

    def test_resume_old_deferred_drawer_only_reads_it_and_continues_other_papers(self):
        queue,backends=self.make_five_queue()
        first=self.roster.records[0]
        backends[first.sa_id].roles_match=False
        partial=run_metadata_batch(self.base,self.roster,[first],queue,self.stop,
                                  client=self.client,excel_builder=self.builder)
        self.assertTrue(partial.roster.records[0].skipped)
        previous_calls=list(backends[first.sa_id].calls)
        original=queue.call
        opened=first.sa_id
        def call(action,payload,timeout=90):
            nonlocal opened
            if action=='status' and opened and opened!=payload['sa_id']:
                raise SafetyStop('网页有未关闭的窗口')
            answer=original(action,payload,timeout)
            if action=='search':
                opened=None if payload.get('close_detail') is True else payload['sa_id']
                answer['detail_released']=payload.get('close_detail') is True
            return answer
        queue.call=call
        result=run_metadata_batch(self.base,partial.roster,partial.roster.records[:5],queue,self.stop,
                                 client=self.client,excel_builder=self.builder)
        self.assertFalse(result.halted,result.reason)
        self.assertEqual([out['status'] for out in result.outcomes],['already_skipped','done','done','done','done'])
        self.assertEqual(backends[first.sa_id].calls[len(previous_calls):],['status','search'])
        self.assertEqual(result.roster.records[0],partial.roster.records[0])
        self.assertIsNone(opened)

    def test_resume_completed_then_pending_paper_only_reads_previous_batch(self):
        for index,pause_at in enumerate(('status','import_excel_upload','import_excel_submit')):
            with self.subTest(pause_at=pause_at):
                if index:self.setUp()
                book=load_workbook(self.roster.path);sheet=book.active
                for cell,value in {'B3':'谭勋策','D3':'Second synthetic paper','E3':'10.1234/second',
                                   'G3':'00001','H3':0,'I3':'','K3':''}.items():sheet[cell]=value
                book.save(self.roster.path);book.close()
                current=read_roster(self.roster.path);first,second=current.records
                _,cr1,oa1=example();cr2=copy.deepcopy(cr1);oa2=copy.deepcopy(oa1)
                cr2.update(DOI=second.doi,title=[second.title]);oa2.update(doi='https://doi.org/'+second.doi,title=second.title)
                self.client.get=Mock(side_effect=[{'message':cr1},oa1,{'message':cr2},oa2])
                other=Backend(second,merge_metadata(second,cr2,oa2));backends={first.sa_id:self.backend,second.sa_id:other}
                if pause_at!='status':other.lost=pause_at
                class Queue:
                    pause=pause_at=='status'
                    def call(queue,action,payload,timeout=90):
                        if queue.pause and payload['sa_id']==second.sa_id:
                            queue.pause=False;raise SafetyStop('第二篇启动前暂停')
                        return backends[payload['sa_id']].call(action,payload,timeout)
                queue=Queue()
                before=run_metadata_batch(self.base,current,[first,second],queue,self.stop,client=self.client,excel_builder=self.builder)
                self.assertTrue(before.halted);self.assertTrue(before.roster.records[0].done)
                previous=len(self.backend.calls)
                after=run_metadata_batch(self.base,before.roster,before.roster.records,queue,self.stop,
                                         client=self.client,excel_builder=self.builder)
                self.assertFalse(after.halted,after.reason);self.assertTrue(all(r.done for r in after.roster.records))
                self.assertEqual(self.backend.calls[previous:],['import_excel_check'] if pause_at=='status' else [])
                for backend in (self.backend,other):
                    for action in ('import_excel_upload','import_excel_submit','import_excel_push','link','complete'):
                        self.assertEqual(backend.calls.count(action),1,(action,backend.calls))


if __name__=='__main__':unittest.main()

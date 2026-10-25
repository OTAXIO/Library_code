"""Human local-status overrides on synthetic workbooks; never the real roster."""
import copy
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import Mock, patch
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from core import SafetyStop, file_hash, read_roster
from manual_review import STATES, write_review
from roster_write import column_name, migrate_status_column, mark_complete
from workflow_service import WorkflowService
from tests.test_roster_write import make_roster


class ManualReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.path = self.base / 'list.xlsx'
        make_roster(self.path, flags=(None, None))
        book = load_workbook(self.path)
        sheet = book.active
        sheet['B2'] = '谭勋策'
        sheet['A2'] = '资料冲突；原备注保留'
        sheet['H2'] = 3  # Local verification covers tasks outside automatic workflows.
        sheet['I2'] = '1234567890123456789,1234567890123456790,1234567890123456791'
        sheet['K2'] = '第一作者不一致'
        sheet['C2'].number_format = '@'
        sheet['C2'].font = Font(name='Arial', bold=True, color='123456')
        sheet['C2'].border = Border(bottom=Side(style='thin', color='000000'))
        source_column = sheet.max_column + 1
        sheet.cell(1, source_column, '数据来源')
        sheet.cell(2, source_column, 'Crossref+OpenAlex')
        book.save(self.path)
        book.close()
        self.roster = migrate_status_column(read_roster(self.path))
        self.service = WorkflowService(self.base)
        self.bridge = SimpleNamespace(pending=None, call=Mock(side_effect=AssertionError('No browser calls')))

    def save(self, state):
        result = self.service.review_selected(self.roster, '谭勋策', self.roster.records[0].sa_id, state, self.bridge)
        self.roster = result.roster
        return result

    def assert_state(self, state):
        record = self.roster.records[0]
        self.assertEqual((record.done, record.skipped), (state == 'done', state == 'skipped'))
        book = load_workbook(self.path)
        try:
            self.assertEqual(book.active.cell(2, self.roster.completion_column).value, STATES[state][1])
            fill = book.active['C2'].fill
            if state == 'pending':
                self.assertIsNone(fill.patternType)
            else:
                self.assertEqual(fill.fgColor.rgb, 'FFFFFF00' if state == 'done' else 'FFFF0000')
        finally:
            book.close()

    def test_every_transition_preserves_notes_source_and_business_identity(self):
        original = self.roster.records[0]
        untouched = self.roster.records[1:]
        for before in STATES:
            self.save(before)
            for after in STATES:
                with self.subTest(before=before, after=after):
                    self.save(before)
                    result = self.save(after)
                    self.assert_state(after)
                    self.assertEqual(self.roster.records[0], replace(original, done=after == 'done', skipped=after == 'skipped'))
                    self.assertEqual(self.roster.records[0].staff_id, '00001')
                    self.assertEqual(self.roster.records[1:], untouched)
                    self.assertIn('仅名单', result.summary)
        self.bridge.call.assert_not_called()

    def test_atomic_status_and_style_change_preserves_other_zip_members(self):
        before_bytes = self.path.read_bytes()
        before = load_workbook(self.path)
        font, border = copy.copy(before.active['C2'].font), copy.copy(before.active['C2'].border)
        values = [[c.value for c in row] for row in before.active]
        before.close()
        saved = write_review(self.roster, self.roster.records[0], 'done')
        self.assertEqual(saved.backup.read_bytes(), before_bytes)
        self.assertEqual(set(saved.cells), {'C2', f'{column_name(self.roster.completion_column)}2'})
        after = load_workbook(self.path)
        try:
            self.assertEqual(after.active['C2'].font, font)
            self.assertEqual(after.active['C2'].border, border)
            self.assertEqual(after.active['C2'].number_format, '@')
            values[1][self.roster.completion_column - 1] = 1
            self.assertEqual([[c.value for c in row] for row in after.active], values)
        finally:
            after.close()
        with zipfile.ZipFile(saved.backup) as a, zipfile.ZipFile(self.path) as b:
            self.assertEqual(a.namelist(), b.namelist())
            self.assertEqual({n for n in a.namelist() if a.read(n) != b.read(n)},
                             {'xl/worksheets/sheet1.xml', 'xl/styles.xml'})

    def test_noop_is_idempotent_but_human_verification_is_logged(self):
        self.save('done')
        checksum = file_hash(self.path)
        saved = write_review(self.roster, self.roster.records[0], 'done')
        self.assertIsNone(saved.backup)
        self.assertEqual(saved.cells, ())
        self.save('done')
        self.assertEqual(file_hash(self.path), checksum)
        entries = self.service.operation_log.read()
        self.assertEqual(len(entries), 2)
        self.assertEqual(entries[-1]['sa_id'], self.roster.records[0].sa_id)
        self.assertIn('已完成 → 已完成', entries[-1]['detail'])
        self.assertNotIn('已入库', entries[-1]['detail'])

    def test_business_intent_journals_are_never_reset_by_reopening_local_task(self):
        journal = self.base / 'runtime/metadata-imports/zero-match.sqlite3'
        journal.parent.mkdir(parents=True)
        marker = b'Existing import_intent / done checkpoints must survive'
        journal.write_bytes(marker)
        self.save('done')
        self.save('pending')
        self.assertEqual(journal.read_bytes(), marker)
        self.assertNotIn('已入库', self.roster.records[0].remark)

    def test_invalid_owner_id_state_duplicate_or_wrong_path_cannot_write(self):
        original = self.path.read_bytes()
        for owner, sid, state in (('其他人', 'demo-001', 'done'), ('', 'demo-001', 'done'),
                ('谭勋策', 'demo-002', 'done'), ('谭勋策', 'absent', 'done'),
                ('谭勋策', 'demo-001', 'invalid'), ('谭勋策', 'demo-001', True)):
            with self.subTest(owner=owner, sid=sid, state=state), self.assertRaises(SafetyStop):
                self.service.review_selected(self.roster, owner, sid, state)
            self.assertEqual(self.path.read_bytes(), original)
        duplicate = copy.copy(self.roster)
        duplicate.records = [*self.roster.records, replace(self.roster.records[0], row=9)]
        with self.assertRaises(SafetyStop):
            self.service.review_selected(duplicate, '谭勋策', 'demo-001', 'done')
        elsewhere = self.base / 'another'
        elsewhere.mkdir()
        (elsewhere / 'list.xlsx').write_bytes(original)
        with self.assertRaises(SafetyStop):
            self.service.review_selected(read_roster(elsewhere / 'list.xlsx'), '谭勋策', 'demo-001', 'done')

    def test_stale_excel_lock_pending_browser_and_commit_failure_preserve_original(self):
        original = self.path.read_bytes()
        self.bridge.pending = {'command': 'unknown receipt'}
        with self.assertRaises(SafetyStop):
            self.save('done')
        self.bridge.pending = None
        lock = self.path.with_name('~$list.xlsx')
        lock.write_bytes(b'Excel lock')
        with self.assertRaises(SafetyStop):
            self.save('done')
        lock.unlink()
        with patch('manual_review._commit_verified', side_effect=PermissionError('locked')), self.assertRaises(SafetyStop):
            self.save('done')
        self.assertEqual(self.path.read_bytes(), original)
        self.assertFalse(list(self.base.glob('.list-review-*')))
        self.assertFalse(list(self.base.glob('*.assistant.lock')))
        self.path.write_bytes(original + b'changed')
        with self.assertRaisesRegex(SafetyStop, '已改变'):
            self.save('done')

    def test_pending_preserves_non_workflow_fill_and_repeated_clear_of_missing_flag(self):
        book = load_workbook(self.path)
        book.active['C2'].fill = PatternFill('solid', fgColor='FF123456')
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)
        self.save('pending')
        book = load_workbook(self.path)
        self.assertEqual(book.active['C2'].fill.fgColor.rgb, 'FF123456')
        book.close()
        self.assertFalse(self.roster.records[0].done)

    def test_id_column_is_discovered_after_column_reordering(self):
        book = load_workbook(self.path)
        book.active.insert_cols(3)
        book.active['C1'] = '额外原字段'
        book.active['C2'] = '不要改'
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)
        self.save('skipped')
        book = load_workbook(self.path)
        self.assertEqual(book.active['D2'].fill.fgColor.rgb, 'FFFF0000')
        self.assertEqual(book.active['C2'].value, '不要改')
        book.close()

    def test_log_failure_reports_saved_result_without_replaying_browser(self):
        with patch.object(self.service.operation_log, 'record', side_effect=OSError('locked')):
            result = self.save('done')
        self.assert_state('done')
        self.assertIn('log.txt 未能保存', result.summary)
        self.bridge.call.assert_not_called()

    def test_original_automatic_completion_guard_is_still_enforced(self):
        self.save('done')
        with self.assertRaises(SafetyStop):
            mark_complete(self.roster, self.roster.records[0], note='已入库')


if __name__ == '__main__':
    unittest.main()

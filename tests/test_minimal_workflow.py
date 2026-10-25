"""Two-task facade + native UI, using only synthetic rosters and fake pages."""
import ast
import tempfile
import threading
import time
import tkinter as tk
import unittest
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from app import App
from core import Record, SafetyStop
from tests import test_zero_match as zero_fixtures
from workflow_service import TaskResult, WorkflowService


def records():
    base = Record(2, "谭勋策", "claim-pending", "Test", "", "", "00001", 1, "100", "待处理", "作者不一致", "1")
    return [base, replace(base, sa_id="claim-skip", row=3, skipped=True),
            replace(base, sa_id="claim-done", row=4, done=True),
            replace(base, sa_id="zero-pending", row=5, matches=0, reason="", item_ids=""),
            replace(base, sa_id="zero-skip", row=6, matches=0, reason="", skipped=True),
            replace(base, sa_id="other-owner", row=7, owner="其他人"),
            replace(base, sa_id="role-issue", row=8, reason="通讯作者不一致")]


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.service = WorkflowService(self.base)
        self.roster = SimpleNamespace(records=records(), assert_unchanged=Mock())

    def test_modes_scopes_and_owner_are_separate(self):
        visible = self.service.visible
        self.assertEqual([r.sa_id for r in visible(self.roster, "claim", "谭勋策", "pending")], ["claim-pending"])
        self.assertEqual([r.sa_id for r in visible(self.roster, "claim", "谭勋策", "skipped")], ["claim-skip"])
        self.assertEqual([r.sa_id for r in visible(self.roster, "claim", "谭勋策", "done")], ["claim-done"])
        self.assertEqual([r.sa_id for r in visible(self.roster, "import", "谭勋策", "pending")], ["zero-pending"])
        self.assertEqual([r.sa_id for r in visible(self.roster, "import", "谭勋策", "skipped")], ["zero-skip"])

    def test_invalid_range_or_other_owner_never_runs(self):
        for owner, count, scope in (("", 1, "pending"), ("其他人", 1, "pending"),
                                    ("谭勋策", 0, "pending"), ("谭勋策", True, "pending"),
                                    ("谭勋策", 101, "pending"), ("谭勋策", 1, "done")):
            with self.assertRaises(SafetyStop):
                self.service.select(self.roster, "claim", owner, count, scope)
        with self.assertRaisesRegex(SafetyStop, "请选择负责人"):
            self.service.select(self.roster, "claim", "", 1, "pending")

    def test_stale_roster_blocked(self):
        self.roster.assert_unchanged.side_effect = SafetyStop("名单已改变")
        with self.assertRaisesRegex(SafetyStop, "名单已改变"):
            self.service.select(self.roster, "claim", "谭勋策", 1, "pending")

    def test_claim_dispatch_has_only_extension_bridge(self):
        bridge = Mock(online=True)
        selected = self.service.select(self.roster, "claim", "谭勋策", 1, "pending")
        result = SimpleNamespace(roster=self.roster, completed_ids=("claim-pending",), synced_ids=(),
                                 skipped={}, halted=False, cancelled=False)
        with patch("workflow_service.run_claim_batch", return_value=result) as run:
            saved = self.service.run(self.roster, selected, "claim", "pending", bridge, threading.Event(), Mock())
        self.assertIs(run.call_args.args[2].transport, bridge)
        self.assertIn("完成 1", saved.summary)

    def test_skip_only_import_explains_reason_without_claiming_website_sync(self):
        bridge = Mock(online=True)
        selected = self.service.select(self.roster, "import", "谭勋策", 1, "pending")
        result = SimpleNamespace(roster=self.roster, remaining=0, halted=False, reason='',
            outcomes=[{'sa_id': selected[0].sa_id, 'status': 'skip', 'message': '未确认交大署名'}])
        with patch('metadata_pipeline.run_metadata_batch', return_value=result):
            saved = self.service.run(self.roster, selected, 'import', 'pending', bridge,
                                     threading.Event(), Mock(), metadata_import=True)
        self.assertIn('跳过原因：未确认交大署名', saved.summary)
        self.assertNotIn('网页与名单已同步', saved.summary)
        self.assertIn('完成 0', saved.summary)

    def test_offline_or_tampered_queue_cannot_dispatch(self):
        for selected, online in (([self.roster.records[5]], True), ([self.roster.records[0]], False)):
            with patch("workflow_service.run_claim_batch") as run, self.assertRaises(SafetyStop):
                self.service.run(self.roster, selected, "claim", "pending", Mock(online=online), threading.Event(), Mock())
            run.assert_not_called()

    def test_manual_selection_can_select_one_nonfirst_owned_zero_match(self):
        later = replace(self.roster.records[3], sa_id="zero-later", row=9)
        self.roster.records.append(later)
        selected = self.service.select_manual(self.roster, "谭勋策", "pending", later.sa_id)
        self.assertEqual(selected, later)
        for sa_id in ("claim-pending", "other-owner", "zero-skip", "missing"):
            with self.assertRaises(SafetyStop):
                self.service.select_manual(self.roster, "谭勋策", "pending", sa_id)

    def test_manual_dispatch_is_single_import_only_and_not_a_bulk_backdoor(self):
        bridge = Mock(online=True)
        record = self.roster.records[3]
        with patch("workflow_service.run_batch") as run:
            for chosen, mode in (([record, record], "import"), ([self.roster.records[0]], "claim")):
                with self.assertRaises(SafetyStop):
                    self.service.run(self.roster, chosen, mode, "pending", bridge,
                                     threading.Event(), Mock(), manual_export=True)
            run.assert_not_called()

    def test_manual_skip_selection_requires_one_pending_owned_visible_record(self):
        selected = self.service.select_skip(self.roster, "import", "谭勋策", "pending", "zero-pending")
        self.assertEqual(selected.sa_id, "zero-pending")
        for mode, owner, scope, sa_id in (("import", "其他人", "pending", "other-owner"),
                ("import", "", "pending", "zero-pending"), ("import", "谭勋策", "skipped", "zero-skip"),
                ("claim", "谭勋策", "pending", "claim-done"),
                ("import", "谭勋策", "pending", "claim-pending"),
                ("import", "谭勋策", "pending", "missing")):
            with self.assertRaises(SafetyStop):
                self.service.select_skip(self.roster, mode, owner, scope, sa_id)
        self.roster.records.append(replace(selected, row=10))
        with self.assertRaises(SafetyStop):
            self.service.select_skip(self.roster, "import", "谭勋策", "pending", "zero-pending")

    def test_all_production_python_has_no_retired_dependency(self):
        root = Path(__file__).resolve().parents[1]
        forbidden = {"wos_browser", "paper_classify", "model_review", "submission_prepare", "settings_panel",
                     "classify_app", "wos_import_panel", "automation_panel", "pilot", "import_channels"}
        for path in root.glob("*.py"):
            tree = ast.parse(path.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                names = [node.module] if isinstance(node, ast.ImportFrom) else [n.name for n in node.names] if isinstance(node, ast.Import) else []
                self.assertFalse(forbidden.intersection(names), path.name)
            self.assertNotIn("bsk.exe", path.read_text(encoding="utf-8"))


class ExtensionOnlyImportTests(unittest.TestCase):
    def setUp(self):
        zero_fixtures.ZeroMatchTests.setUp(self)

    def test_full_facade_downloads_imports_and_writes_both_notes_without_bsk(self):
        self.sa.online = True
        original = self.sa.call
        def call(action, payload, timeout=75):
            if action == "wos_diagnose":
                from tests.test_wos_batch import PreflightTests
                return PreflightTests().capabilities()
            if action.startswith("wos_"):
                return self.download.call(action, payload, timeout)
            return original(action, payload, timeout)
        self.sa.call = call
        service = WorkflowService(self.folder)
        with patch("workflow_service.default_store", return_value=self.download.store):
            result = service.run(self.roster, self.roster.records[:1], "import", "pending", self.sa, self.stop, Mock())
        self.assertIn("完成 1", result.summary)
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(result.roster.records[0].source, "WOS")
        self.assertEqual(result.roster.records[0].remark, self.sa.rows[self.roster.records[0].sa_id]["remark"])
        self.assertTrue((self.folder / "runtime" / "wos-imports" / "zero-match.sqlite3").exists())
        self.assertGreater(len(service.operation_log.read()), 0)

    def test_manual_locate_does_not_import_or_write_roster(self):
        from core import file_hash
        self.sa.online = True
        original = self.sa.call
        def call(action, payload, timeout=75):
            if action == "wos_diagnose":
                from tests.test_wos_batch import PreflightTests
                return PreflightTests().capabilities()
            if action == "wos_search":
                return self.download.call(action, payload, timeout)
            return original(action, payload, timeout)
        self.sa.call = call
        before = file_hash(self.path)
        service = WorkflowService(self.folder)
        record = self.roster.records[1]  # A selected nonfirst row, never a bulk queue.
        with patch("workflow_service.default_store", return_value=self.download.store):
            result = service.locate_import(self.roster, record, "谭勋策", "pending", self.sa, self.stop, Mock())
        self.assertIn("手动导出", result.summary)
        self.assertEqual(file_hash(self.path), before)
        self.assertNotIn("import_upload", self.sa.calls)
        self.assertNotIn("complete", self.sa.calls)
        self.assertEqual(self.download.calls, ["wos_search"])
        self.assertEqual(self.download.store.get(record)["phase"], "export_preparing")

    def test_manual_intake_waits_without_flags_or_automatic_wos_commands(self):
        from core import file_hash
        self.sa.online = True
        before = file_hash(self.path)
        service = WorkflowService(self.folder)
        with patch("workflow_service.default_store", return_value=self.download.store):
            result = service.run(self.roster, [self.roster.records[1]], "import", "pending",
                                 self.sa, self.stop, Mock(), manual_export=True)
        self.assertIn("手动导出", result.summary)
        self.assertEqual(file_hash(self.path), before)
        self.assertFalse(result.roster.records[1].skipped)
        self.assertFalse(result.roster.records[1].done)
        self.assertFalse(any(action.startswith("wos_") for action in self.sa.calls))
        self.assertNotIn("import_upload", self.sa.calls)

    def test_manual_txt_imports_one_selected_nonfirst_paper_and_closes_both_notes(self):
        from tests.test_automation import sample
        self.sa.online = True
        manual = self.download.store.root / "未关联下载" / "savedrecs.txt"
        manual.parent.mkdir()
        manual.write_bytes(sample())
        service = WorkflowService(self.folder)
        selected = self.roster.records[1]
        with patch("workflow_service.default_store", return_value=self.download.store):
            result = service.run(self.roster, [selected], "import", "pending", self.sa,
                                 self.stop, Mock(), manual_export=True)
        self.assertIn("完成 1", result.summary)
        self.assertTrue(result.roster.records[1].done)
        self.assertFalse(result.roster.records[0].done)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertFalse(result.roster.records[2].done)
        self.assertEqual(result.roster.records[1].remark, "已入库")
        self.assertEqual(self.sa.rows[selected.sa_id]["remark"], "已入库")
        self.assertEqual(self.sa.rows[selected.sa_id]["markStatus"], "已处理")
        self.assertEqual(self.sa.calls.count("import_upload"), 1)
        self.assertEqual(self.sa.calls.count("import_submit"), 1)
        self.assertEqual(self.sa.calls.count("import_push"), 1)
        self.assertFalse(any(action.startswith("wos_") for action in self.sa.calls))
        self.assertEqual(manual.read_bytes(), sample())
        self.assertTrue((self.folder / "runtime" / "wos-imports" /
                         __import__("hashlib").sha256(sample()).hexdigest()).with_suffix(".txt").exists())

    def test_manual_locate_nonzero_only_flags_skip_without_notes_or_wos(self):
        self.sa.online = True
        first = self.roster.records[0]
        self.sa.rows[first.sa_id].update(matchCount=2, remark="原有批注")
        service = WorkflowService(self.folder)
        with patch("workflow_service.default_store", return_value=self.download.store):
            result = service.locate_import(self.roster, first, "谭勋策", "pending", self.sa, self.stop, Mock())
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark, "")
        self.assertEqual(self.sa.rows[first.sa_id]["remark"], "原有批注")
        self.assertEqual(self.sa.rows[first.sa_id]["markStatus"], "待处理")
        self.assertFalse(any(action.startswith("wos_") for action in self.sa.calls))

    def test_manual_locate_zero_result_uses_real_miss_note_without_import(self):
        self.sa.online = True
        original = self.sa.call
        def call(action, payload, timeout=75):
            if action == "wos_diagnose":
                from tests.test_wos_batch import PreflightTests
                return PreflightTests().capabilities()
            if action == "wos_search":
                raise SafetyStop("[WOS 已暂停] WOS 未找到记录；这不等于未发表")
            return original(action, payload, timeout)
        self.sa.call = call
        service = WorkflowService(self.folder)
        first = self.roster.records[0]
        with patch("workflow_service.default_store", return_value=self.download.store):
            result = service.locate_import(self.roster, first, "谭勋策", "pending", self.sa, self.stop, Mock())
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark, "wos未收录")
        self.assertFalse(result.roster.records[0].done)
        self.assertNotIn("import_upload", self.sa.calls)


class ManualSkipTests(unittest.TestCase):
    def setUp(self):
        zero_fixtures.ZeroMatchTests.setUp(self)
        self.service = WorkflowService(self.folder)
        self.selected = self.roster.records[1]

    def skip(self, bridge=None):
        return self.service.skip_selected(self.roster, "import", "谭勋策", "pending",
                                          self.selected.sa_id, bridge)

    def test_offline_skip_changes_only_selected_tail_flag_and_keeps_backups(self):
        import zipfile
        from core import read_roster
        from roster_write import worksheet_member, patch_cell, column_name
        with zipfile.ZipFile(self.path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
            member = worksheet_member(archive, self.roster.sheet_name)
        result = self.skip()
        expected = replace(self.selected, skipped=True)
        self.assertEqual(result.roster.records, [self.roster.records[0], expected] + self.roster.records[2:])
        self.assertEqual(self.sa.calls, [])
        self.assertEqual(self.download.calls, [])
        with zipfile.ZipFile(self.path) as archive:
            for name, raw in parts.items():
                wanted = (patch_cell(raw, f"{column_name(self.roster.completion_column)}{self.selected.row}",
                                     self.selected.row, 2) if name == member else raw)
                self.assertEqual(archive.read(name), wanted, name)
        backup = self.folder / "runtime" / "backups" / f"list-{self.roster.sha256}.xlsx"
        self.assertEqual(read_roster(backup).records, self.roster.records)
        self.assertNotIn(self.selected.sa_id,
                         [r.sa_id for r in self.service.select(result.roster, "import", "谭勋策", 5, "pending")])
        self.assertEqual([r.sa_id for r in self.service.select(result.roster, "import", "谭勋策", 5, "skipped")],
                         [self.selected.sa_id])
        self.assertEqual(self.service.operation_log.read()[-1]["sa_id"], self.selected.sa_id)

    def test_unknown_backend_write_or_closure_cannot_be_skipped_and_journals_unchanged(self):
        from core import file_hash
        from automation import ImportStore
        from zero_match import WorkflowStore
        for store, phases in ((ImportStore(self.folder / "runtime" / "wos-imports"),
                ("upload_intent", "uploaded", "import_intent", "imported", "push_intent", "pushed", "unknown")),
                (WorkflowStore(self.folder / "runtime" / "wos-imports"),
                ("link_intent", "linked", "complete_intent", "non_sjtu_complete_intent", "done", "unknown"))):
            for phase in phases:
                store.save(self.selected, {"phase": phase})
                before, journal = file_hash(self.path), file_hash(store.path)
                with self.assertRaisesRegex(SafetyStop, "断点待核验"):
                    self.skip()
                self.assertEqual(file_hash(self.path), before)
                self.assertEqual(file_hash(store.path), journal)
            # Let the next journal be the gate, not a leftover import intent.
            store.save(self.selected, {"phase": "exported" if isinstance(store, ImportStore) else "downloaded"})

    def test_only_downloaded_or_exported_checkpoints_allow_local_skip(self):
        from core import file_hash
        from automation import ImportStore
        from zero_match import WorkflowStore
        stores = [ImportStore(self.folder / "runtime" / "wos-imports"),
                  WorkflowStore(self.folder / "runtime" / "wos-imports")]
        for store, phase in zip(stores, ("exported", "downloaded")):
            store.save(self.selected, {"phase": phase})
        hashes = [file_hash(store.path) for store in stores]
        result = self.skip()
        self.assertTrue(result.roster.records[1].skipped)
        self.assertEqual([file_hash(store.path) for store in stores], hashes)

    def test_busy_bridge_stale_roster_excel_lock_or_corrupt_journal_never_writes(self):
        from core import file_hash
        before = file_hash(self.path)
        with self.assertRaisesRegex(SafetyStop, "仍在执行"):
            self.skip(SimpleNamespace(pending={"action": "status"}))
        self.assertEqual(file_hash(self.path), before)
        lock = self.path.with_name("~$" + self.path.name)
        lock.touch()
        with self.assertRaisesRegex(SafetyStop, "Excel"):
            self.skip()
        lock.unlink()
        with patch.object(self.roster, "assert_unchanged", side_effect=SafetyStop("名单已改变")):
            with self.assertRaisesRegex(SafetyStop, "名单已改变"):
                self.skip()
        journal = self.folder / "runtime" / "wos-imports" / "imports.sqlite3"
        journal.parent.mkdir(parents=True, exist_ok=True)
        journal.write_bytes(b"not sqlite")
        with self.assertRaisesRegex(SafetyStop, "断点无法核验"):
            self.skip()
        self.assertEqual(file_hash(self.path), before)

    def test_foreign_keyed_journal_and_wrong_roster_path_block_skip(self):
        from automation import ImportStore
        store = ImportStore(self.folder / "runtime" / "wos-imports")
        store.save(replace(self.selected, title="Changed title"), {"phase": "exported"})
        with self.assertRaisesRegex(SafetyStop, "断点待核验"):
            self.skip()
        with patch("workflow_service.fixed_roster_path", return_value=self.folder / "another-list.xlsx"):
            with self.assertRaisesRegex(SafetyStop, "当前 code/list"):
                self.skip()

    def test_already_skipped_or_completed_row_remains_readonly(self):
        from core import file_hash
        from roster_write import mark_complete
        result = self.skip()
        before = file_hash(self.path)
        with self.assertRaises(SafetyStop):
            self.service.skip_selected(result.roster, "import", "谭勋策", "pending", self.selected.sa_id)
        self.assertEqual(file_hash(self.path), before)
        first = result.roster.records[0]
        completed = mark_complete(result.roster, first, note="已入库").roster
        before = file_hash(self.path)
        with self.assertRaises(SafetyStop):
            self.service.skip_selected(completed, "import", "谭勋策", "pending", first.sa_id)
        self.assertEqual(file_hash(self.path), before)

    def test_existing_local_and_web_notes_survive_manual_skip(self):
        from roster_write import patch_text_cell, worksheet_member
        from core import read_roster
        import zipfile
        # A disposable fixture already has a human note in both remark columns.
        with zipfile.ZipFile(self.path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
            member = worksheet_member(archive, self.roster.sheet_name)
        changed = patch_text_cell(parts[member], f"A{self.selected.row}", self.selected.row, "人工备注保留")
        with zipfile.ZipFile(self.path, "w") as archive:
            for name, raw in parts.items():
                archive.writestr(name, changed if name == member else raw)
        self.roster = read_roster(self.path)
        self.selected = self.roster.records[1]
        result = self.skip()
        self.assertEqual(result.roster.records[1].remark, "人工备注保留")
        self.assertEqual(result.roster.records[1].source, self.selected.source)
        self.assertEqual(self.sa.calls, [])

    def test_log_failure_reports_saved_skip_without_pretending_it_rolled_back(self):
        with patch.object(self.service.operation_log, "record", side_effect=OSError("test log failure")):
            result = self.skip()
        self.assertTrue(result.roster.records[1].skipped)
        self.assertIn("跳过已保存", result.summary)


class MinimalUITests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.service = WorkflowService(self.temp.name)
        self.root = tk.Tk()
        self.root.withdraw()
        self.app = App(self.root, service=self.service, auto_load=False, bridge=Mock(online=True))
        self.addCleanup(lambda: self.app.close() if not self.app.closed and not self.app.busy else self.root.destroy())
        self.app.roster = SimpleNamespace(records=records(), assert_unchanged=Mock())
        self.app.owner.set("谭勋策")
        self.app.render()
        self.root.update_idletasks()

    def test_only_two_tabs_topmost_and_no_runtime_created_by_opening(self):
        self.assertEqual([self.app.notebook.tab(t, "text") for t in self.app.notebook.tabs()], ["自动认领", "自动导入"])
        self.assertTrue(self.root.attributes("-topmost"))
        self.assertFalse((self.service.base / "runtime").exists())

    def test_completed_view_is_readonly_and_skipped_retry_explicit(self):
        self.app.scope.set("已完成")
        self.app.render()
        self.assertEqual(self.app.trees["claim"].get_children(), ("claim-done",))
        self.assertEqual(str(self.app.start_button["state"]), "disabled")
        self.assertEqual(str(self.app.skip_button["state"]), "disabled")
        self.app.scope.set("重试跳过项")
        self.app.render()
        self.assertEqual(self.app.trees["claim"].get_children(), ("claim-skip",))
        self.assertEqual(self.app.trees["import"].get_children(), ("zero-skip",))
        self.assertEqual(str(self.app.skip_button["state"]), "disabled")

    def test_busy_locks_controls_and_close_requests_safe_pause(self):
        self.app.set_busy(True)
        self.assertEqual(str(self.app.owner_box["state"]), "disabled")
        self.assertEqual(str(self.app.skip_button["state"]), "disabled")
        self.assertEqual(str(self.app.review_button["state"]), "disabled")
        self.assertEqual(self.app.notebook.tab(self.app.pages["import"], "state"), "disabled")
        with patch("app.messages.showwarning"):
            self.app.close()
        self.assertTrue(self.app.stop.is_set())
        self.assertFalse(self.app.closed)
        self.app.set_busy(False)

    def test_manual_review_window_lists_all_owned_tasks_and_filters_three_states(self):
        self.app.open_manual_review()
        dialog = self.app.review_dialog
        self.root.update()
        self.assertEqual(len(dialog.tree.get_children()), 6)
        dialog.query.set('role-issue')
        self.assertEqual(dialog.tree.get_children(), ('8',))
        dialog.query.set('')
        for label, expected in (('已完成', ('4',)), ('已跳过', ('3', '6')), ('未完成', ('2', '5', '8'))):
            dialog.scope.set(label)
            self.assertEqual(dialog.tree.get_children(), expected)
        dialog.close()

    def test_manual_review_save_uses_selected_id_and_refreshes_main_view(self):
        self.app.trees['claim'].selection_set('claim-pending')
        self.app.open_manual_review()
        dialog = self.app.review_dialog
        self.root.update()
        self.assertEqual(dialog.selected().sa_id, 'claim-pending')
        chosen = replace(self.app.roster.records[0], done=True)
        updated = SimpleNamespace(records=[chosen] + self.app.roster.records[1:], assert_unchanged=Mock())
        result = TaskResult(updated, {chosen.sa_id: '手动核验已保存'}, '手动核验已保存')
        dialog.choice.set('已完成')
        with patch.object(self.service, 'review_selected', return_value=result) as review:
            dialog.save_button.invoke()
        self.assertEqual(review.call_args.args[1:4], ('谭勋策', 'claim-pending', 'done'))
        self.assertIs(self.app.roster, updated)
        self.assertNotIn(chosen.sa_id, self.app.trees['claim'].get_children())
        self.assertEqual(dialog.tree.set('2', 'status'), '已完成')
        self.assertFalse(self.app.busy)
        self.assertEqual(self.app.status.get(), '手动核验已保存')
        dialog.close()

    def test_manual_review_cancel_or_failure_keeps_original_state(self):
        original = self.app.roster
        self.app.open_manual_review()
        dialog = self.app.review_dialog
        with patch.object(self.service, 'review_selected') as review:
            dialog.close()
        review.assert_not_called()
        self.app.open_manual_review()
        dialog = self.app.review_dialog
        dialog.tree.selection_set('2')
        self.root.update()
        dialog.choice.set('已跳过')
        with patch.object(self.service, 'review_selected', side_effect=SafetyStop('名单已改变')), patch('manual_review_ui.messages.showwarning'):
            dialog.save_button.invoke()
        self.assertIs(self.app.roster, original)
        self.assertIn('名单已改变', dialog.result.get())
        self.assertEqual(str(dialog.save_button['state']), 'normal')
        self.assertFalse(self.app.busy)
        dialog.close()

    def test_manual_review_cannot_open_while_busy_or_modify_other_owner(self):
        self.app.set_busy(True)
        self.app.open_manual_review()
        self.assertFalse(hasattr(self.app, 'review_dialog'))
        with self.assertRaises(SafetyStop):
            self.app.apply_manual_review('claim-pending', 'done')
        self.app.set_busy(False)
        self.app.owner.set('其他人')
        self.app.render()
        self.assertEqual(str(self.app.review_button['state']), 'disabled')
        with patch('app.messages.showwarning'):
            self.app.open_manual_review()
        self.assertFalse(hasattr(self.app, 'review_dialog'))

    def test_manual_review_real_button_saves_and_reopens_synthetic_workbook_offline(self):
        from openpyxl import load_workbook
        from tests.test_roster_write import make_roster
        path = self.service.base / 'list.xlsx'
        make_roster(path, flags=(None, None))
        book = load_workbook(path)
        book.active['B2'] = '谭勋策'
        book.active['A2'] = '原备注保留'
        book.save(path)
        book.close()
        self.app.roster = self.service.load()
        self.app.bridge = SimpleNamespace(pending=None, online=False, call=Mock(), close=Mock())
        self.app.render()
        self.app.review_button.invoke()
        dialog = self.app.review_dialog
        dialog.tree.selection_set('2')
        self.root.update()
        for label, flag in (('已完成', 1), ('已跳过', 2), ('未完成', None)):
            dialog.choice.set(label)
            dialog.save_button.invoke()
            self.root.update()
            book = load_workbook(path)
            self.assertEqual(book.active.cell(2, self.app.roster.completion_column).value, flag)
            self.assertEqual(book.active['A2'].value, '原备注保留')
            self.assertIsNone(book.active.cell(3, self.app.roster.completion_column).value)
            book.close()
        self.assertEqual(len(self.service.operation_log.read()), 3)
        self.app.bridge.call.assert_not_called()
        dialog.close()

    def test_manual_review_controls_fit_small_window_and_clear_selection_after_filter(self):
        self.app.open_manual_review()
        dialog = self.app.review_dialog
        dialog.window.attributes('-alpha', 0)
        dialog.window.geometry('600x430')
        self.root.update()
        for widget in (dialog.search, dialog.tree, dialog.detail, dialog.save_button, dialog.choices):
            self.assertLessEqual(widget.winfo_rooty() - dialog.window.winfo_rooty() + widget.winfo_height(),
                                 dialog.window.winfo_height())
            self.assertLessEqual(widget.winfo_rootx() - dialog.window.winfo_rootx() + widget.winfo_width(),
                                 dialog.window.winfo_width())
        dialog.tree.selection_set('2')
        self.root.update()
        dialog.query.set('no-match')
        self.assertEqual(dialog.choice.get(), '')
        self.assertEqual(str(dialog.save_button['state']), 'disabled')
        dialog.close()

    def test_threaded_dispatch_finishes_on_main_thread_without_second_start(self):
        held = threading.Event()
        def run(*args):
            held.wait(2)
            return TaskResult(self.app.roster, {"claim-pending": "fixture complete"}, "完成 1")
        with patch.object(self.service, "run", side_effect=run) as job:
            self.app.start()
            self.app.start()
            held.set()
            deadline = time.monotonic() + 3
            while self.app.busy and time.monotonic() < deadline:
                self.root.update()
                time.sleep(.01)
        self.assertEqual(job.call_count, 1)
        self.assertFalse(self.app.busy)
        self.assertEqual(self.app.status.get(), "完成 1")

    def test_other_owner_and_invalid_count_never_start(self):
        for owner, count in (("其他人", "5"), ("", "5"), ("谭勋策", "0"), ("谭勋策", "101")):
            self.app.owner.set(owner)
            self.app.limit.set(count)
            with patch("app.messages.showwarning"), patch.object(self.service, "run") as run:
                self.app.start()
            run.assert_not_called()
            self.assertFalse(self.app.busy)

    def test_manual_import_buttons_require_one_selection_and_disable_batch_limit(self):
        self.app.notebook.select(self.app.pages["import"])
        self.app.import_style.set("手动导出（一次一篇）")
        self.app.render()
        self.root.update_idletasks()
        self.assertEqual(str(self.app.limit_box["state"]), "disabled")
        self.assertIn("定位", self.app.start_button["text"])
        self.assertIn("接回 TXT", self.app.manual_import_button["text"])
        with patch("app.messages.showwarning"), patch.object(self.service, "locate_import") as locate:
            self.app.start()
        locate.assert_not_called()
        self.assertFalse(self.app.busy)

    def test_manual_selected_second_row_dispatches_only_that_record(self):
        later = replace(self.app.roster.records[3], sa_id="zero-later", row=9)
        self.app.roster.records.append(later)
        self.app.notebook.select(self.app.pages["import"])
        self.app.import_style.set("手动导出（一次一篇）")
        self.app.render()
        self.app.trees["import"].selection_set(later.sa_id)
        done = TaskResult(self.app.roster, {later.sa_id: "等待手动导出"}, "等待手动导出")
        with patch.object(self.service, "locate_import", return_value=done) as locate:
            self.app.start()
            deadline = time.monotonic() + 3
            while self.app.busy and time.monotonic() < deadline:
                self.root.update()
                time.sleep(.01)
        self.assertEqual(locate.call_args.args[1].sa_id, later.sa_id)
        self.assertEqual(locate.call_count, 1)
        self.assertFalse(self.app.busy)

    def test_small_window_keeps_actions_and_selected_detail_visible(self):
        self.root.attributes("-alpha", 0)
        self.root.deiconify()
        self.root.geometry("660x430")
        self.root.update()
        for widget in (self.app.start_button, self.app.pause_button, self.app.skip_button,
                       self.app.details["claim"], self.app.scope_box):
            self.assertTrue(widget.winfo_ismapped())
            bottom = widget.winfo_rooty() - self.root.winfo_rooty() + widget.winfo_height()
            right = widget.winfo_rootx() - self.root.winfo_rootx() + widget.winfo_width()
            self.assertLessEqual(bottom, self.root.winfo_height())
            self.assertLessEqual(right, self.root.winfo_width())
        self.root.withdraw()

    def test_manual_small_window_keeps_both_steps_visible_and_locked_during_run(self):
        self.app.notebook.select(self.app.pages["import"])
        self.app.import_style.set("手动导出（一次一篇）")
        self.app.render()
        self.root.attributes("-alpha", 0)
        self.root.deiconify()
        self.root.geometry("660x430")
        self.root.update()
        for widget in (self.app.start_button, self.app.manual_import_button, self.app.pause_button, self.app.skip_button,
                       self.app.import_style_box, self.app.details["import"]):
            self.assertTrue(widget.winfo_ismapped())
            self.assertLessEqual(widget.winfo_rooty() - self.root.winfo_rooty() + widget.winfo_height(),
                                 self.root.winfo_height())
            self.assertLessEqual(widget.winfo_rootx() - self.root.winfo_rootx() + widget.winfo_width(),
                                 self.root.winfo_width())
        self.app.set_busy(True)
        self.assertEqual(str(self.app.import_style_box["state"]), "disabled")
        self.assertEqual(str(self.app.manual_import_button["state"]), "disabled")
        self.app.set_busy(False)
        self.assertEqual(str(self.app.limit_box["state"]), "disabled")
        self.root.withdraw()

    def test_manual_skip_requires_selection_confirmation_and_only_dispatches_that_id(self):
        with patch("app.messages.showwarning"), patch.object(self.service, "skip_selected") as skip:
            self.app.skip_selected()
        skip.assert_not_called()
        self.app.trees["claim"].selection_set("claim-pending")
        self.app.show_selected("claim")
        self.assertEqual(str(self.app.skip_button["state"]), "normal")
        with patch("app.messages.askyesno", return_value=False), patch.object(self.service, "skip_selected") as skip:
            self.app.skip_selected()
        skip.assert_not_called()
        chosen = replace(self.app.roster.records[0], skipped=True)
        saved = SimpleNamespace(records=[chosen] + self.app.roster.records[1:], assert_unchanged=Mock())
        result = TaskResult(saved, {chosen.sa_id: "已手动跳过"}, "已手动跳过")
        with patch("app.messages.askyesno", return_value=True), patch.object(self.service, "skip_selected", return_value=result) as skip:
            self.app.skip_selected()
        self.assertEqual(skip.call_args.args[4], "claim-pending")
        self.assertEqual(self.app.status.get(), "已手动跳过")
        self.assertFalse(self.app.busy)
        self.assertNotIn(chosen.sa_id, self.app.trees["claim"].get_children())
        self.assertEqual(self.app.trees["claim"].selection(), ())
        self.assertEqual(str(self.app.skip_button["state"]), "disabled")

    def test_manual_skip_is_locked_while_running_and_backend_failure_stays_visible(self):
        self.app.trees["claim"].selection_set("claim-pending")
        self.app.set_busy(True)
        with patch.object(self.service, "skip_selected") as skip:
            self.app.skip_selected()
        skip.assert_not_called()
        self.app.set_busy(False)
        with patch("app.messages.askyesno", return_value=True), patch("app.messages.showwarning"), \
                patch.object(self.service, "skip_selected", side_effect=SafetyStop("原操作尚待核验")):
            self.app.skip_selected()
        self.assertIn("尚待核验", self.app.status.get())
        self.assertIn("claim-pending", self.app.trees["claim"].get_children())
        self.assertFalse(self.app.busy)

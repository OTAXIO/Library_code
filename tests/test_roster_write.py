"""Synthetic XLSX fixtures only; never change the user's list.xlsx."""
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from core import HEADERS, QUERY_HEADER, SafetyStop, file_hash, read_roster
from roster_write import mark_complete


def make_roster(path, flags=(None, 1, "1", 0, True)):
    book = Workbook()
    sheet = book.active
    sheet.title = "名单"
    sheet.append(["备注"] + list(HEADERS.values()) + [QUERY_HEADER, "备注"])
    for i, flag in enumerate(flags, 1):
        values = dict(owner="测试员", sa_id=f"demo-{i:03}", title=f"Synthetic {i}", doi="", wos="",
                      staff_id="00001", matches=1, item_ids="1234567890123456789",
                      mark="待处理", reason="作者不一致")
        sheet.append([flag] + [values[k] for k in HEADERS] + ["1", "网页备注保留"])
    sheet['A2'].fill = PatternFill('solid', fgColor='FFF4C2')
    sheet.freeze_panes = "C2"
    sheet.column_dimensions['A'].width = 12
    extra = book.create_sheet("保留页")
    extra['A1'] = "=1+1"
    extra['B1'] = "不要修改"
    book.save(path)
    book.close()


class RosterWriteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "list.xlsx"
        make_roster(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_only_numeric_one_counts(self):
        roster = read_roster(self.path)
        self.assertEqual([r.done for r in roster.records], [False, True, False, False, False])
        self.assertEqual(roster.completion_column, 1)

    def test_write_preserves_other_zip_parts_and_notes(self):
        roster = read_roster(self.path)
        before = self.path.read_bytes()
        with zipfile.ZipFile(self.path) as z:
            parts = {name: z.read(name) for name in z.namelist()}
        result = mark_complete(roster, roster.records[0])
        self.assertEqual(result.backup.read_bytes(), before)
        self.assertEqual(result.cell, "A2")
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(result.roster.records[0].key, roster.records[0].key)
        with zipfile.ZipFile(self.path) as z:
            changed = [name for name in z.namelist() if z.read(name) != parts[name]]
        self.assertEqual(changed, ["xl/worksheets/sheet1.xml"])
        book = load_workbook(self.path)
        try:
            self.assertEqual(book['名单']['A2'].value, 1)
            self.assertEqual(book['名单']['A2'].data_type, 'n')
            self.assertEqual(book['名单']['M2'].value, "网页备注保留")
            self.assertEqual(book['名单']['A2'].fill.fgColor.rgb, '00FFF4C2')
            self.assertEqual(book['名单'].freeze_panes, 'C2')
            self.assertEqual(book['保留页']['A1'].value, '=1+1')
        finally:
            book.close()

    def test_sequential_completion_updates_fingerprint(self):
        roster = read_roster(self.path)
        first = mark_complete(roster, roster.records[0])
        second = mark_complete(first.roster, first.roster.records[2])
        self.assertTrue(second.roster.records[0].done)
        self.assertTrue(second.roster.records[2].done)
        second.roster.assert_unchanged()

    def test_duplicate_completion_stops(self):
        roster = read_roster(self.path)
        with self.assertRaises(SafetyStop):
            mark_complete(roster, roster.records[1])

    def test_changed_file_stops_without_overwrite(self):
        roster = read_roster(self.path)
        make_roster(self.path, [0])
        current = file_hash(self.path)
        with self.assertRaises(SafetyStop):
            mark_complete(roster, roster.records[0])
        self.assertEqual(file_hash(self.path), current)

    def test_changed_during_prepare_stops(self):
        roster = read_roster(self.path)
        def changed(path):
            verified = read_roster(path)
            make_roster(self.path, [0])
            return verified
        with patch('roster_write.read_roster', side_effect=changed):
            with self.assertRaises(SafetyStop):
                mark_complete(roster, roster.records[0])
        self.assertEqual(len(read_roster(self.path).records), 1)

    def test_excel_lock_stops(self):
        roster = read_roster(self.path)
        self.path.with_name('~$list.xlsx').touch()
        with self.assertRaisesRegex(SafetyStop, "Excel"):
            mark_complete(roster, roster.records[0])
        roster.assert_unchanged()

    def test_assistant_lock_stops(self):
        roster = read_roster(self.path)
        self.path.with_name('.list.xlsx.assistant.lock').touch()
        with self.assertRaises(SafetyStop):
            mark_complete(roster, roster.records[0])
        roster.assert_unchanged()

    def test_replace_failure_leaves_original_and_cleans_temp(self):
        roster = read_roster(self.path)
        with patch('roster_write.os.replace', side_effect=PermissionError):
            with self.assertRaises(SafetyStop):
                mark_complete(roster, roster.records[0])
        roster.assert_unchanged()
        self.assertEqual(list(self.path.parent.glob('.list-write-*')), [])
        self.assertFalse(self.path.with_name('.list.xlsx.assistant.lock').exists())

    def test_backup_failure_leaves_original(self):
        roster = read_roster(self.path)
        blocked = self.path.parent / 'blocked'
        blocked.write_text('fixture', encoding='utf-8')
        with self.assertRaises(OSError):
            mark_complete(roster, roster.records[0], blocked)
        roster.assert_unchanged()

    def test_missing_flag_cell_is_inserted_in_order(self):
        make_roster(self.path, [None, 1, "1", None])
        roster = read_roster(self.path)
        result = mark_complete(roster, roster.records[3])
        self.assertTrue(result.roster.records[3].done)

    def test_existing_note_is_backed_up(self):
        make_roster(self.path, ["未核对原文"])
        roster = read_roster(self.path)
        result = mark_complete(roster, roster.records[0])
        self.assertEqual(result.previous, "未核对原文")
        self.assertEqual(read_roster(result.backup).records[0].remark, "未核对原文")

    def test_missing_or_ambiguous_header_stops(self):
        for header in ('其他', '备注'):
            make_roster(self.path)
            book = load_workbook(self.path)
            book.active['A1'] = header
            book.active['B1'] = '负责人' if header == '其他' else '备注'
            if header == '其他':
                book.active['M1'] = '其他备注'
            book.save(self.path)
            book.close()
            with self.assertRaises(SafetyStop):
                read_roster(self.path)

    def test_formula_flag_stops(self):
        make_roster(self.path, ['=1'])
        with self.assertRaises(SafetyStop):
            read_roster(self.path)

    def test_protected_sheet_stops(self):
        book = load_workbook(self.path)
        book.active.protection.sheet = True
        book.save(self.path)
        book.close()
        roster = read_roster(self.path)
        with self.assertRaises(SafetyStop):
            mark_complete(roster, roster.records[0])
        roster.assert_unchanged()

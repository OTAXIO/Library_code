"""Cleanup tests use only synthetic workbooks and never reset real journals."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

from core import SafetyStop, read_roster
from maintenance.cleanup_import_notes import apply_candidate, cleanup_plan
from tests.test_roster_write import make_roster


class CleanupImportNotesTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / 'list.xlsx'
        make_roster(self.path, [None] * 7)
        book = load_workbook(self.path)
        sheet = book['名单']
        sheet['N1'] = '是否识别'
        cases = [
            ('谭勋策', 'WOS 下载未完成：[WOS 已暂停] WOS 检索结果超时', 2),
            ('谭勋策', 'wos未查询到', 2),
            ('谭勋策', '已补录/关联，第一单位待核验，暂不处理', 2),
            ('谭勋策', '非交大', 1),
            ('其他负责人', 'WOS 下载未完成：合成异常', 2),
            ('谭勋策', '人工暂存说明', 2),
            ('谭勋策', 'WOS 下载未完成：[WOS 已暂停] WOS 未找到记录；不自动结案', 2),
        ]
        for row, (owner, note, flag) in enumerate(cases, 2):
            sheet.cell(row, 1, note)
            sheet.cell(row, 2, owner)
            sheet.cell(row, 8, 0)  # matches column in the original fixture
            sheet.cell(row, 9).value = None
            sheet.cell(row, 14, flag)
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)

    def tearDown(self):
        self.folder.cleanup()

    def candidate(self):
        path = self.path.with_name('authored.xlsx')
        book = load_workbook(self.path)
        sheet = book['名单']
        for change in cleanup_plan(self.roster):
            sheet.cell(change['row'], 1).value = change['after'] or None
            if change['clear_flag']:
                sheet.cell(change['row'], 14).value = None
        book.save(path)
        book.close()
        return path

    def test_plan_excludes_complete_other_owner_unknown_and_real_zero(self):
        changes = cleanup_plan(self.roster)
        self.assertEqual([(c['row'], c['after'], c['clear_flag']) for c in changes],
            [(2, '', True), (3, 'wos未收录', False), (4, '第一单位待核验', False),
             (8, 'wos未收录', False)])

    def test_apply_changes_only_approved_cells_and_backs_up_full_old_notes(self):
        result = apply_candidate(self.roster, self.candidate(), self.roster.sha256)
        self.assertFalse(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark, '')
        self.assertTrue(result.roster.records[1].skipped)
        self.assertEqual(result.roster.records[1].remark, 'wos未收录')
        self.assertEqual(result.roster.records[3:6], self.roster.records[3:6])
        self.assertEqual(read_roster(result.backup).records, self.roster.records)
        self.assertEqual([r.key for r in result.roster.records], [r.key for r in self.roster.records])

    def test_conflicting_candidate_and_changed_source_never_write(self):
        candidate = self.candidate()
        book = load_workbook(candidate)
        book['名单']['N5'] = None  # completed row is not an allowed cleanup target
        book.save(candidate)
        book.close()
        with self.assertRaises(SafetyStop):
            apply_candidate(self.roster, candidate, self.roster.sha256)
        with self.assertRaises(SafetyStop):
            apply_candidate(self.roster, candidate, '0' * 64)
        self.roster.assert_unchanged()

    def test_authoring_export_metadata_is_never_copied_into_original(self):
        candidate = self.candidate()
        book = load_workbook(candidate)
        book['名单']['I2'] = '323'  # emulate an empty-shared-string export bug
        book.save(candidate)
        book.close()
        result = apply_candidate(self.roster, candidate, self.roster.sha256)
        self.assertEqual(result.roster.records[0].item_ids, '')
        self.assertEqual(result.roster.records[0].key, self.roster.records[0].key)


if __name__ == '__main__':
    unittest.main()

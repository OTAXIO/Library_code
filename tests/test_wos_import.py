import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from automation import ImportStore, WOSFlow
from core import SafetyStop
from tests.test_automation import record as base_record, sample, result as sa_result
from wos_import import ImportPlan, build_plan, run_import_plan


def record(**kwargs):
    return base_record(owner="谭勋策", **kwargs)


class Roster:
    def __init__(self, records):
        self.records = records
        self.sha256 = "synthetic-fingerprint"
        self.changed = False

    def assert_unchanged(self):
        if self.changed:
            raise SafetyStop("名单改变")


class Bridge:
    def __init__(self, records):
        self.records = {r.sa_id: r for r in records}
        self.calls = []
        self.batches = {}
        self.fail = None
        self.processed = set()

    def call(self, action, payload, timeout=75):
        self.calls.append((action, payload["sa_id"]))
        if action == self.fail:
            raise SafetyStop("synthetic timeout")
        ident = payload["sa_id"]
        if action in ("status", "search"):
            r = self.records[ident]
            return sa_result(saLzkId=ident, titleValue=r.title, doiValue=r.doi, wosValue=r.wos,
                             markStatus="已处理" if ident in self.processed else "待处理")
        if action == "import_scan":
            return {"batches": [self.batches[ident]] if ident in self.batches else []}
        if action == "import_upload":
            return {"uploaded": True, "sha256": payload["candidate"]["sha256"]}
        if action == "import_submit":
            self.batches[ident] = dict(id="batch-" + ident, batchNumber="synthetic-" + ident, status=1, modelId="m")
            return {"submitted": True}
        if action == "import_check":
            if ident not in self.batches:
                raise SafetyStop("批次未完成")
            return {"verified": True, "batch": dict(self.batches[ident])}
        if action == "import_push":
            self.batches[ident]["status"] = 2
            return {"submitted": True}
        raise AssertionError("Unexpected command: " + action)


class ImportQueueTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inbox = self.root / "inbox"
        self.inbox.mkdir()
        self.file = self.inbox / "one.txt"
        self.file.write_bytes(sample())
        self.store = ImportStore(self.root / "imports")
        self.roster = Roster([record()])
        self.bridge = Bridge(self.roster.records)

    def plan(self, **kwargs):
        return build_plan(self.roster, "谭勋策", kwargs.get("limit", 5), self.inbox, self.store,
                          scope=kwargs.get("scope", "pending"), download_store=kwargs.get("download_store"))

    def run_plan(self, plan=None, **kwargs):
        return run_import_plan(plan or self.plan(), self.roster, self.bridge, self.store, reviewed=True, **kwargs)

    def actions(self):
        return [a for a, _ in self.bridge.calls]

    def test_plan_is_local_only_and_does_not_stage_import(self):
        plan = self.plan()
        self.assertEqual([item.status for item in plan.items], ["ready"])
        self.assertIsNone(self.store.get(record()))
        self.assertEqual(self.bridge.calls, [])

    def test_stale_download_does_not_hide_independently_valid_inbox_file(self):
        from automation import parse_wos
        downloads = ImportStore(self.root / 'downloads')
        old = record(title='Previous title')
        downloads.archive(sample())
        downloads.save(old, {'phase': 'downloaded', 'candidate': parse_wos(sample()), 'identity_confirmed': True})
        plan = build_plan(self.roster, '谭勋策', 5, self.inbox, self.store, download_store=downloads)
        self.assertEqual(plan.items[0].status, 'ready')
        self.assertIn('下载存档', plan.file_errors[0])
        self.file.unlink()
        plan = build_plan(self.roster, '谭勋策', 5, self.inbox, self.store, download_store=downloads)
        self.assertEqual(plan.items, ())
        self.assertIn('存档不可用', plan.excluded[0].message)

    def test_owner_scope_done_skipped_and_matches(self):
        self.roster.records += [replace(record(sa_id="other"), owner="另一位"), record(sa_id="done", done=True),
                                record(sa_id="skip", skipped=True), record(sa_id="matched", matches=1)]
        self.assertEqual([i.record.sa_id for i in self.plan().items], ["demo-001"])
        for owner in ("", "另一位"):
            with self.assertRaises(SafetyStop):
                build_plan(self.roster, owner, 5, self.inbox, self.store)
        for limit in (0, 101, True, "5"):
            with self.assertRaises(SafetyStop):
                self.plan(limit=limit)

    def test_requires_explicit_review_and_unchanged_roster(self):
        plan = self.plan()
        with self.assertRaises(SafetyStop):
            run_import_plan(plan, self.roster, self.bridge, self.store)
        self.roster.sha256 = "changed"
        with self.assertRaises(SafetyStop):
            self.run_plan(plan)
        self.assertEqual(self.bridge.calls, [])

    def test_explicit_skipped_scope_imports_without_clearing_skip_or_approving(self):
        self.roster.records = [record(skipped=True, remark="2")]
        self.bridge = Bridge(self.roster.records)
        self.assertEqual(self.plan().items, ())
        plan = self.plan(scope="skipped")
        self.assertEqual(plan.scope, "skipped")
        self.assertEqual(len(plan.items), 1)
        result = self.run_plan(plan)
        self.assertEqual(result.outcomes[0]["status"], "pushed")
        self.assertTrue(result.roster.records[0].skipped)
        self.assertFalse(result.roster.records[0].done)
        self.assertNotIn("complete", self.actions())

    def test_scope_mismatch_or_unknown_scope_rejected_before_browser(self):
        self.roster.records = [record(skipped=True)]
        for scope in ("all", "", None):
            with self.assertRaises(SafetyStop):
                self.plan(scope=scope)
        plan = self.plan(scope="skipped")
        for scope in ("pending", "all"):
            with self.assertRaises(SafetyStop):
                self.run_plan(replace(plan, scope=scope))
        self.assertEqual(self.bridge.calls, [])

    def test_limit_applies_to_usable_exports_not_first_missing_roster_rows(self):
        missing = [record(sa_id=f"missing-{i}", title=f"Missing {i}", doi=f"10.1234/missing{i}")
                   for i in range(7)]
        self.roster.records = missing + [record()]
        plan = self.plan(limit=1)
        self.assertEqual([item.record.sa_id for item in plan.items], ["demo-001"])
        self.assertEqual(len(plan.excluded), 7)
        self.assertTrue(all(item.status == "deferred" for item in plan.excluded))

    def test_completed_imports_do_not_starve_later_files_and_extra_ready_is_queued(self):
        self.run_plan()
        for n in (2, 3):
            self.roster.records.append(record(sa_id=f"later-{n}", title=f"Later {n}", doi=f"10.1234/later{n}"))
            (self.inbox / f"later{n}.txt").write_bytes(sample(TI=f"Later {n}", DI=f"10.1234/later{n}", UT=f"WOS:00012345678902{n}"))
        plan = self.plan(limit=1)
        self.assertEqual([item.record.sa_id for item in plan.items], ["later-2"])
        self.assertEqual([item.status for item in plan.excluded], ["pushed", "queued"])

    def test_resume_has_priority_over_new_upload_under_execution_limit(self):
        resumed = record(sa_id="resume", title="Resume", doi="10.1234/resume")
        self.roster.records.append(resumed)
        local = self.inbox / "resume.txt"
        local.write_bytes(sample(TI="Resume", DI="10.1234/resume", UT="WOS:000123456789013"))
        state = WOSFlow(self.bridge, self.store).prepare_file(resumed, local)
        self.store.save(resumed, {**state, "phase": "import_intent"})
        plan = self.plan(limit=1)
        self.assertEqual([(item.record.sa_id, item.status) for item in plan.items], [("resume", "resume")])
        self.assertEqual(plan.excluded[0].status, "queued")

    def test_weak_title_only_download_is_visible_for_review_without_import(self):
        from automation import parse_wos
        self.file.unlink()
        r = record(doi="")
        self.roster.records = [r]
        downloads = ImportStore(self.root / "downloads")
        raw = sample()
        sha = downloads.archive(raw)
        downloads.save(r, {"phase": "downloaded", "candidate": parse_wos(raw), "identity_confirmed": False})
        plan = self.plan(download_store=downloads)
        self.assertEqual(plan.items, ())
        item = plan.excluded[0]
        self.assertEqual(item.status, "deferred")
        self.assertEqual(item.sha256, sha)
        self.assertEqual(Path(item.path).read_bytes(), raw)
        self.assertEqual(item.candidate["title"], "Synthetic paper")
        self.assertEqual(self.bridge.calls, [])

    def test_download_archive_never_adopts_another_sa_identity_or_ignores_conflict(self):
        from automation import parse_wos
        self.file.unlink()
        downloads = ImportStore(self.root / "downloads")
        raw = sample()
        downloads.archive(raw)
        downloads.save(record(sa_id="other"), {"phase": "downloaded", "candidate": parse_wos(raw), "identity_confirmed": True})
        self.assertFalse(self.plan(download_store=downloads).excluded[0].path)
        conflicting = sample(DI="10.1234/wrong")
        downloads.archive(conflicting)
        downloads.save(record(), {"phase": "downloaded", "candidate": parse_wos(conflicting), "identity_confirmed": True})
        plan = self.plan(download_store=downloads)
        self.assertEqual(plan.items, ())
        self.assertFalse(plan.excluded[0].path)
        self.assertIn("冲突", plan.excluded[0].message)

    def test_foreign_record_plan_rejected_before_browser(self):
        plan = self.plan()
        item = replace(plan.items[0], record=replace(record(), owner="另一位"))
        with self.assertRaises(SafetyStop):
            self.run_plan(replace(plan, items=(item,)))
        self.assertEqual(self.bridge.calls, [])

    def test_local_file_imports_once_without_wos_or_approval_commands(self):
        result = self.run_plan()
        self.assertEqual(result.outcomes[0]["status"], "pushed")
        for action in ("import_upload", "import_submit", "import_push"):
            self.assertEqual(self.actions().count(action), 1)
        self.assertFalse(any(a.startswith("wos_") for a in self.actions()))
        self.assertEqual(set(self.actions()) - {"search", "import_scan", "import_upload", "import_submit", "import_push", "import_check"}, set())
        self.assertFalse(result.roster.records[0].done)
        second = self.run_plan()
        self.assertEqual(second.outcomes, [])
        self.assertEqual(self.plan().excluded[0].status, "pushed")
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_audit_uses_existing_operation_log_schema(self):
        from operation_log import OperationLog
        log = OperationLog(self.root / "log.txt")
        self.run_plan(audit=log.record)
        events = log.read()
        self.assertIn("import_submit", [event["action"] for event in events])
        self.assertEqual(events[-1]["action"], "WOS 导入队列")
        self.assertEqual(events[-1]["result"], "已执行")

    def test_backend_done_syncs_before_any_import(self):
        self.bridge.processed.add("demo-001")
        updated = Roster([record(done=True)])
        from types import SimpleNamespace
        with patch("wos_import.reconcile_processed", return_value=SimpleNamespace(roster=updated)) as sync:
            result = self.run_plan()
        self.assertEqual(self.actions(), ["search"])
        self.assertEqual(result.outcomes[0]["status"], "synced")
        self.assertIs(result.roster, updated)
        self.assertEqual(sync.call_args.args[2]["markStatus"], "已处理")

    def test_missing_file_is_shown_without_consuming_queue_or_visiting_browser(self):
        self.file.unlink()
        self.bridge.processed.add("demo-001")
        plan = self.plan()
        self.assertEqual(plan.items, ())
        self.assertEqual(plan.excluded[0].status, "deferred")
        self.assertEqual(self.run_plan(plan).outcomes, [])
        self.assertEqual(self.actions(), [])

    def test_no_affiliation_or_weak_identity_never_uploads(self):
        for raw, r in ((sample(C1="Other University"), record()), (sample(), record(doi="")),
                       (sample(DI="10.1234/conflict"), record()), (sample(TI="Different title"), record())):
            with self.subTest(record=r, content=raw[:30]):
                self.file.write_bytes(raw)
                self.roster = Roster([r])
                self.bridge = Bridge(self.roster.records)
                self.assertEqual(self.plan().excluded[0].status, "deferred")
                result = self.run_plan()
                self.assertEqual(result.outcomes, [])
                self.assertNotIn("import_upload", self.actions())

    def test_changed_file_after_plan_does_not_upload(self):
        plan = self.plan()
        self.file.write_bytes(sample(AF="Another author"))
        self.assertEqual(self.run_plan(plan).outcomes[0]["status"], "deferred")
        self.assertNotIn("import_upload", self.actions())

    def test_distinct_exports_ambiguous_identical_bytes_deduplicated(self):
        other = self.inbox / "copy.txt"
        other.write_bytes(self.file.read_bytes())
        self.assertEqual(self.plan().items[0].status, "ready")
        other.write_bytes(sample(AF="Other author"))
        self.assertEqual(self.plan().excluded[0].status, "deferred")

    def test_invalid_and_multirecord_files_do_not_replace_valid_file(self):
        (self.inbox / "invalid.txt").write_bytes(b"<html>login</html>")
        raw = sample()
        (self.inbox / "many.txt").write_bytes(raw + raw.split(b"\r\n")[1] + b"\r\n")
        plan = self.plan()
        self.assertEqual(len(plan.file_errors), 2)
        self.assertEqual(plan.items[0].status, "ready")

    def test_duplicate_paper_different_sa_id_never_reimports(self):
        self.roster.records.append(record(sa_id="second-id", row=3))
        self.bridge = Bridge(self.roster.records)
        self.assertEqual([i.status for i in self.plan().items], ["ready"])
        self.assertEqual([i.status for i in self.plan().excluded], ["deferred"])
        self.run_plan()
        next_plan = self.plan()
        self.assertEqual(next_plan.items, ())
        self.assertEqual([i.status for i in next_plan.excluded], ["pushed", "deferred"])
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_import_timeout_resume_only_reads_existing_submission(self):
        self.bridge.fail = "import_check"
        result = self.run_plan()
        self.assertTrue(result.halted)
        self.assertEqual(self.store.get(record())["phase"], "import_intent")
        self.bridge.fail = None
        self.assertEqual(self.plan().items[0].status, "resume")
        resumed = self.run_plan()
        self.assertEqual(resumed.outcomes[0]["status"], "pushed")
        self.assertEqual(self.actions().count("import_upload"), 1)
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_uncertain_push_is_not_sent_twice(self):
        self.bridge.fail = "import_push"
        self.assertTrue(self.run_plan().halted)
        self.bridge.fail = None
        self.assertTrue(self.run_plan().halted)
        self.assertEqual(self.actions().count("import_push"), 1)

    def test_uncertain_upload_is_deferred_without_second_upload(self):
        self.bridge.fail = "import_upload"
        self.assertTrue(self.run_plan().halted)
        self.bridge.fail = None
        self.assertEqual(self.plan().excluded[0].status, "deferred")
        self.run_plan()
        self.assertEqual(self.actions().count("import_upload"), 1)

    def test_disconnection_stops_instead_of_visiting_next_record(self):
        self.roster.records.append(record(sa_id="second-id", row=3))
        self.bridge.fail = "search"
        result = self.run_plan()
        self.assertTrue(result.halted)
        self.assertEqual(len(result.outcomes), 1)
        self.assertEqual(len(self.bridge.calls), 1)

    def test_stop_and_changed_roster_before_run(self):
        result = self.run_plan(stop=lambda: True)
        self.assertTrue(result.cancelled)
        self.assertEqual(self.bridge.calls, [])
        plan = self.plan()
        self.roster.changed = True
        with self.assertRaises(SafetyStop):
            self.run_plan(plan)

    def test_prepare_file_stages_weak_evidence_without_confirming(self):
        flow = WOSFlow(self.bridge, self.store)
        state = flow.prepare_file(record(doi=""), self.file)
        self.assertEqual(state["phase"], "exported")
        self.assertFalse(state["identity_confirmed"])
        self.assertEqual(self.bridge.calls, [])

    def test_existing_archive_cannot_be_silently_replaced(self):
        flow = WOSFlow(self.bridge, self.store)
        flow.prepare_file(record(), self.file)
        self.file.write_bytes(sample(AF="Changed author"))
        with self.assertRaisesRegex(SafetyStop, "已有不同"):
            flow.prepare_file(record(), self.file)

    def test_same_paper_guard_also_protects_existing_single_record_flow(self):
        self.run_plan()
        r = record(sa_id="another-id", row=3)
        flow = WOSFlow(self.bridge, self.store)
        flow.prepare_file(r, self.file)
        with self.assertRaisesRegex(SafetyStop, "同一论文已有导入"):
            flow.proceed(r)
        self.assertEqual(self.actions().count("import_submit"), 1)

    def test_real_synthetic_workbook_sync_keeps_next_import_pending(self):
        from openpyxl import load_workbook
        from core import HEADERS, read_roster
        from tests.test_roster_write import make_roster
        path = self.root / "list.xlsx"
        make_roster(path, flags=(None, None, None))
        book = load_workbook(path)
        sheet = book.active
        columns = {key: 2 + list(HEADERS).index(key) for key in HEADERS}
        for n in (2, 3):
            for key, value in {"owner": "谭勋策", "matches": 0, "item_ids": "", "reason": ""}.items():
                sheet.cell(n, columns[key], value)
        sheet.cell(3, columns["title"], "Synthetic paper")
        sheet.cell(3, columns["doi"], "10.1234/test")
        sheet.cell(2, columns["doi"], "10.1234/first")
        (self.inbox / "first.txt").write_bytes(sample(TI="Synthetic 1", DI="10.1234/first", UT="WOS:000123456789013"))
        book.save(path)
        book.close()
        self.roster = read_roster(path)
        self.bridge = Bridge(self.roster.records)
        self.bridge.processed.add("demo-001")
        plan = self.plan()
        self.assertEqual(len(plan.items), 2)
        result = self.run_plan(plan)
        self.assertEqual([o["status"] for o in result.outcomes], ["synced", "pushed"])
        verified = read_roster(path)
        self.assertTrue(verified.records[0].done)
        self.assertFalse(verified.records[1].done)
        self.assertEqual(verified.records[2], self.roster.records[2])

    def test_import_failure_records_two_and_resume_does_not_repeat_submission(self):
        from openpyxl import load_workbook
        from core import HEADERS, read_roster
        from tests.test_roster_write import make_roster
        path = self.root / 'list.xlsx'
        make_roster(path, flags=(None, None, None))
        book = load_workbook(path)
        sheet = book.active
        columns = {key: 2 + list(HEADERS).index(key) for key in HEADERS}
        sheet.cell(1, 14, '是否识别')
        for row, title, identifier in ((2, 'Synthetic paper', '10.1234/test'),
                                       (3, 'Next paper', '10.1234/next')):
            for key, value in {'owner': '谭勋策', 'matches': 0, 'item_ids': '', 'reason': '',
                               'title': title, 'doi': identifier}.items():
                sheet.cell(row, columns[key], value)
        book.save(path)
        book.close()
        (self.inbox / 'next.txt').write_bytes(sample(TI='Next paper', DI='10.1234/next', UT='WOS:000123456789013'))
        self.roster = read_roster(path)
        untouched = self.roster.records[1:]
        self.bridge = Bridge(self.roster.records)
        self.bridge.fail = 'import_check'
        result = self.run_plan()
        self.assertTrue(result.halted)
        self.assertEqual(len(result.outcomes), 1)
        self.assertTrue(result.roster.records[0].skipped)
        self.assertIn('synthetic timeout', result.roster.records[0].remark)
        self.assertEqual(result.roster.records[1:], untouched)
        self.assertEqual(self.actions().count('import_submit'), 1)
        self.roster = result.roster
        self.bridge.fail = None
        resumed = self.run_plan(self.plan(scope='skipped'))
        self.assertEqual(resumed.outcomes[0]['status'], 'pushed')
        self.assertTrue(resumed.roster.records[0].skipped)
        self.assertFalse(resumed.roster.records[0].done)
        self.assertEqual(self.actions().count('import_submit'), 1)

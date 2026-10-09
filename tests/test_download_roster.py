"""WOS download outcomes use synthetic workbooks only."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from core import SafetyStop, file_hash, read_roster
from tests.test_roster_write import make_roster
from wos_batch import persist_download_outcomes


class DownloadRosterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'list.xlsx'
        make_roster(self.path, [None] * 6)
        book = load_workbook(self.path)
        sheet = book['名单']
        sheet.cell(1, 14, '数据来源')
        sheet.cell(1, 15, '是否识别')
        for row in range(2, 8):
            sheet.cell(row, 2, '谭勋策')
            sheet.cell(row, 8, 0)
            sheet.cell(row, 9).value = None
        # A duplicate paper, a different owner and a conflicting WOS ID must
        # not borrow an outcome indiscriminately from the selected row.
        sheet.cell(5, 4, 'Synthetic 1')
        sheet.cell(6, 4, 'Synthetic 1')
        sheet.cell(6, 2, 'Other')
        sheet.cell(7, 4, 'Synthetic 1')
        sheet.cell(7, 6, 'WOS:000123456789012')
        book.save(self.path)
        book.close()
        self.roster = read_roster(self.path)
        self.targets = self.roster.records[:3]

    def result(self):
        a, b, _c = self.targets
        return {'attempted': 2,
                'exported': [{'sa_id': a.sa_id, 'row': a.row}],
                'unconfirmed': [],
                'failed': {b.sa_id: {'row': b.row, 'error': 'WOS 未找到记录'}}}

    def test_success_is_provenance_only_failed_row_is_two_and_tail_unchanged(self):
        update = persist_download_outcomes(self.roster, self.targets, self.result())
        a, b, c, duplicate, other, conflict = update.roster.records
        self.assertEqual(update.source_count, 2)
        self.assertEqual(update.skipped_count, 1)
        self.assertEqual(a.source, 'WOS')
        self.assertFalse(a.done)
        self.assertFalse(a.skipped)
        self.assertTrue(b.skipped)
        self.assertIn('未找到记录', b.remark)
        self.assertEqual(c, self.roster.records[2])
        self.assertEqual(duplicate.source, 'WOS')
        self.assertEqual(other, self.roster.records[4])
        self.assertEqual(conflict, self.roster.records[5])
        book = load_workbook(self.path)
        self.assertEqual(book['名单']['A3'].value, b.remark)
        self.assertEqual(book['名单']['O3'].value, 2)
        self.assertEqual(book['名单']['M3'].value, '网页备注保留')
        book.close()

    def test_weak_identity_preserves_archive_but_records_manual_review(self):
        a = self.targets[0]
        result = {'attempted': 1, 'exported': [], 'failed': {},
                  'unconfirmed': [{'sa_id': a.sa_id, 'row': a.row, 'archive': 'kept.txt'}]}
        update = persist_download_outcomes(self.roster, self.targets, result)
        self.assertTrue(update.roster.records[0].skipped)
        self.assertIn('身份待核验', update.roster.records[0].remark)
        self.assertEqual(update.roster.records[0].source, '')
        self.assertFalse(update.roster.records[0].done)

    def test_disconnect_only_records_the_attempted_row(self):
        a = self.targets[0]
        result = {'attempted': 1, 'exported': [], 'unconfirmed': [],
                  'failed': {a.sa_id: {'row': a.row, 'error': '浏览器通信已断开；结果不明'}},
                  'remaining': 2, 'disconnected': True}
        update = persist_download_outcomes(self.roster, self.targets, result)
        self.assertTrue(update.roster.records[0].skipped)
        self.assertIn('结果不明', update.roster.records[0].remark)
        self.assertEqual(update.roster.records[1:], self.roster.records[1:])

    def test_unattempted_or_duplicate_outcomes_cannot_write(self):
        before = file_hash(self.path)
        result = self.result()
        result['attempted'] = 1
        with self.assertRaises(SafetyStop):
            persist_download_outcomes(self.roster, self.targets, result)
        result = self.result()
        result['failed'][self.targets[0].sa_id] = {'row': 2, 'error': 'contradictory'}
        with self.assertRaises(SafetyStop):
            persist_download_outcomes(self.roster, self.targets, result)
        self.assertEqual(file_hash(self.path), before)

    def test_second_write_failure_keeps_the_first_transaction_and_its_roster(self):
        with patch('roster_write.mark_skipped_many', side_effect=SafetyStop('文件被占用')):
            update = persist_download_outcomes(self.roster, self.targets, self.result())
        self.assertEqual(update.roster.sha256, file_hash(self.path))
        self.assertEqual(update.source_count, 2)
        self.assertEqual(update.skipped_count, 0)
        self.assertIn('文件被占用', update.workflow_error)
        self.assertFalse(update.roster.records[1].skipped)

    def test_retry_can_replace_reason_without_clearing_the_skip_flag(self):
        b = self.targets[1]
        result = {'attempted': 1, 'exported': [], 'unconfirmed': [],
                  'failed': {b.sa_id: {'row': b.row, 'error': '第一次失败'}}}
        first = persist_download_outcomes(self.roster, [b], result)
        latest = first.roster.records[1]
        result['failed'][b.sa_id]['error'] = '重新核验仍有多个结果'
        second = persist_download_outcomes(first.roster, [latest], result)
        self.assertTrue(second.roster.records[1].skipped)
        self.assertIn('重新核验', second.roster.records[1].remark)
        self.assertNotIn('第一次失败', second.roster.records[1].remark)


if __name__ == '__main__':
    unittest.main()

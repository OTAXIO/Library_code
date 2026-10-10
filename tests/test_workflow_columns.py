"""Synthetic fixtures for separate status/remarks; no user roster is mutated."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from core import SafetyStop, read_roster
from roster_write import mark_complete, mark_skipped_many, clear_skipped_many, reconcile_processed, migrate_status_column
from roster_write import cleanup_skipped_notes
from tests.test_roster_write import make_roster


class WorkflowColumnsTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / 'list.xlsx'
        make_roster(self.path, ['原说明', '非工作项', None])
        book = load_workbook(self.path)
        sheet = book['名单']
        sheet.cell(1,14,'是否识别')
        sheet.cell(2,2,'谭勋策')
        sheet.cell(3,14,1)
        sheet.cell(4,14,2)
        book.save(self.path)
        book.close()

    def tearDown(self):
        self.folder.cleanup()

    def test_new_column_is_the_only_status_source(self):
        roster = read_roster(self.path)
        self.assertTrue(roster.status_separate)
        self.assertEqual(roster.completion_column,14)
        self.assertEqual(roster.remark_column,1)
        self.assertEqual([r.remark for r in roster.records],['原说明','非工作项',''])
        self.assertEqual([r.done for r in roster.records],[False,True,False])
        self.assertEqual([r.skipped for r in roster.records],[False,False,True])

    def test_note_and_status_are_saved_atomically(self):
        roster=read_roster(self.path)
        result=mark_complete(roster,roster.records[0],note='已认领')
        self.assertEqual(result.cell,'N2')
        self.assertEqual(result.roster.records[0].remark,'已认领')
        self.assertTrue(result.roster.records[0].done)
        self.assertEqual(result.roster.records[1:],roster.records[1:])
        book=load_workbook(self.path)
        self.assertEqual(book['名单']['M2'].value,'网页备注保留')
        self.assertEqual(book['名单']['A2'].value,'已认领')
        book.close()

    def test_skip_reason_and_retry_do_not_clear_notes(self):
        roster=read_roster(self.path)
        result=mark_skipped_many(roster,[roster.records[0]],reasons={roster.records[0].sa_id:'无法确认文献归属'})
        self.assertTrue(result.roster.records[0].skipped)
        self.assertEqual(result.roster.records[0].remark,'无法确认文献归属')
        retried=clear_skipped_many(result.roster,[result.roster.records[0]])
        self.assertFalse(retried.roster.records[0].skipped)
        self.assertEqual(retried.roster.records[0].remark,'无法确认文献归属')

    def test_existing_skip_can_update_its_reason(self):
        roster=read_roster(self.path)
        item=roster.records[2]
        result=mark_skipped_many(roster,[item],reasons={item.sa_id:'重新核验仍无匹配'})
        self.assertTrue(result.roster.records[2].skipped)
        self.assertEqual(result.roster.records[2].remark,'重新核验仍无匹配')

    def test_explicit_note_cleanup_clears_only_the_skipped_row(self):
        import zipfile
        roster = read_roster(self.path)
        result = mark_skipped_many(roster, [roster.records[0]],
            reasons={roster.records[0].sa_id: 'WOS 下载未完成：合成页面问题'})
        before = result.roster
        with zipfile.ZipFile(self.path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        cleared = clear_skipped_many(before, [before.records[0]], clear_notes=True)
        self.assertFalse(cleared.roster.records[0].skipped)
        self.assertEqual(cleared.roster.records[0].remark, '')
        self.assertEqual(cleared.roster.records[1:], before.records[1:])
        self.assertEqual(read_roster(cleared.backup).records, before.records)
        with zipfile.ZipFile(self.path) as archive:
            self.assertEqual([name for name in parts if archive.read(name) != parts[name]],
                             ['xl/worksheets/sheet1.xml'])

    def test_combined_cleanup_retains_completed_and_other_owner_rows(self):
        from dataclasses import replace
        roster = read_roster(self.path)
        result = mark_skipped_many(roster, [roster.records[0]],
            reasons={roster.records[0].sa_id: '旧错误'})
        # Row 3 is completed and row 4 belongs to another owner.
        before = result.roster
        cleared = cleanup_skipped_notes(before, [before.records[0]], {})
        self.assertEqual(cleared.roster.records,
            [replace(before.records[0], skipped=False, remark='')] + before.records[1:])
        for record in cleared.roster.records[1:]:
            with self.assertRaises(SafetyStop):
                cleanup_skipped_notes(cleared.roster, [record], {})

    def test_cleanup_short_note_keeps_numeric_two(self):
        roster = read_roster(self.path)
        result = mark_skipped_many(roster, [roster.records[0]],
            reasons={roster.records[0].sa_id: '已补录/关联，第一单位待核验，暂不处理'})
        before = result.roster
        simplified = cleanup_skipped_notes(before, [], {before.records[0].sa_id: '第一单位待核验'})
        self.assertTrue(simplified.roster.records[0].skipped)
        self.assertEqual(simplified.roster.records[0].remark, '第一单位待核验')
        self.assertEqual(simplified.roster.records[1:], before.records[1:])

    def test_state_requires_an_explanation(self):
        roster=read_roster(self.path)
        with self.assertRaises(SafetyStop):
            mark_complete(roster,roster.records[0])
        with self.assertRaises(SafetyStop):
            mark_skipped_many(roster,[roster.records[0]])

    def test_backend_sync_uses_real_note(self):
        roster=read_roster(self.path)
        item=roster.records[0]
        result=reconcile_processed(roster,item,{'saLzkId':item.sa_id,'markStatus':'已处理','remark':'已认领'})
        self.assertEqual(result.roster.records[0].remark,'已认领')

    def test_bad_status_cannot_be_read_as_pending(self):
        book=load_workbook(self.path)
        book['名单']['N2']='1'
        book.save(self.path)
        book.close()
        with self.assertRaises(SafetyStop):
            read_roster(self.path)

    def test_migration_preserves_history_and_unrelated_package_parts(self):
        import zipfile
        make_roster(self.path,[1,2,'2',True,'说明'])
        before=read_roster(self.path)
        with zipfile.ZipFile(self.path) as z:
            parts={name:z.read(name) for name in z.namelist()}
        result=migrate_status_column(before)
        self.assertTrue(result.status_separate)
        self.assertEqual(result.completion_column,14)
        self.assertEqual([r.done for r in result.records],[r.done for r in before.records])
        self.assertEqual([r.skipped for r in result.records],[r.skipped for r in before.records])
        self.assertEqual([r.remark for r in result.records],['','','2','True','说明'])
        with zipfile.ZipFile(self.path) as z:
            self.assertEqual([name for name in parts if z.read(name)!=parts[name]],['xl/worksheets/sheet1.xml'])
        self.assertEqual(migrate_status_column(result).sha256,result.sha256)

if __name__=='__main__':
    unittest.main()

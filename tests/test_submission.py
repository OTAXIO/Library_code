import io
import json
import os
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
import openpyxl
from core import SafetyStop, file_hash
from submission_prepare import (KINDS,REQUIRED,zero_roster,templates,crossref_source,validate_metadata,
                                assess,original_export,prepare,fetch_bytes,MetadataClient)


class SubmissionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
        self.path=self.root/'list.xlsx'
        b=openpyxl.Workbook()
        b.active.append(['题名','DOI','匹配到的条目数量'])
        b.active.append(['Target','10.1000/test',0])
        b.active.append(['Target','10.1000/test',1])
        b.active.append(['Target','10.1000/test','0'])
        b.active.append(['Existing','',2])
        b.save(self.path)
        b.close()
        self.tdir=self.root/'templates'
        self.tdir.mkdir()
        for kind in KINDS:
            b=openpyxl.Workbook()
            headers=list(dict.fromkeys(REQUIRED[kind]+['DOI','URL','摘要']))+['发表日期']
            b.active.append(['说明']*len(headers))
            b.active.append(headers)
            b.save(self.tdir/(kind+'.xlsx'))
            b.close()
        self.spec=templates(self.tdir)
        self.paper=zero_roster(self.path)['papers'][0]

    def source(self,paper=None):
        p=paper or self.paper
        return {'id':'C1','provider':'Synthetic','url':'https://example.org/article','kind':'期刊论文','identity_verified':True,
                'fields':{'题名':p['title'],'DOI':p['doi'],'作者':'A(1);B(1)','作者单位':'(1)University','发表日期':'2024-03-01',
                          '发表期刊':'Journal','摘要':'=This must remain text'}}

    def response(self,sources=None):
        s=(sources or [self.source()])[-1]
        return {'kind':'期刊论文','kind_evidence':['C1'],'fields':{k:{'value':v,'source_ids':['C1']} for k,v in s['fields'].items()},'issues':[],'missing':[]}

    def client(self):
        return Mock(fill=Mock(side_effect=lambda p,s,templates,model:self.response(s)))

    def test_only_zero_rows_and_dedup_after_filter(self):
        papers=zero_roster(self.path)['papers']
        self.assertEqual(len(papers),1)
        self.assertEqual(papers[0]['rows'],[2,4])

    def manifest(self,exports):
        path=self.root/'sources.json'
        path.write_text(json.dumps({'exports':{self.paper['id']:exports},'download_hosts':['example.org']}),encoding='utf-8')
        return path

    def test_download_first_skips_metadata_ai_and_key_even_on_resume(self):
        entry={'channel':'WOS','format':'txt','url':'https://example.org/file'}
        manifest=self.manifest([entry])
        fetch=Mock(return_value=(b'DI 10.1000/test\nER\nEF','text/plain'))
        provider=Mock(side_effect=AssertionError('must skip metadata'))
        with patch('submission_prepare.KeyStore',side_effect=AssertionError('must not need key')):
            folder=prepare(self.path,self.tdir,self.root/'out',manifest=manifest,provider=provider,fetch=fetch,progress=lambda _:None)
            prepare(self.path,self.tdir,self.root/'out',manifest=manifest,provider=provider,fetch=fetch,retry_incomplete=True,progress=lambda _:None)
        self.assertEqual(fetch.call_count,1)
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['status'],'原始导出待核验')
        self.assertNotIn('metadata',result)
        Path(result['export']['file']).unlink()
        prepare(self.path,self.tdir,self.root/'out',manifest=manifest,provider=provider,fetch=fetch,progress=lambda _:None)
        self.assertEqual(fetch.call_count,2)

    def test_download_failure_falls_back_to_template(self):
        manifest=self.manifest({'channel':'WOS','format':'xlsx','url':'https://example.org/file'})
        client=self.client()
        folder=prepare(self.path,self.tdir,self.root/'out',manifest=manifest,client=client,provider=self.source,
                       fetch=Mock(return_value=(b'PKbroken','application/octet-stream')),progress=lambda _:None)
        client.fill.assert_called_once()
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['route'],'模板填写')
        self.assertTrue(result['retrieval_notes'])

    def test_additional_template_requires_explicit_key_fields(self):
        from submission_prepare import required_fields
        b=openpyxl.Workbook()
        b.active.append(['说明','说明'])
        b.active.append(['题名','DOI'])
        b.save(self.tdir/'专利.xlsx')
        b.close()
        spec=templates(self.tdir)
        self.assertIn('专利',spec)
        with self.assertRaises(SafetyStop):
            required_fields(spec,{})
        self.assertEqual(required_fields(spec,{'required':{'专利':['题名','DOI']}})['专利'],['题名','DOI'])

    def test_unknown_match_count_is_not_zero(self):
        for value in (None,False,'=0','unknown',-1,0.5):
            b=openpyxl.load_workbook(self.path)
            b.active['C2']=value
            b.save(self.path)
            b.close()
            with self.subTest(value=value),self.assertRaises(SafetyStop):
                zero_roster(self.path)

    def test_original_wos_is_preserved_only_for_zero_rows(self):
        b=openpyxl.load_workbook(self.path)
        b.active['D1']='WOS_ID'
        b.active['D2']='WOS:original-zero'
        b.active['D3']='WOS:nonzero-must-not-merge'
        b.save(self.path)
        b.close()
        paper=zero_roster(self.path)['papers'][0]
        self.assertEqual(paper['original_fields'],{'WOS记录号':'WOS:original-zero'})

    def test_crossref_requires_unique_identity(self):
        record={'title':['Target'],'DOI':'10.1000/test','type':'journal-article','author':[{'given':'A','family':'Name','affiliation':[{'name':'Uni'}]}],
                'published':{'date-parts':[[2024,3,1]]},'container-title':['Journal']}
        fetch=Mock(return_value=(json.dumps({'message':record}).encode(),'application/json'))
        source=crossref_source(self.paper,fetch)
        self.assertEqual(source['fields']['作者'],'A Name(1)')
        self.assertEqual(source['fields']['作者单位'],'(1)Uni')
        record['title']=['Another paper']
        fetch.return_value=(json.dumps({'message':record}).encode(),'application/json')
        with self.assertRaises(SafetyStop):
            crossref_source(self.paper,fetch)

    def test_fabricated_fields_and_citations_rejected(self):
        for field,value in [('作者','Invented'),('DOI','10.1/wrong'),('WOS记录号','WOS:madeup')]:
            response=self.response()
            response['fields'][field]={'value':value,'source_ids':['C1']}
            with self.subTest(field=field),self.assertRaises(SafetyStop):
                validate_metadata(json.dumps(response),self.paper,[self.source()],self.spec)
        response=self.response()
        response['fields']['作者']['source_ids']=['unknown']
        with self.assertRaises(SafetyStop):
            validate_metadata(json.dumps(response),self.paper,[self.source()],self.spec)

    def test_missing_core_fields_conflicts_and_invalid_dates_block_ready(self):
        response=self.response()
        del response['fields']['作者单位']
        self.assertEqual(assess(response,self.paper,[self.source()],REQUIRED)['status'],'待补资料')
        response=self.response()
        response['fields']['发表日期']['value']='2024-02-31'
        self.assertIn('发表日期无效',assess(response,self.paper,[self.source()],REQUIRED)['issues'])
        response=self.response()
        response['fields']['作者']['value']='A(9)'
        self.assertIn('作者与单位编号关联缺失或无效',assess(response,self.paper,[self.source()],REQUIRED)['issues'])

    def test_pipeline_preserves_templates_rows_provenance_and_resume(self):
        before=file_hash(self.path)
        hashes={k:file_hash(v['path']) for k,v in self.spec.items()}
        client=self.client()
        folder=prepare(self.path,self.tdir,self.root/'out',client=client,provider=self.source,progress=lambda _:None)
        self.assertEqual(client.fill.call_count,1)
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['status'],'字段齐备待核验')
        self.assertEqual(result['rows'],[2,4])
        b=openpyxl.load_workbook(folder/'字段齐备'/'期刊论文.xlsx')
        headers=[c.value for c in b.active[2]]
        self.assertEqual(b.active['A1'].value,'说明')
        dates=[i+1 for i,h in enumerate(headers) if h=='发表日期']
        self.assertTrue(all(b.active.cell(3,c).value=='2024-03-01' for c in dates))
        cell=b.active.cell(3,headers.index('摘要')+1)
        self.assertEqual(cell.data_type,'s')
        b.close()
        self.assertEqual(file_hash(self.path),before)
        self.assertEqual({k:file_hash(v['path']) for k,v in self.spec.items()},hashes)
        fresh=self.client()
        provider=Mock(side_effect=AssertionError('should not fetch again'))
        prepare(self.path,self.tdir,self.root/'out',client=fresh,provider=provider,progress=lambda _:None)
        fresh.fill.assert_not_called()

    def test_missing_units_stays_only_in_draft(self):
        source=self.source()
        del source['fields']['作者单位']
        folder=prepare(self.path,self.tdir,self.root/'out',client=self.client(),provider=lambda p:source,progress=lambda _:None)
        for directory,count in [('字段齐备',2),('待补草稿',3)]:
            b=openpyxl.load_workbook(folder/directory/'期刊论文.xlsx')
            self.assertEqual(b.active.max_row,count)
            b.close()

    def test_omitted_but_sourced_important_content_is_completed(self):
        client=self.client()
        response=self.response()
        del response['fields']['摘要']
        client.fill.side_effect=None
        client.fill.return_value=response
        folder=prepare(self.path,self.tdir,self.root/'out',client=client,provider=self.source,progress=lambda _:None)
        saved=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))
        self.assertEqual(saved['records'][0]['metadata']['fields']['摘要']['value'],'=This must remain text')

    def test_original_download_preserves_bytes_rejects_login_wrong_identity_and_host(self):
        entry={'url':'https://example.org/file','format':'txt','channel':'WOS'}
        raw=b'DI 10.1000/test\nER\nEF'
        fetch=Mock(return_value=(raw,'text/plain'))
        saved=original_export(entry,self.paper,self.root,fetch,{'example.org'})
        self.assertEqual(Path(saved['file']).read_bytes(),raw)
        for bad,mime in [(b'<html>login</html>','text/html'),(b'Another paper','text/plain')]:
            fetch.return_value=(bad,mime)
            with self.assertRaises(SafetyStop):
                original_export(entry,self.paper,self.root,fetch,{'example.org'})
        with self.assertRaises(SafetyStop):
            fetch_bytes('https://not-allowed.test/export',{'example.org'})

    def classification(self,channel='WOS',kind='期刊论文'):
        """A saved classification run for this exact roster, as the UI would leave it."""
        folder=self.root/'classification'/'task'
        folder.mkdir(parents=True,exist_ok=True)
        (folder/'分类结果.json').write_text(json.dumps({'source':{'sha256':zero_roster(self.path)['sha256']},
            'records':[{'id':self.paper['id'],'classification':{'type':kind},
                        'import_route':{'recommended_channel':channel}}]}),encoding='utf-8')
        return self.root/'classification'

    def test_most_complete_classification_wins_over_a_newer_partial_run(self):
        from submission_prepare import saved_classification
        sha=zero_roster(self.path)['sha256']
        def write(name,sha256,records):
            folder=self.root/'classification'/name
            folder.mkdir(parents=True,exist_ok=True)
            path=folder/'分类结果.json'
            path.write_text(json.dumps({'source':{'sha256':sha256},'records':records}),encoding='utf-8')
            return path
        write('complete',sha,[{'id':self.paper['id'],'classification':{'type':'期刊论文'}},
                              {'id':'second','classification':{'type':'会议论文'}}])
        # A later, abandoned run must not shadow the complete one just by being newer.
        partial=write('newer-partial',sha,[{'id':'other','classification':{'type':'会议论文'}}])
        os.utime(partial,(9_999_999_999,9_999_999_999))
        # A run for a different roster is never mixed in.
        write('foreign','different-sha',[{'id':self.paper['id'],'classification':{'type':'科技论文'}}])
        saved=saved_classification(sha,self.root/'classification')
        self.assertIn(self.paper['id'],saved)
        self.assertEqual(saved[self.paper['id']]['classification']['type'],'期刊论文')

    def test_classification_preselects_template_and_prefills_the_by_hand_file(self):
        folder=prepare(self.path,self.tdir,self.root/'out',client=Mock(fill=Mock(side_effect=SafetyStop('无模型'))),
                       provider=Mock(side_effect=SafetyStop('无公开元数据')),
                       classification_dir=self.classification(),progress=lambda _:None)
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['template_kind'],'期刊论文')
        self.assertEqual(result['template_kind_source'],'AI 分类建议')
        self.assertIn('AI 分类建议',''.join(result['issues']))
        self.assertIn('未经出版来源核实',''.join(result['issues']))
        # No download and no usable metadata still yields a fillable, prefilled file.
        b=openpyxl.load_workbook(folder/'待补草稿'/'期刊论文.xlsx')
        headers=[c.value for c in b.active[2]]
        self.assertEqual(b.active.cell(3,headers.index('题名')+1).value,'Target')
        self.assertEqual(b.active.cell(3,headers.index('DOI')+1).value,'10.1000/test')
        self.assertIsNone(b.active.cell(3,headers.index('作者')+1).value)
        b.close()

    def test_export_channel_order_follows_the_classified_channel(self):
        manifest=self.manifest([{'channel':'EI','format':'txt','url':'https://example.org/ei'},
                                {'channel':'WOS','format':'txt','url':'https://example.org/wos'}])
        calls=[]
        def fetch(url,hosts,limit=0):
            calls.append(url)
            if url.endswith('/ei'):
                raise SafetyStop('该渠道不可用')
            return b'DI 10.1000/test\nER\nEF','text/plain'
        prepare(self.path,self.tdir,self.root/'out',manifest=manifest,fetch=fetch,
                classification_dir=self.classification('WOS'),progress=lambda _:None)
        self.assertTrue(calls[0].endswith('/wos'),calls)

    def test_inbox_export_is_adopted_matched_and_reported(self):
        inbox=self.root/'inbox'
        inbox.mkdir()
        (inbox/'wos-export.txt').write_text('DI 10.1000/test\nTI Target\nER\nEF',encoding='utf-8')
        (inbox/'unrelated.txt').write_text('DI 10.9999/other\nTI Nothing here',encoding='utf-8')
        provider=Mock(side_effect=AssertionError('a downloaded export must skip metadata and AI'))
        folder=prepare(self.path,self.tdir,self.root/'out',provider=provider,fetch=Mock(),
                       inbox=inbox,progress=lambda _:None)
        provider.assert_not_called()
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['status'],'原始导出待核验')
        self.assertEqual(result['route'],'待收目录导出文件')
        self.assertTrue(Path(result['export']['file']).is_file())
        self.assertEqual(Path(result['export']['file']).read_bytes(),(inbox/'wos-export.txt').read_bytes())
        unmatched=json.loads((folder/'待收未匹配.json').read_text(encoding='utf-8'))
        self.assertEqual([u['file'] for u in unmatched['unmatched']],['unrelated.txt'])
        status=json.loads((folder/'status.json').read_text(encoding='utf-8'))
        self.assertEqual((status['downloaded'],status['inbox_unmatched']),(1,1))

    def test_xlsx_intake_is_resaved_for_the_import_page(self):
        inbox=self.root/'inbox'
        inbox.mkdir()
        book=openpyxl.Workbook()
        book.active.append(['题名','DOI'])
        book.active.append(['Target','10.1000/test'])
        exported=inbox/'CNKI-export.xlsx'
        book.save(exported)
        book.close()
        folder=prepare(self.path,self.tdir,self.root/'out',client=self.client(),provider=self.source,
                       inbox=inbox,progress=lambda _:None)
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['status'],'原始导出待核验')
        # The original bytes are kept, and a re-saved copy is produced for the import page.
        self.assertEqual(Path(result['export']['file']).read_bytes(),exported.read_bytes())
        normalized=Path(result['export']['import_file'])
        self.assertTrue(normalized.is_file())
        reopened=openpyxl.load_workbook(normalized)
        self.assertEqual(reopened.active.cell(2,1).value,'Target')
        reopened.close()
        audit=openpyxl.load_workbook(folder/'提交准备与来源.xlsx')
        headers=[c.value for c in audit['提交准备总表'][1]]
        self.assertEqual(audit['提交准备总表'].cell(2,headers.index('可导入文件')+1).value,str(normalized))
        audit.close()

    def test_txt_intake_gains_no_resaved_copy(self):
        inbox=self.root/'inbox'
        inbox.mkdir()
        (inbox/'wos.txt').write_text('DI 10.1000/test\nTI Target\nER\nEF',encoding='utf-8')
        folder=prepare(self.path,self.tdir,self.root/'out',client=self.client(),provider=self.source,
                       inbox=inbox,progress=lambda _:None)
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertNotIn('import_file',result['export'])
        self.assertEqual(list((folder/'原始导出').glob('另存-*')),[])

    def test_inbox_rejects_html_and_keeps_it_out_of_the_queue(self):
        inbox=self.root/'inbox'
        inbox.mkdir()
        (inbox/'login.txt').write_text('<html><body>DI 10.1000/test sign in</body></html>',encoding='utf-8')
        folder=prepare(self.path,self.tdir,self.root/'out',client=self.client(),provider=self.source,
                       inbox=inbox,progress=lambda _:None)
        result=json.loads((folder/'提交准备.json').read_text(encoding='utf-8'))['records'][0]
        self.assertEqual(result['status'],'字段齐备待核验')
        self.assertNotIn('待收-',json.dumps(result.get('export') or {}))

    def test_model_payload_has_no_workflow_identifiers(self):
        client=MetadataClient(Mock())
        client._request=Mock(return_value={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(self.response())}}]})
        paper={**self.paper,'owner':'private-owner','staff_id':'secret','rows':[2]}
        client.fill(paper,[self.source()],self.spec,'deepseek-chat')
        payload=client._request.call_args.args[2]
        text=json.dumps(payload)
        self.assertNotIn('private-owner',text)
        self.assertNotIn('secret',text)
        self.assertEqual(set(json.loads(payload['messages'][1]['content'])['paper']),{'id','title','doi'})

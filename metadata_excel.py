"""Build and independently verify code-local general-import workbooks."""
import os
from pathlib import Path
import shutil
import subprocess

from core import SafetyStop, file_hash


def _build_log(folder, message):
    from run_diagnostics import safe_text
    try:
        (folder / "excel-build.log").write_text(safe_text(message) + "\n", encoding="utf-8")
    except OSError:
        pass  # Keep the original build error/result; never replay it for logging.


def build_excel(folder, candidate):
    folder = Path(folder).resolve()
    output = folder / "paper.xlsx"
    if not output.exists():
        bundled = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node"
        modules = Path(os.environ.get("SA_ARTIFACT_NODE_MODULES", str(bundled / "node_modules")))
        node = os.environ.get("SA_NODE_EXECUTABLE") or (str(bundled / "bin/node.exe")
                if (bundled / "bin/node.exe").is_file() else shutil.which("node"))
        if not node or not (modules / "@oai/artifact-tool").is_dir():
            raise SafetyStop("通用 Excel 生成运行库未配置；请配置 SA_NODE_EXECUTABLE 与 SA_ARTIFACT_NODE_MODULES。")
        try:
            completed = subprocess.run([node, str(Path(__file__).with_suffix('.mjs')), str(folder / "metadata.json"),
                str(output), str(modules)], capture_output=True, text=True, timeout=90, encoding="utf-8", errors="replace")
        except subprocess.TimeoutExpired as exc:
            _build_log(folder, "Node 生成超过 90 秒：" + str(exc.stderr or ""))
            raise SafetyStop("通用 Excel 生成超时（90 秒），未上传；诊断保存在同篇目录 excel-build.log。") from None
        except OSError as exc:
            _build_log(folder, type(exc).__name__ + ": " + str(exc))
            raise SafetyStop("通用 Excel 运行库无法启动，未上传；诊断保存在同篇目录 excel-build.log。") from None
        _build_log(folder, f"exit_code={completed.returncode}\n{completed.stdout}\n{completed.stderr}")
        if completed.returncode or not output.is_file():
            raise SafetyStop("通用 Excel 生成失败，未上传；具体原因保存在同篇目录 excel-build.log。")
    from openpyxl import load_workbook
    book = load_workbook(output, read_only=True, data_only=False)
    try:
        if len(book.worksheets) != 1:
            raise SafetyStop("通用 Excel 工作表数量不符。")
        sheet = book.active
        rows = list(sheet.iter_rows())
        expected = [list(candidate["fields"]), list(candidate["fields"].values())]
        # OOXML dimension is optional (Artifact Tool legitimately omits it).
        if len(rows) != 2 or any(len(row) != len(expected[0]) for row in rows):
            raise SafetyStop("通用 Excel 不是恰好一篇文献。")
        for cells, values in zip(rows, expected):
            if [c.value or "" for c in cells] != values or any(c.data_type == "f" for c in cells):
                raise SafetyStop("通用 Excel 与核验元数据不一致，未上传。")
    finally:
        book.close()
    if output.is_symlink() or output.stat().st_size > 524288:
        raise SafetyStop("通用 Excel 路径或大小异常。")
    return output, file_hash(output)

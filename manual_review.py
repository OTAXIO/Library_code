"""Explicit, single-row LOCAL verification; no browser or workflow-journal writes."""
from dataclasses import replace
from pathlib import Path
import copy
import os
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET

from core import SafetyStop, file_hash, read_roster
from roster_annotations import OWNER, YellowStyles, _yellow_id
from roster_write import (NS, RosterUpdate, _commit_verified, column_name, patch_cell,
                          worksheet_member, write_lock)

STATES = {'done': ('已完成', 1), 'pending': ('未完成', None), 'skipped': ('已跳过', 2)}


def state_of(record):
    return 'done' if record.done else 'skipped' if record.skipped else 'pending'


def select_review(roster, owner, sa_id):
    if not roster:
        raise SafetyStop('请先读取 code/list.xlsx。')
    if owner != OWNER:
        raise SafetyStop('手动核验仅修改谭勋策的名单。')
    selected = [r for r in roster.records if r.sa_id == sa_id]
    if not sa_id or len(selected) != 1 or selected[0].owner != OWNER:
        raise SafetyStop('请选择一条 ID 唯一的谭勋策记录。')
    roster.assert_unchanged()
    return selected[0]


def write_review(roster, record, state, backup_dir=None):
    """Atomically save flag AND ID fill, preserving every note and business field.

    This is a human decision, not proof of a new admission or remote closure.
    Automatic completion APIs retain their independent evidence requirements.
    """
    if not isinstance(state, str) or state not in STATES:
        raise SafetyStop('核验状态只能是已完成、未完成或已跳过。')
    if select_review(roster, OWNER, record.sa_id) != record:
        raise SafetyStop('所选记录已变化，请重读名单。')
    if (roster.path.name.lower() != 'list.xlsx' or not roster.status_separate
            or not roster.completion_column or not roster.remark_column
            or roster.completion_column == roster.remark_column):
        raise SafetyStop('请重读名单，确认独立的是否识别列。')
    value = STATES[state][1]
    path, temporary = roster.path, None
    try:
        with write_lock(path):
            roster.assert_unchanged()
            if path.with_name('~$' + path.name).exists():
                raise SafetyStop('请保存并关闭 Excel，再保存核验结果。')
            with zipfile.ZipFile(path) as original:
                member = worksheet_member(original, roster.sheet_name)
                sheet = changed = original.read(member)
                # Discover the original ID column; never assume it is column C.
                from openpyxl import load_workbook
                from openpyxl.utils.cell import range_boundaries
                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    labels = [c.value for c in next(book[roster.sheet_name].iter_rows(min_row=1, max_row=1))]
                finally:
                    book.close()
                if labels.count('sa_lzk表ID') != 1:
                    raise SafetyStop('ID 表头不唯一，未保存核验。')
                id_col = labels.index('sa_lzk表ID') + 1
                id_ref = column_name(id_col) + str(record.row)
                flag_ref = column_name(roster.completion_column) + str(record.row)
                document = ET.fromstring(sheet)
                for cell in document.findall(f'{{{NS}}}mergeCells/{{{NS}}}mergeCell'):
                    left, top, right, bottom = range_boundaries(cell.get('ref'))
                    if left <= id_col <= right and top <= record.row <= bottom:
                        raise SafetyStop('ID 单元格已合并，未保存核验。')
                # patch_cell validates worksheet protection even when clearing a blank.
                flag = document.find(f'{{{NS}}}sheetData/{{{NS}}}row[@r="{record.row}"]/{{{NS}}}c[@r="{flag_ref}"]')
                if value is None and flag is None:
                    # Validate protection without adding an empty flag cell.
                    patch_cell(sheet, flag_ref, record.row, 1)
                else:
                    changed = patch_cell(changed, flag_ref, record.row, value)
                flag_changed = changed != sheet
                styles = YellowStyles(original.read('xl/styles.xml'),
                                      'FFFF0000' if state == 'skipped' else 'FFFFFF00')
                changed, color_changed = _yellow_id(changed, id_ref, styles, pending=state == 'pending')
                if changed == sheet:
                    return RosterUpdate(roster, None, (), {record.sa_id: state_of(record)})
                ET.fromstring(changed)
                ET.fromstring(styles.raw)
                fd, name = tempfile.mkstemp(prefix='.list-review-', suffix='.xlsx', dir=path.parent)
                os.close(fd)
                temporary = Path(name)
                with zipfile.ZipFile(temporary, 'w') as output:
                    output.comment = original.comment
                    for entry in original.infolist():
                        data = (changed if entry.filename == member else styles.raw.encode('utf-8')
                                if entry.filename == 'xl/styles.xml' else original.read(entry))
                        output.writestr(copy.copy(entry), data)
            verified = read_roster(temporary)
            expected = [replace(r, done=state == 'done', skipped=state == 'skipped')
                        if r == record else r for r in roster.records]
            if verified.records != expected:
                raise SafetyStop('核验回读发现备注或其他任务变化，原名单未修改。')
            backups = Path(backup_dir) if backup_dir else path.parent / 'runtime/backups'
            backups.mkdir(parents=True, exist_ok=True)
            backup = backups / ('list-' + roster.sha256 + '.xlsx')
            if not backup.exists():
                with path.open('rb') as source, backup.open('xb') as target:
                    shutil.copyfileobj(source, target)
                    target.flush()
                    os.fsync(target.fileno())
            if file_hash(backup) != roster.sha256:
                raise SafetyStop('名单备份校验失败，未保存核验。')
            with temporary.open('rb+') as stream:
                os.fsync(stream.fileno())
            _commit_verified(temporary, roster, verified.sha256)
            temporary = None
            verified.path = path
            verified.mtime_ns = path.stat().st_mtime_ns
            verified.assert_unchanged()
            cells = ((flag_ref,) if flag_changed else ()) + ((id_ref,) if color_changed else ())
            return RosterUpdate(verified, backup, cells, {record.sa_id: state_of(record)})
    except PermissionError as exc:
        raise SafetyStop('名单被占用，未保存核验；请关闭 Excel 后重读。') from exc
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)

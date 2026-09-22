"""Read-only roster ingestion, review rules and durable local checkpoints."""
from __future__ import annotations

import hashlib
import json
import re
import sqlite3
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path


class SafetyStop(Exception):
    """A condition requiring explicit human attention, never an automatic retry."""


HEADERS = {
    "owner": "负责人", "sa_id": "sa_lzk表ID", "title": "题名",
    "doi": "DOI", "wos": "WOS_ID", "staff_id": "工号",
    "matches": "匹配到的条目数量", "item_ids": "平台唯一号",
    "mark": "标记状态", "reason": "标记为待处理原因",
}
QUERY_HEADER = "查询方式（DOI和WOS_ID：1   题名： 2）"


def text(value):
    return "" if value is None else str(value).strip()


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def file_hash(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


@dataclass(frozen=True)
class Record:
    row: int
    owner: str
    sa_id: str
    title: str
    doi: str
    wos: str
    staff_id: str
    matches: int
    item_ids: str
    mark: str
    reason: str
    query: str

    @property
    def key(self):
        # Changed task facts invalidate an old completion, even for the same ID.
        return digest(asdict(self))


@dataclass
class Roster:
    path: Path
    sha256: str
    records: list[Record]
    mtime_ns: int

    def assert_unchanged(self):
        if self.path.stat().st_mtime_ns != self.mtime_ns or file_hash(self.path) != self.sha256:
            raise SafetyStop("名单文件已改变。请重新读取名单，再核验当前记录。")


def latest_roster(folder):
    candidates = [p for p in Path(folder).glob("数据比对结果*.xlsx")
                  if p.is_file() and not p.name.startswith("~$")]
    if not candidates:
        raise SafetyStop("未找到「数据比对结果*.xlsx」，请用“选择名单”指定文件。")
    # Latest means most recently modified, not largest filename/date range.
    return max(candidates, key=lambda p: (p.stat().st_mtime_ns, p.name))


def read_roster(path):
    from openpyxl import load_workbook
    path = Path(path).resolve()
    stamp, checksum = path.stat().st_mtime_ns, file_hash(path)
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        choices = []
        for sheet in book:
            row = next(sheet.iter_rows(min_row=1, max_row=1), ())
            labels = [text(c.value) for c in row]
            if "sa_lzk表ID" in labels:
                choices.append((sheet, labels))
        if len(choices) != 1:
            raise SafetyStop("应恰好有一个包含 sa_lzk表ID 表头的工作表。")
        sheet, labels = choices[0]
        mapping = {}
        for key, label in HEADERS.items():
            if labels.count(label) != 1:
                raise SafetyStop(f"缺少或重复必要表头：{label}")
            mapping[key] = labels.index(label)
        query_cols = [i for i, label in enumerate(labels) if label.startswith("查询方式")]
        if len(query_cols) != 1:
            raise SafetyStop("无法唯一识别查询方式列。")
        mapping["query"] = query_cols[0]
        records, seen = [], set()
        for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if all(c.value is None for c in cells):
                continue
            values = {}
            for key, index in mapping.items():
                cell = cells[index]
                if cell.data_type == "f":
                    raise SafetyStop(f"第 {number} 行 {labels[index]} 是公式，请先人工确认并转为文本值。")
                if key in ("sa_id", "staff_id", "item_ids") and cell.value is not None and not isinstance(cell.value, str):
                    raise SafetyStop(f"第 {number} 行 {labels[index]} 不是文本，可能丢失前导零或长编号精度。")
                values[key] = text(cell.value)
            sa_id = values["sa_id"]
            if not sa_id or len(sa_id) > 160 or re.search(r"[\x00-\x1f]", sa_id):
                raise SafetyStop(f"第 {number} 行名单 ID 无效。")
            if sa_id in seen:
                raise SafetyStop(f"第 {number} 行名单 ID 重复，停止建立队列。")
            seen.add(sa_id)
            try:
                count = int(values["matches"])
            except ValueError as exc:
                raise SafetyStop(f"第 {number} 行匹配数量不是整数。") from exc
            ids = [x.strip() for x in values["item_ids"].strip(",").split(",") if x.strip()]
            if count < 0 or count != len(ids) or len(ids) != len(set(ids)):
                raise SafetyStop(f"第 {number} 行匹配数量与平台唯一号不一致。")
            if not values["owner"]:
                values["owner"] = "（未分配）"
            values["matches"] = count
            records.append(Record(row=number, **values))
        if not records:
            raise SafetyStop("名单没有记录。")
    finally:
        book.close()
    result = Roster(path, checksum, records, stamp)
    result.assert_unchanged()
    return result


def guide(record):
    lines = []
    if record.query not in {"1", "2"}:
        lines.append(f"【未知情况】查询方式={record.query or '空'}，PPT 未定义。先向负责人确认含义，记录证据后人工决定。")
    if record.matches == 0:
        lines.extend([
            "【无匹配】核验交大成果归属、题名和发表情况；前台找到同一论文后填写平台唯一号。",
            "有原文/数据库记录但本库缺失：人工导入，说明填 SA补充-当前名单ID；确认查重与合并设置后推送。",
            "未查到或未发表：记录“未查询到该文献（平台不处理）”，先确认状态规则；不得直接等同已处理。",
        ])
    elif record.matches > 1:
        lines.append(f"【多条匹配：{record.matches}】逐条核验是否同一文献；只有确认重复后才能人工选择主条目合并。工具不自动合并/删除。")
    else:
        lines.append("【单条匹配】打开对比详情，结合论文原文逐项核验；红色差异不等于本库错误。")
    if "作者不一致" in record.reason:
        lines.append("作者不一致：按工号核实学者身份，必要时在原网页认领；别名新增必须核对身份。")
    if "第一作者" in record.reason or "通讯作者" in record.reason:
        lines.append("作者/第一单位标记：核对原文署名和单位；SA 正确则在本库编辑，本库正确则说明依据，不反向修改本库。")
    if "DOI" in record.reason or "WOS" in record.reason:
        lines.append("DOI/WOS 差异：先证明是同一篇论文，再核对正确标识；不能仅凭题名相似直接关闭。")
    lines.append("完成前：保存原文/页面证据、核验修改后的页面，再填写备注并设置已处理。遇到登录、报错或不认识的控件，暂停人工处理。")
    return "\n\n".join(lines)


class Journal:
    def __init__(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(path)
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.executescript("""
            CREATE TABLE IF NOT EXISTS progress (
                task_key TEXT PRIMARY KEY, sa_id TEXT NOT NULL, state TEXT NOT NULL,
                note TEXT NOT NULL, updated TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT, time TEXT NOT NULL,
                task_key TEXT NOT NULL, action TEXT NOT NULL, detail TEXT NOT NULL);
        """)

    def state(self, record):
        value = self.db.execute("SELECT state FROM progress WHERE task_key=?", (record.key,)).fetchone()
        return value[0] if value else "待开始"

    def save(self, record, state, note="", detail=None):
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        with self.db:
            self.db.execute("INSERT OR REPLACE INTO progress VALUES (?,?,?,?,?)",
                            (record.key, record.sa_id, state, note, now))
            self.db.execute("INSERT INTO events(time,task_key,action,detail) VALUES (?,?,?,?)",
                            (now, record.key, state, json.dumps(detail or {"note": note}, ensure_ascii=False)))

    def close(self):
        self.db.close()
